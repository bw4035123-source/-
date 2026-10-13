"""다이수 테스트: python -m unittest discover -s tests"""
from __future__ import annotations

import datetime as dt
import glob
import io
import math
import os
import shutil
import sys
import tempfile
import time
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ.setdefault("DAISU_WORKSPACE", tempfile.mkdtemp(prefix="daisu_test_"))

from daisu import config, engine, qa, reports, schedule  # noqa: E402
from daisu.readers import read  # noqa: E402
from daisu.render import FORMATS  # noqa: E402
from daisu.sources import doc_kinds  # noqa: E402

BASE = config.SAMPLE


def _bytes(p):
    with open(p, "rb") as f:
        return f.read()


def sample_project(overrides=None, base="2026-10-12"):
    files = [engine.file_entry(p, os.path.relpath(p, BASE)) for p in sorted(glob.glob(BASE + "/**/*.*", recursive=True))]
    return {"id": "test0000", "files": files, "base_date": base, "overrides": overrides or {"person": {}, "course": {}}}


class TestReading(unittest.TestCase):
    def test_kinds(self):
        kinds = {}
        for p in glob.glob(BASE + "/**/*.*", recursive=True):
            kinds[os.path.relpath(p, BASE)] = doc_kinds(read(p))
        got = sorted(k for ks in kinds.values() for k in ks)
        self.assertEqual(got.count("roster"), 1)
        self.assertEqual(got.count("records"), 1)
        self.assertEqual(got.count("sign"), 3)  # 한글·엑셀·워드 서명부
        self.assertEqual(got.count("cert"), 3)
        self.assertEqual(got.count("shift"), 1)
        self.assertEqual(got.count("survey"), 1)
        self.assertEqual(got.count("guide"), 1)


class TestEngine(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res = engine.compute(sample_project())
        cls.sm = {s["key"]: s for s in cls.res["summary"]}
        cls.by_name = {}
        for p in cls.res["people"]:
            cls.by_name.setdefault(p["name"], []).append(p)

    def cell(self, pid, key):
        return self.res["cells"][pid][key]

    def test_planted_findings(self):
        self.assertEqual(self.sm["sr@H1"]["overdue"], 12)   # 상반기 현장 안전교육: 한 회차만 듣거나 불참
        self.assertEqual(self.sm["sr@H1"]["partial"], 7)
        self.assertEqual(self.sm["hire"]["overdue"], 1)     # 10/5 입사자 채용 시 교육 없음
        types = sorted(i["type"] for i in self.res["issues"])
        self.assertEqual(types, ["ambiguous", "course", "unknown"])
        unk = next(i for i in self.res["issues"] if i["type"] == "unknown")
        self.assertEqual(unk["candidates"][0]["name"], "이영희")  # 오타 후보 제시

    def test_combined_course_split(self):
        """4대폭력 통합 4시간 → 성희롱·성매매·성폭력·가정폭력 각 1시간으로 인정."""
        rec = next(r for r in self.res["records"] if "통합" in r["title"] and r["done"] and r["pid"])
        for k in ("sh", "pr", "sv", "dv"):
            c = self.cell(rec["pid"], k)
            self.assertEqual(c["s"], "done")
            self.assertIn(rec["rid"], c["recs"])
            self.assertAlmostEqual(c["h"], 1.0)

    def test_partial_hours(self):
        short = [r for r in self.res["records"] if r["title"].startswith("청렴 첫걸음")]
        self.assertTrue(short)
        c = self.cell(short[0]["pid"], "int")
        self.assertEqual((c["s"], c["h"], c["need"]), ("partial", 1.0, 2))

    def test_job_hours_and_supervisor(self):
        field = next(p for p in self.res["people"] if p["job"] == "현장" and not p["flags"]["관리감독자"])
        office = next(p for p in self.res["people"] if p["job"] == "사무")
        self.assertEqual(self.cell(field["id"], "sr@H2")["need"], 12)
        self.assertEqual(self.cell(office["id"], "sr@H2")["need"], 6)
        sup = [p for p in self.res["people"] if p["flags"]["관리감독자"]]
        self.assertEqual(len(sup), 9)
        for p in sup:
            self.assertEqual(self.cell(p["id"], "sr@H2")["s"], "na")  # 관리감독자는 별도 과정
        certs = [r for r in self.res["records"] if r["kind"] == "cert"]
        self.assertTrue(all(r["pid"] and r["rules"] == ["sup"] for r in certs))
        for r in certs:
            self.assertEqual(self.cell(r["pid"], "sup")["s"], "done")

    def test_leave_left_flags(self):
        leave = next(p for p in self.res["people"] if p.get("leave") and p["leave"][0] == "2026-03-01")
        self.assertEqual(self.cell(leave["id"], "int")["s"], "leave")
        left = next(p for p in self.res["people"] if p.get("left"))
        self.assertTrue(all(c["s"] == "left" for c in self.res["cells"][left["id"]].values()))
        nonpi = next(p for p in self.res["people"] if not p["flags"]["개인정보취급자"])
        self.assertEqual(self.cell(nonpi["id"], "pi")["s"], "na")

    def test_unsigned_not_credited(self):
        unsigned = [r for r in self.res["records"] if r["kind"] == "sign" and not r["done"]]
        self.assertTrue(unsigned)
        for r in unsigned:
            for key, c in self.res["cells"].get(r["pid"], {}).items():
                self.assertNotIn(r["rid"], c["recs"])

    def test_overrides(self):
        amb = next(i for i in self.res["issues"] if i["type"] == "ambiguous")
        target = amb["candidates"][0]["id"]
        before = self.cell(target, "sr@H1")["h"]
        course = next(i for i in self.res["issues"] if i["type"] == "course")
        res2 = engine.compute(sample_project({"person": {amb["key"]: target}, "course": {course["key"]: []}}))
        self.assertEqual(res2["cells"][target]["sr@H1"]["h"], before + 4)
        self.assertEqual([i["type"] for i in res2["issues"]], ["unknown"])

    def test_base_date_changes_overdue(self):
        before = {s["key"]: s for s in engine.compute(sample_project(base="2026-06-29"))["summary"]}
        after = {s["key"]: s for s in engine.compute(sample_project(base="2026-07-01"))["summary"]}
        self.assertEqual(before["sr@H1"]["overdue"], 0)  # 상반기 기한(6/30) 전에는 지남 아님
        self.assertEqual(after["sr@H1"]["overdue"], 12)  # 지나면 기한 지남
        prj = sample_project(base="2027-01-20")
        self.assertEqual(engine.compute(prj)["summary"][0]["done"], 0)  # 새해에는 새 연도 기준(지난해 기록 미인정)
        prj["year"] = 2026
        last = {s["key"]: s for s in engine.compute(prj)["summary"]}
        self.assertEqual(last["sh"]["done"], 107)  # 판정 연도를 2026으로 두면 지난해 실적 그대로
        self.assertEqual(last["sh"]["overdue"], last["sh"]["none"])  # 못 마친 사람은 모두 기한 지남


class TestSchedule(unittest.TestCase):
    def test_constraints(self):
        res = engine.compute(sample_project())
        st = config.settings()["schedule"]
        idx, days = schedule._shift_index(res["shifts"], res["people"])
        people = {p["id"]: p for p in res["people"]}
        for key in ("sr@H2", "dis", "int"):
            pl = schedule.plan(res, key)
            self.assertTrue(pl["sessions"])
            seen = {}
            for s in pl["sessions"]:
                d = dt.date.fromisoformat(s["date"])
                self.assertLess(d.weekday(), 5)
                self.assertNotIn(s["date"], st["holidays"])
                self.assertLessEqual(s["count"], st["room_capacity"])
                for a in s["people"]:
                    p = people[a["id"]]
                    if p["shift"] == "교대":
                        self.assertIn(idx[p["id"]][s["date"]], st["day_shift_codes"])  # 비번·휴무·야간일 배정 없음
                    self.assertNotIn((a["id"], s["date"]), seen)  # 하루 한 회차
                    seen[(a["id"], s["date"])] = 1
                for dept, n in s["by_dept"].items():
                    on = sum(1 for p in people.values() if p["dept"] == dept and not p.get("left") and
                             (p["shift"] != "교대" or idx.get(p["id"], {}).get(s["date"]) in st["day_shift_codes"]))
                    self.assertLessEqual(n, max(1, math.floor(on * (1 - st["min_on_duty_ratio"]))))
            if key == "dis":
                self.assertFalse(pl["unplaced"])
                self.assertGreater(pl["naive"]["absent"], 0)


class TestReports(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res = engine.compute(sample_project())
        cls.org = config.settings()["org"]

    def test_documents_roundtrip(self):
        pl = schedule.plan(self.res, "dis")
        sign = next(s for s in self.res["signs"] if "2회차" in s["title"])
        d = tempfile.mkdtemp()
        for name, blocks, must in (("status", reports.status_blocks(self.res, self.org, pl), "교육별 이수 현황"),
                                   ("result", reports.result_blocks(self.res, self.org, sign), "만족도"),
                                   ("plan", reports.plan_blocks(self.res, self.org, pl), "회차")):
            for fmt, (fn, _) in FORMATS.items():
                path = os.path.join(d, f"{name}.{fmt}")
                with open(path, "wb") as f:
                    f.write(fn(blocks))
                if fmt in ("hwpx", "docx"):
                    doc = read(path)
                    text = " ".join(t for t, _ in doc.texts) + " ".join(" ".join(" ".join(r) for r in t.rows) for t in doc.tables)
                    self.assertIn(must, text, (name, fmt))
        shutil.rmtree(d)

    def test_survey_and_improvements(self):
        sign = next(s for s in self.res["signs"] if "2회차" in s["title"])
        sv = reports.survey_summary(reports.survey_for(self.res, sign["title"] + " " + sign["context"]))
        self.assertEqual(sv["n"], 46)
        self.assertIn("시간", sv["lowest"]["item"])
        imp = reports.improvements(sv)
        self.assertTrue(any("4시간 이내" in x for x in imp))
        first = next(s for s in self.res["signs"] if "1회차" in s["title"])
        self.assertEqual(reports.survey_for(self.res, first["title"] + " " + first["context"]), [])  # 다른 회차 응답 섞지 않음
        dis = next(s for s in self.res["signs"] if "장애인" in s["title"])
        imp = reports.improvements(reports.survey_summary(reports.survey_for(self.res, dis["title"])))
        self.assertFalse(any("4시간 이내" in x for x in imp))  # “짧고 좋았다”는 불만이 아님

    def test_reminders_excel_evidence(self):
        msgs = reports.reminders(self.res, self.org)
        over = [m for m in msgs if m["overdue"]]
        self.assertEqual(len(over), 13)
        self.assertIn("기한이 지났습니다", over[0]["text"])
        self.assertIn("(목)", next(m for m in msgs if not m["overdue"])["text"])  # 2026-12-31은 목요일
        wb = reports.excel(self.res)
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(io.BytesIO(wb)).sheetnames, ["교육별 요약", "이수현황", "미이수자", "판정 근거", "확인 필요"])
        files = {f["rel"]: f["path"] for f in sample_project()["files"]}
        z = zipfile.ZipFile(io.BytesIO(reports.evidence_zip(self.res, files)))
        names = z.namelist()
        self.assertIn("00_증빙목록.xlsx", names)
        self.assertTrue(any("수료증_" in n and "관리감독자" in n for n in names))


class TestQA(unittest.TestCase):
    def test_answers(self):
        res = engine.compute(sample_project())
        a = qa.ask(res, "김민수 뭐 남았어?")
        self.assertEqual(a["mode"], "person")
        self.assertIn("같은 이름이 여러 명", a["answer"])
        a = qa.ask(res, "기간제도 청렴교육 대상이야?")
        self.assertIn("기간제 근로자를 포함한 전 직원", a["answer"])
        self.assertIn("제81조의2", a["answer"])
        self.assertEqual(qa.ask(res, "체육시설팀 미이수자")["mode"], "dept")
        self.assertEqual(qa.ask(res, "급한 거 알려줘")["mode"], "due")
        self.assertEqual(qa.ask(res, "주차 등록 방법")["mode"], "none")  # 근거 없으면 모른다고


class TestServer(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient

        from daisu.server import app
        cls.c = TestClient(app)

    def test_security(self):
        c = self.c
        self.assertEqual(c.post("/api/settings", headers={"Origin": "https://evil.example", "Content-Type": "application/json"}, content="{}").status_code, 403)
        self.assertEqual(c.post("/api/settings", headers={"Content-Type": "text/plain"}, content='{"org":{"name":"X"}}').status_code, 415)
        self.assertEqual(c.get("/api/info", headers={"Host": "evil.example"}).status_code, 403)
        self.assertEqual(c.post("/api/projects/sample", headers={"Origin": "null"}).status_code, 403)
        s = c.get("/api/settings").json()
        self.assertTrue(all("api_key" not in v for v in s["profiles"].values()))
        self.assertEqual(c.get("/api/projects/../../etc").status_code, 404)

    def test_flow(self):
        c = self.c
        pid = c.post("/api/projects/sample").json()["id"]
        d = c.get(f"/api/projects/{pid}").json()
        self.assertTrue(d["integrity"])
        self.assertEqual(len(d["issues"]), 3)
        amb = next(i for i in d["issues"] if i["type"] == "ambiguous")
        self.assertEqual(c.post(f"/api/projects/{pid}/resolve", json={"type": "person", "key": amb["key"], "value": "없는사람"}).status_code, 400)
        c.post(f"/api/projects/{pid}/resolve", json={"type": "person", "key": amb["key"], "value": amb["candidates"][1]["id"]})
        self.assertEqual(len(c.get(f"/api/projects/{pid}").json()["issues"]), 2)
        c.delete(f"/api/projects/{pid}/resolve", params={"type": "person", "key": amb["key"]})
        self.assertEqual(len(c.get(f"/api/projects/{pid}").json()["issues"]), 3)
        self.assertEqual(c.patch(f"/api/projects/{pid}", json={"base_date": "2026-02-30"}).status_code, 400)
        self.assertEqual(c.get(f"/api/projects/{pid}/plan", params={"col": "없음"}).status_code, 400)
        for q in ("doc=status&fmt=hwpx&col=sr@H2", "doc=result&fmt=docx&sign=0", "doc=plan&fmt=md&col=int", "doc=excel",
                  "doc=evidence", "doc=reminders&within=30"):
            r = c.get(f"/api/projects/{pid}/export?{q}")
            self.assertEqual(r.status_code, 200, q)
            self.assertGreater(len(r.content), 500, q)
        self.assertEqual(c.get(f"/api/projects/{pid}/export?doc=result&sign=99").status_code, 400)
        self.assertEqual(c.post(f"/api/projects/{pid}/ask", json={"q": ""}).status_code, 400)
        self.assertEqual(len(c.get(f"/api/projects/{pid}/survey").json()), 3)
        rec = c.get(f"/api/projects/{pid}/records", params={"person": d["people"][0]["id"]}).json()
        self.assertTrue(rec and all("file" in r for r in rec))
        c.delete(f"/api/projects/{pid}")

    def test_upload_filters_and_path(self):
        c = self.c
        roster = glob.glob(BASE + "/01_인사/*.xlsx")[0]
        cyber = glob.glob(BASE + "/02_이수기록/*.xlsx")[0]
        files = [("files", ("01_인사/명단.xlsx", _bytes(roster), "application/octet-stream")),
                 ("files", ("02/이수.xlsx", _bytes(cyber), "application/octet-stream")),
                 ("files", ("사진.jpg", b"\xff\xd8\xff", "image/jpeg")),
                 ("files", ("../../escape.xlsx", _bytes(cyber), "application/octet-stream"))]
        r = c.post("/api/projects/upload", files=files, data={"name": "업로드", "base_date": "2026-10-12"})
        self.assertEqual(r.status_code, 200)
        pid = r.json()["id"]
        d = c.get(f"/api/projects/{pid}").json()
        self.assertEqual(sorted(f["rel"] for f in d["files"]), ["01_인사/명단.xlsx", "02/이수.xlsx", "escape.xlsx"])
        self.assertEqual(len(d["people"]), 120)
        self.assertEqual(c.post("/api/projects/upload", files=[("files", ("a.jpg", b"x", "image/jpeg"))]).status_code, 400)
        r = c.post("/api/projects/path", json={"path": BASE})
        self.assertEqual(len(c.get(f"/api/projects/{r.json()['id']}").json()["files"]), 11)
        self.assertEqual(c.post("/api/projects/path", json={"path": "/없는/폴더"}).status_code, 400)

    def test_rules_edit(self):
        c = self.c
        try:
            rules = c.get("/api/rules").json()
            bad = dict(rules, courses=[dict(rules["courses"][0], hours=-1)])
            self.assertEqual(c.put("/api/rules", json=bad).status_code, 400)
            dup = dict(rules, courses=rules["courses"] + [rules["courses"][0]])
            self.assertEqual(c.put("/api/rules", json=dup).status_code, 400)
            changed = dict(rules, courses=[dict(x, hours=3) if x["id"] == "int" else x for x in rules["courses"]])
            self.assertEqual(c.put("/api/rules", json=changed).status_code, 200)
            res = engine.compute(sample_project())
            self.assertTrue(all(row["int"]["need"] in (0, 3) for row in res["cells"].values()))
        finally:
            c.post("/api/rules/reset")
        self.assertFalse(os.path.exists(os.path.join(config.CONFIG, "rules.local.json")))


if __name__ == "__main__":
    unittest.main()


PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


class TestOperations(unittest.TestCase):
    def test_fees(self):
        from daisu import fees
        cfg = config.settings()["fees"]
        r = fees.compute({"grade": "일반1", "income": "기타소득"}, 1, cfg)
        self.assertEqual((r["total"], r["withholding"], r["net"]), (250000, 22000, 228000))  # 실효 8.8%
        r = fees.compute({"grade": "일반2", "income": "기타소득"}, 1, cfg)
        self.assertEqual(r["withholding"], 13200)
        r = fees.compute({"grade": "내부", "income": "기타소득"}, 1, cfg)
        self.assertEqual(r["withholding"], 0)  # 10만원: 기타소득금액 4만원 ≤ 5만원 과세최저한
        r = fees.compute({"grade": "일반2", "income": "사업소득", "pages": 5, "travel": 30000}, 4, cfg)
        self.assertEqual((r["lecture"], r["income_tax"], r["local_tax"], r["total"]), (515000, 15450, 1540, 545000))

    def test_gongmun_numbering(self):
        from daisu import gongmun
        b = gongmun.blocks({"name": "기관"}, "제목", ["가", ["나", "다", ["라"]], "마"], ["붙임1 1부.", "붙임2 1부."])
        text = [x[1] for x in b if x[0] == "p"]
        self.assertIn("제목  제목", text)
        self.assertIn("  가. 나", text)
        self.assertIn("    1) 라", text)
        self.assertIn("2. 마", text)
        self.assertTrue(any(t.endswith("붙임2 1부.  끝.") for t in text))
        b = gongmun.blocks({"name": "기관"}, "제목", ["하나"])
        self.assertIn("1. 하나  끝.", [x[1] for x in b if x[0] == "p"])

    def test_josa(self):
        from daisu.util import josa
        self.assertEqual(josa("「청렴교육」"), "「청렴교육」을")
        self.assertEqual(josa("「교육 1회차」"), "「교육 1회차」를")

    def test_planning(self):
        from daisu import planning
        prj = sample_project()
        res = engine.compute(prj)
        st = planning.stats(prj, res, config.settings()["fees"])
        self.assertEqual(st["people"], 119)
        self.assertGreater(st["avg_hours"], 5)
        tasks = planning.tasks(prj, res, {"first_days": 30, "boss_days": 14})
        titles = " ".join(t["title"] for t in tasks)
        self.assertIn("확인 필요 3건", titles)
        self.assertIn("기한(6. 30.(화)) 지난", titles)
        self.assertIn("채용 시 교육", titles)
        md = __import__("daisu.render", fromlist=["to_md"]).to_md(planning.next_plan_blocks(prj, res, config.settings()["org"], config.settings()["fees"])).decode()
        self.assertIn("2027년 법정의무교육 운영 계획(안)", md)
        self.assertIn("기한 지난 미이수 12명", md)
        self.assertIn("4시간 이내", md)        # 만족도 의견 반영
        self.assertNotIn("하반기): 현장직 이수율", md)  # 진행 중인 반기는 문제로 잡지 않음
        card = planning.card_blocks(res, res["people"][0]["id"], config.settings()["org"])
        self.assertEqual(card[0][0], "title")


class TestSessionFlow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient

        from daisu.server import app
        cls.c = TestClient(app)

    def test_flow(self):
        c = self.c
        pid = c.post("/api/projects/sample").json()["id"]
        ids = c.post(f"/api/projects/{pid}/sessions", json={"col": "dis", "place": "대회의실"}).json()["ids"]
        self.assertGreaterEqual(len(ids), 2)
        sid = ids[0]
        s = c.get(f"/api/projects/{pid}/sessions/{sid}").json()
        person = s["targets"][0]["id"]
        before = c.get(f"/api/projects/{pid}").json()["cells"][person]["dis"]["s"]
        self.assertIn(before, ("none", "partial"))
        self.assertEqual(c.post(f"/api/projects/{pid}/sessions/{sid}/attend", json={"person": person, "sign": "data:image/png;base64,AAAA"}).status_code, 400)
        self.assertEqual(c.post(f"/api/projects/{pid}/sessions/{sid}/attend", json={"person": person}).status_code, 400)
        self.assertEqual(c.post(f"/api/projects/{pid}/sessions/{sid}/attend", json={"person": "없음", "sign": PNG}).status_code, 400)
        self.assertTrue(c.post(f"/api/projects/{pid}/sessions/{sid}/attend", json={"person": person, "sign": PNG}).json()["ok"])
        self.assertTrue(c.post(f"/api/projects/{pid}/sessions/{sid}/attend", json={"person": person, "sign": PNG}).json()["already"])
        d = c.get(f"/api/projects/{pid}").json()
        self.assertEqual(d["cells"][person]["dis"]["s"], "done")  # 전자 서명 → 바로 이수
        scores = {k: 4 for k in __import__("daisu.sessions", fromlist=["SURVEY_ITEMS"]).SURVEY_ITEMS}
        self.assertEqual(c.post(f"/api/projects/{pid}/sessions/{sid}/survey", json={"scores": {"x": 9}}).status_code, 400)
        self.assertEqual(c.post(f"/api/projects/{pid}/sessions/{sid}/survey", json={"scores": scores, "comment": "좋음"}).json()["n"], 1)
        self.assertEqual(c.patch(f"/api/projects/{pid}/sessions/{sid}", json={"instructor": {"grade": "없는등급"}}).status_code, 400)
        self.assertEqual(c.get(f"/api/projects/{pid}/export?doc=session_fee&fmt=hwpx&session={sid}").status_code, 400)  # 강사 없음
        c.patch(f"/api/projects/{pid}/sessions/{sid}", json={"instructor": {"name": "박강사", "org": "가상기관", "grade": "일반2"}, "start": "10:00"})
        v = c.get(f"/api/projects/{pid}/sessions/{sid}").json()
        self.assertEqual(v["fee"]["net"], 136800)
        tasks = {t["key"]: t for t in c.get(f"/api/projects/{pid}/tasks").json()}
        self.assertFalse(tasks[f"plan:{sid}"]["done"])
        for doc, fmt in (("session_plan", "hwpx"), ("signsheet", "docx"), ("signsheet", "hwpx"), ("session_result", "hwpx"), ("session_fee", "docx")):
            r = c.get(f"/api/projects/{pid}/export?doc={doc}&fmt={fmt}&session={sid}")
            self.assertEqual(r.status_code, 200, doc)
        tasks = {t["key"]: t for t in c.get(f"/api/projects/{pid}/tasks").json()}
        self.assertTrue(tasks[f"plan:{sid}"]["done"] and tasks[f"result:{sid}"]["done"] and tasks[f"run:{sid}"]["done"])  # 자동 완료
        self.assertTrue(c.post(f"/api/projects/{pid}/tasks/issues").json()["done"])
        for q in ("doc=nextplan&fmt=hwpx", "doc=stats", f"doc=card&fmt=docx&person={person}"):
            self.assertEqual(c.get(f"/api/projects/{pid}/export?{q}").status_code, 200, q)
        st = c.get(f"/api/projects/{pid}/stats").json()
        self.assertEqual(st["spend"], 150000)
        c.delete(f"/api/projects/{pid}/sessions/{sid}/attend/{person}")
        self.assertNotEqual(c.get(f"/api/projects/{pid}").json()["cells"][person]["dis"]["s"], "done")
        c.delete(f"/api/projects/{pid}/sessions/{sid}")
        self.assertEqual(c.get(f"/api/projects/{pid}/sessions/{sid}").status_code, 404)
        self.assertEqual(c.post(f"/api/projects/{pid}/sessions", json={"manual": True, "col": "int", "date": "2026-11-20"}).status_code, 200)
