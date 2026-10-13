"""제출 양식 렌더링 테스트: [예시] 양식과 기관이 따로 만든 양식(열 순서·이름이 다른 표) 모두 채워지는지."""
from __future__ import annotations

import io
import zipfile
from decimal import Decimal as D

import pytest

from daisu.rulekit.calc import compute_all, parse_judgment
from daisu.rulekit.official import DEFAULT_MAP, TableMap, TemplateMap, build_default_template, render, report_data

from .test_calc import J, KIM, LEE, PARK, RS, it


@pytest.fixture(scope="module")
def results():
    js = [J("P1", it("BASIC-A5-1", "기본교육", 40, "r1"), it("PROF-A6-1", "전문교육", 30, "r2")),
          J("P2", it("BASIC-A5-1", "기본교육", 10, "r1")),
          J("P3", it("BASIC-A5-1", "기본교육", 40, "r1"), it("PROF-A6-2", "전문교육", 50, "r2"), it("SELF-A7-1", "자기개발", 9, "r3"))]
    return compute_all([KIM, LEE, PARK], [parse_judgment(j) for j in js], RS)


def _hwpx_tables(b: bytes):
    import tempfile

    from daisu.readers import read
    with tempfile.NamedTemporaryFile(suffix=".hwpx", delete=False) as f:
        f.write(b)
    doc = read(f.name)
    return doc, [t.rows for t in doc.tables]


def test_report_data(results):
    d = report_data(results, "가상기관", 2026, "2026-12-31")
    assert d["people"] == 3 and d["total_hours"] == D("174.5")
    tot = d["depts_total"]
    assert tot["dept"] == "합계" and tot["met_count"] == 2 and tot["met_rate"] == D("66.7")
    assert [p["shortfall"] for p in d["persons"]] == [D(0), D(50), D(0)]


def test_default_hwpx(results):
    b = render(results, "가상기관", 2026, "2026-12-31", "hwpx")
    z = zipfile.ZipFile(io.BytesIO(b))
    assert z.namelist()[0] == "mimetype" and z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
    doc, tables = _hwpx_tables(b)
    text = "\n".join(t["text"] if isinstance(t, dict) else str(t) for t in doc.texts)
    assert "{{" not in text and "가상기관" in text and "2026년" in text
    dept, person = tables
    assert dept[0][1] == "부서" and len(dept) == 1 + 2 + 1  # 머리글 + 부서 2 + 합계
    assert dept[-1][1] == "합계" and dept[-1][2] == "3"
    assert len(person) == 4 and person[2][2] == "이팀장" and person[2][-2] == "50"
    sec = z.read("Contents/section0.xml").decode()
    assert 'rowCnt="4"' in sec and 'rowCnt="4"' in sec


def test_default_docx(results):
    from docx import Document
    b = render(results, "가상기관", 2026, "2026-12-31", "docx")
    d = Document(io.BytesIO(b))
    assert all("{{" not in p.text for p in d.paragraphs)
    dept, person = d.tables
    assert [c.text for c in dept.rows[-1].cells][1:3] == ["합계", "3"]
    assert [c.text for c in person.rows[3].cells][2] == "박시설"
    assert [c.text for c in person.rows[3].cells][7] == "4.5"  # 자기개발 9시간 × 1/2


def test_custom_docx_template_column_order(results):
    """기관 양식: 열 순서가 다르고 매핑에 없는 열(비고)이 있어도 머리글 글자로 찾아 채운다."""
    from docx import Document
    t = Document()
    t.add_paragraph("{{기관명}} 교육훈련 실적({{연도}})")
    tb = t.add_table(rows=2, cols=5)
    for i, h in enumerate(["성 명", "부서", "합계", "미달시간", "비고"]):
        tb.rows[0].cells[i].text = h
    buf = io.BytesIO()
    t.save(buf)
    mp = TemplateMap(placeholders=DEFAULT_MAP.placeholders,
                     tables=[TableMap(source="persons", columns={"성명": "name", "부서": "dept", "합계": "total", "미달시간": "shortfall"})])
    out = Document(io.BytesIO(render(results, "가상기관", 2026, "2026-12-31", "docx", template=buf.getvalue(), mapping=mp)))
    assert out.paragraphs[0].text == "가상기관 교육훈련 실적(2026)"
    rows = [[c.text for c in r.cells] for r in out.tables[0].rows]
    assert rows[1] == ["김주무", "총무팀", "70", "0", ""]
    assert len(rows) == 4


def test_custom_hwpx_template(results):
    """다른 프로그램으로 만든 hwpx 양식(머리글 행 + 빈 서식 행)에도 채운다."""
    from daisu.render import to_hwpx
    tpl = to_hwpx([("title", "{{기관명}} 부서별 실적"), ("table", ["부서", "인원", "총 이수시간", "충족률(%)"], [["", "", "", ""]])])
    mp = TemplateMap(placeholders={"{{기관명}}": "org"},
                     tables=[TableMap(source="depts", columns={"부서": "dept", "인원": "headcount", "총 이수시간": "total_hours",
                                                               "충족률(%)": "met_rate"}, total_row=True)])
    _, tables = _hwpx_tables(render(results, "가상기관", 2026, "2026-12-31", "hwpx", template=tpl, mapping=mp))
    assert tables[0][1:] == [["총무팀", "2", "80", "50"], ["시설팀", "1", "94.5", "100"], ["합계", "3", "174.5", "66.7"]]


def test_default_template_is_marked_as_example():
    from docx import Document
    d = Document(io.BytesIO(build_default_template("docx")))
    assert any("[예시]" in p.text for p in d.paragraphs)


def test_bad_format(results):
    with pytest.raises(ValueError):
        render(results, "가상기관", 2026, "2026-12-31", "pdf")
