"""자료 읽기: 엑셀·CSV·한글(HWPX/HWP)·워드·PDF·텍스트에서 '표'와 '글'을 꺼낸다.

원본 파일은 읽기만 한다. 표는 행·열 그대로(문자열) 돌려주고, 표 바로 앞의 글(교육명·일시 등)을 함께 붙인다.
"""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
import struct
import zipfile
import zlib
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field

SUPPORTED = (".xlsx", ".xlsm", ".csv", ".hwpx", ".hwp", ".docx", ".pdf", ".txt", ".md")


@dataclass
class Table:
    rows: list[list[str]]
    where: str                       # 예: "시트 '이수현황'", "표 1"
    before: list[str] = field(default_factory=list)  # 표 앞의 글(제목·교육명·일시)


@dataclass
class Doc:
    path: str
    ext: str
    tables: list[Table] = field(default_factory=list)
    texts: list[tuple[str, str]] = field(default_factory=list)  # (문단, 위치)
    note: str = ""


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute) == (0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read(path: str) -> Doc:
    ext = os.path.splitext(path)[1].lower()
    doc = Doc(path, ext)
    try:
        {".xlsx": _xlsx, ".xlsm": _xlsx, ".csv": _csv, ".hwpx": _hwpx, ".hwp": _hwp, ".docx": _docx,
         ".pdf": _pdf, ".txt": _txt, ".md": _txt}[ext](path, doc)
    except KeyError:
        doc.note = "지원하지 않는 형식"
    except Exception as e:  # 손상 파일 하나 때문에 전체가 멈추지 않도록
        doc.note = f"읽기 실패: {type(e).__name__}"
    return doc


def _xlsx(path, doc):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        while rows and not any(rows[-1]):
            rows.pop()
        if rows:
            doc.tables.append(Table(rows, f"시트 '{ws.title}'"))
    wb.close()


def _csv(path, doc):
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, newline="", encoding=enc) as f:
                doc.tables.append(Table([[c.strip() for c in r] for r in csv.reader(f)], "표"))
            return
        except UnicodeDecodeError:
            continue
    doc.note = "인코딩 판별 실패"


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _hwpx(path, doc):
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
                       key=lambda n: int(re.findall(r"\d+", n)[-1]))
        recent: list[str] = []
        n_tbl = 0

        def text_of(el) -> str:
            return "".join(t.text or "" for t in el.iter() if _local(t.tag) == "t").strip()

        def walk(el):
            nonlocal n_tbl
            tag = _local(el.tag)
            if tag == "tbl":
                n_tbl += 1
                rows = []
                for tr in (x for x in el.iter() if _local(x.tag) == "tr"):
                    rows.append([", ".join(t for t in (text_of(p) for p in tc.iter() if _local(p.tag) == "p") if t)
                                 for tc in tr if _local(tc.tag) == "tc"])
                doc.tables.append(Table(rows, f"표 {n_tbl}", recent[-6:]))
                return
            if tag == "p" and not any(_local(x.tag) == "tbl" for x in el.iter()):
                t = text_of(el)
                if t:
                    recent.append(t)
                    doc.texts.append((t, f"문단 {len(doc.texts) + 1}"))
                return
            for child in el:
                walk(child)

        for name in names:
            walk(ET.fromstring(z.read(name)))


def _hwp(path, doc):
    """HWP 5.x: 본문 레코드의 글자만 꺼낸다(표는 칸별 문단으로 나옴)."""
    import olefile

    with olefile.OleFileIO(path) as ole:
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0] if len(header) >= 40 else 0
        if flags & 2:
            doc.note = "암호가 걸린 한글 문서"
            return
        for entry in sorted((e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"),
                            key=lambda e: int(re.sub(r"\D", "", e[1]) or 0)):
            raw = ole.openstream(entry).read()
            data = zlib.decompress(raw, -15) if flags & 1 else raw
            pos = 0
            while pos + 4 <= len(data):
                (h,) = struct.unpack_from("<I", data, pos)
                pos += 4
                tag, size = h & 0x3FF, (h >> 20) & 0xFFF
                if size == 0xFFF:
                    (size,) = struct.unpack_from("<I", data, pos)
                    pos += 4
                if tag == 67:
                    body = data[pos:pos + size]
                    chars = []
                    i = 0
                    while i < len(body) - 1:
                        ch = body[i] | (body[i + 1] << 8)
                        if ch < 32:
                            i += 16 if ch in {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23} else 2
                            continue
                        chars.append(chr(ch))
                        i += 2
                    t = "".join(chars).strip()
                    if t:
                        doc.texts.append((t, f"문단 {len(doc.texts) + 1}"))
                pos += size


def _docx(path, doc):
    from docx import Document

    d = Document(path)
    recent: list[str] = []
    n_tbl = 0
    for el in d.element.body.iterchildren():
        tag = _local(el.tag)
        if tag == "p":
            t = "".join(x.text or "" for x in el.iter() if _local(x.tag) == "t").strip()
            if t:
                recent.append(t)
                doc.texts.append((t, f"문단 {len(doc.texts) + 1}"))
        elif tag == "tbl":
            n_tbl += 1
            rows = []
            for tr in (x for x in el.iter() if _local(x.tag) == "tr"):
                rows.append(["".join(t.text or "" for t in tc.iter() if _local(t.tag) == "t").strip()
                             for tc in tr if _local(tc.tag) == "tc"])
            doc.tables.append(Table(rows, f"표 {n_tbl}", recent[-6:]))


def _pdf(path, doc):
    from pypdf import PdfReader

    r = PdfReader(path)
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            doc.note = "암호가 걸린 PDF"
            return
    for i, page in enumerate(r.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            doc.texts.append((t, f"{i}쪽"))
    if not doc.texts:
        doc.note = "글자가 없는 PDF(스캔본은 OCR 필요)"


def _txt(path, doc):
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc) as f:
                lines = f.read().splitlines()
            break
        except UnicodeDecodeError:
            continue
    else:
        doc.note = "인코딩 판별 실패"
        return
    for i, line in enumerate(lines, 1):
        if line.strip():
            doc.texts.append((line.strip(), f"{i}행"))
