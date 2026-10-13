"""운영 지원: 할 일 목록, 교육훈련 실적 집계(경영평가 대비), 직원별 교육이력 카드, 다음 해 운영계획(안)."""
from __future__ import annotations

import datetime as dt
import io
import math
import re

from . import fees, gongmun
from .reports import improvements, survey_for, survey_summary
from .util import dept_core, fmt_h, kdate


def _d(s):
    return dt.date.fromisoformat(s) if s else None


# ───────────── 할 일
def tasks(project: dict, res: dict, remind: dict) -> list[dict]:
    base = _d(res["base"])
    done = set(project.get("done_tasks", []))
    exported = project.get("exported", {})
    out = []

    def add(key, when, title, detail="", link="", auto_done=False):
        out.append({"key": key, "date": when.isoformat() if when else None, "title": title, "detail": detail, "link": link,
                    "done": key in done or auto_done, "auto": auto_done,
                    "dday": (when - base).days if when else None})

    if res["issues"]:
        add("issues", base, f"확인 필요 {len(res['issues'])}건 처리", "동명이인·오타·미분류 과정을 정하면 이수 현황이 정확해집니다", "issues")
    by_due: dict[str, list] = {}
    for s in res["summary"]:
        if s["mandatory"] and s["due"]:
            by_due.setdefault(s["due"], []).append(s)
    for due, ss in sorted(by_due.items()):
        left = [s for s in ss if s["partial"] + s["none"]]
        dd = _d(due)
        if not left:
            continue
        n = sum(s["partial"] + s["none"] for s in left)
        names = ", ".join(s["label"] for s in left[:4]) + (" 외" if len(left) > 4 else "")
        if dd < base:
            add(f"over:{due}", base, f"기한({kdate(dd, False)}) 지난 교육 보충·사유 확인 – {n}명·건", names, "matrix/over")
            continue
        add(f"remind:{due}", dd - dt.timedelta(days=remind.get("first_days", 30)), f"미이수자 본인 안내 발송 ({n}명·건)",
            f"기한 {kdate(dd, False)}: {names}", "notify")
        add(f"boss:{due}", dd - dt.timedelta(days=remind.get("boss_days", 14)), "부서장 미이수 통보", f"기한 {kdate(dd, False)}", "notify")
        add(f"due:{due}", dd, f"이수 기한 – {names}", f"남은 {n}명·건", "board")
    for s in project.get("sessions", []):
        sd = _d(s["date"])
        sid = s["id"]
        add(f"plan:{sid}", sd - dt.timedelta(days=7), f"실시계획 기안 – {s['title']}", kdate(sd), f"session/{sid}",
            auto_done=f"{sid}:plan" in exported)
        if not s["instructor"].get("name"):
            add(f"inst:{sid}", sd - dt.timedelta(days=14), f"강사 섭외 – {s['title']}", "강사 정보를 넣으면 강사료가 계산됩니다", f"session/{sid}")
        add(f"run:{sid}", sd, f"교육 실시·출석 키오스크 – {s['title']}", f"{s['start']}~{s['end']} {s.get('place') or ''}",
            f"session/{sid}", auto_done=bool(s["attend"]))
        add(f"result:{sid}", sd + dt.timedelta(days=3), f"결과 보고 기안 – {s['title']}", "", f"session/{sid}",
            auto_done=f"{sid}:result" in exported)
        if s["instructor"].get("name") and s["instructor"].get("income") != "근로":
            add(f"fee:{sid}", sd + dt.timedelta(days=5), f"강사료 지급 기안 – {s['title']}", "", f"session/{sid}",
                auto_done=f"{sid}:fee" in exported)
    hires = [p for p in res["people"] if res["cells"][p["id"]].get("hire", {}).get("s") in ("none", "partial")]
    for p in hires:
        add(f"hire:{p['id']}", _d(res["cells"][p["id"]]["hire"]["due"]), f"신규입사자 채용 시 교육 – {p['name']}({p['dept']})",
            "업무 배치 전 8시간", "matrix/hire")
    if any(s["group"] == "폭력예방" for s in res["summary"]):
        add(f"report:{res['year']}", dt.date(res["year"] + 1, 2, 28), f"{res['year']}년 폭력예방교육 실적 등록(주무부처 시스템)",
            "보고·증빙 탭의 이수현황 엑셀 활용", "report")
    out.sort(key=lambda t: (t["done"], t["date"] or "9999"))
    return out


# ───────────── 교육훈련 실적 집계
def stats(project: dict, res: dict, fee_cfg: dict) -> dict:
    y0, y1 = f"{res['year']}-01-01", f"{res['year']}-12-31"
    active = [p for p in res["people"] if not (p.get("left") and p["left"] <= res["base"])]
    hours = {p["id"]: {"legal": 0.0, "other": 0.0} for p in active}
    seen = set()
    for r in res["records"]:
        pid = r.get("pid")
        if not pid or pid not in hours or not r["done"] or (r.get("date") and not y0 <= r["date"] <= y1):
            continue
        key = (pid, r["title"], r.get("date"))
        if key in seen:  # 같은 과정·같은 날 중복 기록은 한 번만
            continue
        seen.add(key)
        hours[pid]["legal" if r.get("rules") else "other"] += float(r["hours"] or 0)
    total = sum(h["legal"] + h["other"] for h in hours.values())
    n = max(1, len(active))

    def rate(ps):
        cells = [c for p in ps for col in res["columns"] if col["mandatory"] for c in [res["cells"][p["id"]][col["key"]]]
                 if c["s"] in ("done", "partial", "none")]
        return round(100 * sum(1 for c in cells if c["s"] == "done") / max(1, len(cells)), 1)

    groups = []
    for label, ps in [("전체", active)] + [(d, [p for p in active if p["dept"] == d]) for d in res["depts"]] + \
                     [(f"직무: {j}", [p for p in active if p["job"] == j]) for j in ("사무", "현장")] + \
                     [(f"근무: {s}", [p for p in active if p["shift"] == s]) for s in ("일근", "교대")]:
        if not ps:
            continue
        hs = [hours[p["id"]]["legal"] + hours[p["id"]]["other"] for p in ps]
        groups.append({"group": label, "n": len(ps), "avg_hours": round(sum(hs) / len(ps), 1), "rate": rate(ps),
                       "legal": round(sum(hours[p["id"]]["legal"] for p in ps) / len(ps), 1),
                       "other": round(sum(hours[p["id"]]["other"] for p in ps) / len(ps), 1)})
    spend = sum(fees.compute(s["instructor"], s["hours"], fee_cfg)["total"] for s in project.get("sessions", [])
                if s["instructor"].get("name"))
    return {"year": res["year"], "people": len(active), "avg_hours": round(total / n, 1), "total_hours": round(total, 1),
            "groups": groups, "spend": spend, "sessions": len(project.get("sessions", [])), "hours": hours}


def stats_excel(st: dict, res: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "집계"
    ws.append([f"{st['year']}년 교육훈련 실적 집계 (기준일 {res['base']})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["재직 인원", st["people"], "1인당 교육시간", st["avg_hours"], "총 교육시간", st["total_hours"], "집합교육 강사료", st["spend"]])
    ws.append([])
    ws.append(["구분", "인원", "1인당 교육시간", "법정교육(1인당)", "그 밖의 교육(1인당)", "법정교육 이수율(%)"])
    for g in st["groups"]:
        ws.append([g["group"], g["n"], g["avg_hours"], g["legal"], g["other"], g["rate"]])
    ws2 = wb.create_sheet("직원별")
    ws2.append(["사번", "성명", "소속", "직무", "근무", "법정교육 시간", "그 밖의 교육 시간", "합계"])
    for p in res["people"]:
        h = st["hours"].get(p["id"])
        if h:
            ws2.append([p["id"], p["name"], p["dept"], p["job"], p["shift"], h["legal"], h["other"], h["legal"] + h["other"]])
    ws.column_dimensions["A"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────── 직원별 교육이력 카드
def card_blocks(res: dict, pid: str, org: dict) -> list:
    p = next(x for x in res["people"] if x["id"] == pid)
    names = {r["id"]: r["name"] for r in res["rules"]["courses"]}
    recs = sorted([r for r in res["records"] if r.get("pid") == pid], key=lambda r: r.get("date") or "")
    B = [("title", f"교육훈련 이력 카드 – {p['name']}"),
         ("table", ["사번", "소속", "직위", "고용형태", "근무", "입사일"],
          [[p["id"], p["dept"], p["title"], p["emp"], p["shift"] + (f"({p['team']}조)" if p.get("team") else ""), p["hired"]]], [1, 2, 1, 1, 1, 1]),
         ("h1", f"{res['year']}년 법정의무교육")]
    rows = []
    for c in res["columns"]:
        x = res["cells"][pid][c["key"]]
        if x["s"] == "na":
            continue
        st = {"done": "이수", "partial": "부족", "none": "미이수", "leave": "휴직", "left": "퇴사"}[x["s"]]
        rows.append([c["label"], st + (" (기한 지남)" if x.get("overdue") else ""), f"{fmt_h(x['h'])}/{fmt_h(x['need'])}시간",
                     kdate(_d(x["due"]), False)])
    B.append(("table", ["교육", "상태", "인정/필요", "기한"], rows, [4, 2, 2, 2]))
    B.append(("h1", "교육 이수 기록"))
    B.append(("table", ["날짜", "과정", "시간", "인정 교육", "근거"],
              [[r.get("date") or "", r["title"], fmt_h(r["hours"]) if r["done"] else "미이수",
                ", ".join(names.get(x, x) for x in r.get("rules", [])) or "그 밖의 교육", f"{r['file']} {r['where']}"] for r in recs],
              [1.4, 4, 1, 2.5, 4]))
    B.append(("note", f"{org.get('name', '')} {org.get('dept', '')} · 다이수 자동 작성(기준일 {res['base']})"))
    return B


# ───────────── 다음 해 운영계획(안)
MONTH_HINT = {"폭력예방": 3, "정보보호": 5, "인식개선": 6, "청렴": 9}


def next_plan_blocks(project: dict, res: dict, org: dict, fee_cfg: dict, cap_per_session: int = 18) -> list:
    y = res["year"] + 1
    rules = {r["id"]: r for r in res["rules"]["courses"]}
    active = [p for p in res["people"] if not (p.get("left") and p["left"] <= res["base"])]
    field = [p for p in active if p["job"] == "현장"]
    office = [p for p in active if p["job"] == "사무"]

    def rate_of(ps, key):
        cs = [res["cells"][p["id"]][key] for p in ps if res["cells"][p["id"]][key]["s"] in ("done", "partial", "none")]
        return round(100 * sum(1 for c in cs if c["s"] == "done") / len(cs), 1) if cs else None

    # 1) 올해 결과 표 + 방식(사이버/집합) 비중
    rows, findings = [], []
    rec = {r["rid"]: r for r in res["records"]}
    for s in res["summary"]:
        if not s["mandatory"]:
            continue
        kinds = {"cyber": 0, "group": 0}
        for p in active:
            for rid in res["cells"][p["id"]][s["key"]]["recs"]:
                kinds["cyber" if rec[rid]["kind"] == "cyber" else "group"] += 1
        tot = max(1, sum(kinds.values()))
        fr, of = rate_of(field, s["key"]), rate_of(office, s["key"])
        rows.append([s["label"], f"{s['rate']}%" if s["rate"] is not None else "-", f"{of if of is not None else '-'} / {fr if fr is not None else '-'}",
                     f"사이버 {round(100 * kinds['cyber'] / tot)}% · 집합·외부 {round(100 * kinds['group'] / tot)}%", f"{s['overdue']}명" if s["overdue"] else "-"])
        ongoing = s["cycle"] == "half" and s["due"] and s["due"] >= res["base"]  # 진행 중인 반기는 비교하지 않음
        if fr is not None and of is not None and of - fr >= 10 and not ongoing:
            findings.append(f"{s['label']}: 현장직 이수율({fr}%)이 사무직({of}%)보다 {round(of - fr, 1)}%p 낮음")
        if s["overdue"]:
            findings.append(f"{s['label']}: 기한 지난 미이수 {s['overdue']}명")
    sv_imp = []
    for sg in res["signs"]:
        sv = survey_summary(survey_for(res, sg["title"] + " " + sg.get("context", "")))
        for x in improvements(sv):
            sv_imp.append(x)
    for s in project.get("sessions", []):
        sv = survey_summary([{"course": s["title"], "dept": "", "scores": r["scores"], "comment": r.get("comment", "")} for r in s["survey"]])
        for x in improvements(sv):
            sv_imp.append(x)

    # 2) 운영 방향(근거가 있는 것만)
    direction = []
    sv_imp = list(dict.fromkeys(re.sub(r" \(관련 의견 \d+건\)$", "", x) for x in sv_imp if "점수가 가장 낮음" not in x))
    if any("현장직 이수율" in f for f in findings) or (res["shifts"] and any(p["shift"] == "교대" for p in active)):
        direction.append("현장직: 교대근무표 기반 집합교육(주간 근무일 배정·팀별 현장 인원 유지) + 사이버교육 병행")
    if any("상반기" in f and "기한" in f for f in findings):
        direction.append("상반기 산업안전보건교육: 4~5월 조기 실시, 6월은 보충 회차만 운영")
    if any("4시간 이내" in x for x in sv_imp):
        direction.append("장시간 집합교육은 4시간 이내 회차로 분할(만족도 의견 반영)")
    if any("실습" in x for x in sv_imp):
        direction.append("팀 업무와 연결된 현장 실습·사례 중심 교안")
    direction.append("출석은 전자 서명(키오스크)으로 받아 이수 현황에 즉시 반영, 결과보고·강사료 기안 자동 작성")
    direction.append(f"미이수자 관리: 기한 30일 전 본인 안내, 14일 전 부서장 통보, 기한 경과 시 보고")

    # 3) 월별 일정(안)
    months: dict[int, list] = {m: [] for m in range(1, 13)}
    for r in rules.values():
        if not r.get("mandatory", True):
            continue
        if r.get("cycle") == "half":
            months[4].append(f"{r['name']}(상반기) – 사무 사이버, 현장 집합")
            months[5].append(f"{r['name']}(상반기) 집합 회차")
            months[6].append(f"{r['name']}(상반기) 보충 회차·기한 6/30")
            months[10].append(f"{r['name']}(하반기) 집합 회차")
            months[11].append(f"{r['name']}(하반기) 집합 회차")
        elif r.get("cycle") == "hire":
            continue
        else:
            months[MONTH_HINT.get(r.get("group"), 7)].append(r["name"])
    for g in ("폭력예방",):
        names = [r["name"] for r in rules.values() if r.get("group") == g]
        if names:
            months[3] = [x for x in months[3] if x not in names] + [f"4대폭력 예방 통합교육(사이버 4시간: {', '.join(n.replace(' 예방교육', '') for n in names)})"]
    months[12].append("미이수자 보충 교육, 연간 실적 정리")
    month_rows = [[f"{m}월", "\n".join(v)] for m, v in months.items() if v]

    # 4) 소요 예산(추정): 올해 회차 운영 방식으로 현장 집합교육 회차 수 산정
    budget_rows, total = [], 0
    tier = (fee_cfg.get("tiers") or [{}])[min(2, len(fee_cfg.get("tiers") or [1]) - 1)]
    for r in rules.values():
        if r.get("cycle") != "half" or not r.get("mandatory", True):
            continue
        need = r["hours"].get("현장", 12) if isinstance(r["hours"], dict) else r["hours"]
        tgt = [p for p in field if not p["flags"].get((r.get("target") or {}).get("exclude_flag", ""), False)]
        per = math.ceil(need / 4)
        n_sess = math.ceil(len(tgt) / cap_per_session) * per * 2
        fee = fees.compute({"grade": tier.get("grade"), "income": "기타소득"}, 4, fee_cfg)["total"]
        budget_rows.append([f"{r['name']}(현장 집합)", f"{len(tgt)}명 × {per}회(4시간) × 반기 2회 ÷ 회차당 {cap_per_session}명",
                            f"{n_sess}회", f"{n_sess * fee:,}원"])
        total += n_sess * fee
    budget_rows.append(["합계", "외부 강사 기준 추정(내부 강사 활용 시 감액)", "", f"{total:,}원"])

    body = ["관련: 산업안전보건법, 양성평등기본법 등 법정의무교육 관계 법령(교육별 근거는 교육 규칙표) 및 기관 운영지침",
            f"{res['year']}년 법정의무교육 운영 결과를 분석하여 {y}년 운영 계획(안)을 아래와 같이 수립하고자 합니다.",
            [f"{res['year']}년 운영 결과: 재직 {len(active)}명, 법정교육 평균 이수율 "
             f"{round(sum(s['rate'] or 0 for s in res['summary'] if s['mandatory']) / max(1, sum(1 for s in res['summary'] if s['mandatory'])), 1)}%(기준일 {res['base']})",
             "주요 문제: " + ("; ".join(findings[:5]) if findings else "특이사항 없음"),
             "교육생 의견: " + ("; ".join(sv_imp[:4]) if sv_imp else "특이사항 없음")],
            f"{y}년 운영 방향", direction,
            f"세부 일정(안)과 소요 예산(안): 붙임 참조(예산은 기관 강사료 기준 {tier.get('grade', '')} 등급 4시간 기준 추정)"]
    tables = [("h2", f"붙임 1. {res['year']}년 교육별 운영 결과"),
              ("table", ["교육", "이수율", "사무 / 현장 이수율(%)", "이수 방식", "기한 지남"], rows, [4, 1.2, 2, 2.6, 1.2]),
              ("h2", f"붙임 2. {y}년 월별 일정(안)"), ("table", ["시기", "교육"], month_rows, [1, 7]),
              ("h2", f"붙임 3. {y}년 소요 예산(안)"), ("table", ["항목", "산출 근거", "회차", "금액"], budget_rows, [3, 5, 1, 2])]
    return gongmun.blocks(org, f"{y}년 법정의무교육 운영 계획(안)", body,
                          [f"{res['year']}년 교육별 운영 결과 1부.", f"{y}년 월별 일정(안) 1부.", f"{y}년 소요 예산(안) 1부."], tables=tables)
