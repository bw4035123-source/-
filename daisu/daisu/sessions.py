"""집합교육 회차 운영: 일정 확정 → 실시계획 기안 → 출석(키오스크 서명) → 퇴실 만족도 → 결과보고 → 강사료 지급.

출석은 곧바로 이수 기록이 되어 이수 현황에 반영된다(서명부를 다시 입력할 필요 없음).
"""
from __future__ import annotations

import base64
import datetime as dt
import io
import re
import uuid

from . import fees, gongmun
from .reports import improvements, survey_summary
from .util import dept_core, fmt_h, josa, kdate

SURVEY_ITEMS = ["교육 내용 만족(5점)", "강사 만족(5점)", "업무 도움(5점)", "시간·장소 적절(5점)"]
MAX_SIGN = 300_000  # 서명 이미지 최대 크기(바이트, base64 문자열 기준)


def new_id() -> str:
    return "s" + uuid.uuid4().hex[:7]


def from_plan(pl: dict, rule_law: str, place: str = "") -> list[dict]:
    out = []
    n = len(pl["sessions"])
    for k, s in enumerate(pl["sessions"], 1):
        m = re.search(r"(\d{1,2}:\d{2})~(\d{1,2}:\d{2})", s["label"])
        out.append({"id": new_id(), "col": pl["col"]["key"], "rule": pl["col"]["rule"], "law": rule_law,
                    "title": f"{pl['col']['label']} {k}회차" if n > 1 else pl["col"]["label"],
                    "date": s["date"], "start": m.group(1) if m else "09:00", "end": m.group(2) if m else "",
                    "hours": pl["hours"], "place": place, "targets": [p["id"] for p in s["people"]],
                    "instructor": {"name": "", "org": "", "grade": "", "income": "기타소득", "pages": 0, "travel": 0},
                    "attend": {}, "survey": [], "created": dt.datetime.now().strftime("%Y-%m-%d %H:%M")})
    return out


def records(project: dict, people_by_id: dict) -> list[dict]:
    """출석(서명)한 사람 → 이수 기록."""
    out = []
    for s in project.get("sessions", []):
        for pid, a in s.get("attend", {}).items():
            p = people_by_id.get(pid)
            if not p:
                continue
            out.append({"kind": "kiosk", "title": s["title"], "context": "", "date": s["date"], "hours": float(s["hours"]),
                        "done": True, "name": p["name"], "emp_id": pid, "dept": p["dept"], "file": "전자출석",
                        "where": f"{s['title']} ({s['date']}) {a.get('at', '')[11:16]} 서명", "note": "",
                        "preset_rules": [s["rule"]] if s.get("rule") else []})
    return out


def check_sign(data: str) -> str:
    if not isinstance(data, str) or not data.startswith("data:image/png;base64,") or len(data) > MAX_SIGN:
        raise ValueError("서명 이미지 형식이 올바르지 않습니다.")
    raw = base64.b64decode(data.split(",", 1)[1], validate=True)
    # PNG 머리(시그니처+IHDR)와 끝(IEND)이 모두 있어야 서명부 문서에 넣을 수 있다
    if not (raw.startswith(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR") and raw.rstrip(b"\x00").endswith(b"IEND\xaeB`\x82")):
        raise ValueError("서명 이미지 형식이 올바르지 않습니다.")
    return data


def view(s: dict, people_by_id: dict, fee_cfg: dict) -> dict:
    """화면용: 서명 이미지는 빼고 출석 여부·시각만."""
    tg = []
    for pid in s["targets"]:
        p = people_by_id.get(pid, {"name": "?", "dept": "", "team": ""})
        a = s["attend"].get(pid)
        tg.append({"id": pid, "name": p["name"], "dept": p["dept"], "team": p.get("team", ""), "at": a["at"] if a else None})
    extra = [{"id": pid, "name": people_by_id.get(pid, {}).get("name", "?"), "dept": people_by_id.get(pid, {}).get("dept", ""),
              "team": people_by_id.get(pid, {}).get("team", ""), "at": a["at"], "walkin": True}
             for pid, a in s["attend"].items() if pid not in s["targets"]]
    return {k: s[k] for k in ("id", "col", "rule", "title", "date", "start", "end", "hours", "place", "instructor", "law")} | {
        "targets": tg + extra, "attended": len(s["attend"]), "n_targets": len(s["targets"]), "survey_n": len(s["survey"]),
        "fee": fees.compute(s["instructor"], s["hours"], fee_cfg) if s["instructor"].get("name") else None,
        "survey": survey_summary(_survey_rows(s))}


def _survey_rows(s):
    return [{"course": s["title"], "dept": "", "scores": r["scores"], "comment": r.get("comment", "")} for r in s.get("survey", [])]


def _time(s):
    return f"{kdate(dt.date.fromisoformat(s['date']))} {s['start']}~{s['end']} ({fmt_h(s['hours'])}시간)"


def _by_dept(pids, people_by_id):
    cnt: dict[str, int] = {}
    for pid in pids:
        d = people_by_id.get(pid, {}).get("dept", "")
        cnt[d] = cnt.get(d, 0) + 1
    return ", ".join(f"{dept_core(d)} {n}명" for d, n in cnt.items())


def fee_table(fee: dict) -> tuple:
    rows = [["강사 등급", f"{fee['grade']} {fee['grade_desc']}".strip()], ["강의 시간", f"{fmt_h(fee['hours'])}시간"],
            ["강의료", f"{fee['first'] + fee['extra']:,}원"], ["원고료", f"{fee['manuscript']:,}원"], ["교통비(실비)", f"{fee['travel']:,}원"],
            ["지급 총액(세전)", f"{fee['total']:,}원"], ["소득 구분", fee["income"]],
            ["소득세", f"{fee['income_tax']:,}원"], ["지방소득세", f"{fee['local_tax']:,}원"], ["실지급액", f"{fee['net']:,}원"]]
    return ("table", ["항목", "금액·내용"], rows, [2, 5])


def plan_doc(s: dict, people_by_id: dict, org: dict, fee_cfg: dict) -> list:
    inst = s["instructor"]
    fee = fees.compute(inst, s["hours"], fee_cfg) if inst.get("name") else None
    body = [f"관련: {s.get('law') or '관계 법령'}, 기관 법정의무교육 운영지침",
            f"아래와 같이 {josa('「' + s['title'] + '」')} 실시하고자 합니다.",
            [f"교육명: {s['title']}", f"일시: {_time(s)}", f"장소: {s.get('place') or '(미정)'}",
             f"대상: {len(s['targets'])}명({_by_dept(s['targets'], people_by_id)})",
             f"강사: {inst['name']}({inst.get('org', '')})" if inst.get("name") else "강사: (섭외 중)",
             "방법: 집합교육(교육 시간은 근무로 인정), 출석은 전자 서명으로 확인"],
            f"소요 예산: 강사료 {fee['total']:,}원(원천징수 {fee['withholding']:,}원 포함, 산출 내역 붙임)" if fee else "소요 예산: 강사 확정 후 산출",
            "협조 사항: 각 팀장은 대상자가 교육에 참석할 수 있도록 근무를 조정하여 주시기 바랍니다"
            "(대상자는 교대근무표상 주간 근무일에 배정하였으며, 팀별 현장 근무 인원을 유지함)."]
    names = [["연번", "소속", "성명", "근무조"]]
    rows = [[str(i), people_by_id.get(pid, {}).get("dept", ""), people_by_id.get(pid, {}).get("name", ""),
             people_by_id.get(pid, {}).get("team", "") or "일근"] for i, pid in enumerate(s["targets"], 1)]
    tables = [("h2", "붙임 1. 교육 대상자 명단"), ("table", names[0], rows, [1, 3, 2, 1])]
    att = ["교육 대상자 명단 1부."]
    if fee:
        tables += [("h2", "붙임 2. 강사료 산출 내역"), fee_table(fee)] + [("bullet", x) for x in fee["steps"]]
        att.append("강사료 산출 내역 1부.")
    return gongmun.blocks(org, f"{s['title']} 실시 계획", body, att, tables=tables)


def result_doc(s: dict, people_by_id: dict, org: dict) -> list:
    att = s["attend"]
    absent = [pid for pid in s["targets"] if pid not in att]
    sv = survey_summary(_survey_rows(s))
    body = [f"관련: {org.get('dept', '')} 「{s['title']} 실시 계획」",
            f"{josa('「' + s['title'] + '」')} 아래와 같이 실시하였기에 그 결과를 보고합니다.",
            [f"일시·장소: {_time(s)}, {s.get('place') or '(미정)'}",
             f"참석: 대상 {len(s['targets'])}명 중 {sum(1 for p in s['targets'] if p in att)}명"
             f"(참석률 {round(100 * sum(1 for p in s['targets'] if p in att) / max(1, len(s['targets'])), 1)}%)"
             + (f", 대상 외 참석 {len(att) - sum(1 for p in s['targets'] if p in att)}명" if len(att) > sum(1 for p in s['targets'] if p in att) else ""),
             f"이수 처리: 출석자 {len(att)}명에게 {fmt_h(s['hours'])}시간 인정(전자 서명 기준, 이수 현황 반영 완료)"],
            "만족도 조사 결과",
            [f"응답 {sv['n']}명, 전체 평균 {sv['overall']}점(5점 만점), 만족(4점 이상) {sv['satisfied']}%",
             "문항별: " + ", ".join(f"{re.sub(r'[(（].*', '', k)} {v}점" for k, v in sv["avg"].items())]
            + ([f"주요 의견: " + "; ".join(f"{t['theme']} {t['count']}건(“{t['examples'][0]}”)" for t in sv["themes"][:3])] if sv["themes"] else [])
            if sv else ["응답 없음"],
            "개선 사항", (improvements(sv) or ["특이사항 없음"]),
            "불참자 조치: " + (f"{len(absent)}명({', '.join(people_by_id.get(p, {}).get('name', '') for p in absent[:15])}"
                             f"{' 외' if len(absent) > 15 else ''})은 다음 회차 또는 사이버교육으로 이수하도록 개별 안내" if absent else "불참자 없음")]
    return gongmun.blocks(org, f"{s['title']} 결과 보고", body, ["전자 서명부 1부.", "만족도 조사 결과 1부."] if sv else ["전자 서명부 1부."])


def fee_doc(s: dict, org: dict, fee_cfg: dict) -> list:
    inst = s["instructor"]
    fee = fees.compute(inst, s["hours"], fee_cfg)
    body = [f"관련: {org.get('dept', '')} 「{s['title']} 실시 계획」 및 「결과 보고」",
            "위 교육의 강사료를 아래와 같이 지급하고자 합니다.",
            [f"강사: {inst['name']}({inst.get('org', '')})", f"강의: {_time(s)}",
             f"지급 총액: {fee['total']:,}원(강의료 {fee['first'] + fee['extra']:,}원, 원고료 {fee['manuscript']:,}원, 교통비 {fee['travel']:,}원)",
             f"원천징수: {fee['withholding']:,}원(소득세 {fee['income_tax']:,}원, 지방소득세 {fee['local_tax']:,}원, {fee['income']})",
             f"실지급액: {fee['net']:,}원", f"지급 근거: 기관 강사료 지급 기준({fee['grade']} 등급)"]]
    tables = [("h2", "붙임 1. 강사료 산출 내역"), fee_table(fee)] + [("bullet", x) for x in fee["steps"]]
    return gongmun.blocks(org, f"{s['title']} 강사료 지급", body, ["강사료 산출 내역 1부.", "강사 신분증·통장 사본 각 1부(별도 제출)."],
                          tables=tables)


def signsheet_docx(s: dict, people_by_id: dict, org: dict) -> bytes:
    """전자 서명부(워드): 서명 이미지를 그대로 넣는다."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(10)
    d.add_heading(f"{s['title']} 참석자 서명부", 1)
    d.add_paragraph(f"일시: {_time(s)} / 장소: {s.get('place') or '-'} / {org.get('name', '')} {org.get('dept', '')}")
    order = list(s["targets"]) + [p for p in s["attend"] if p not in s["targets"]]
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for c, name in enumerate(["연번", "소속", "성명", "서명", "서명 시각"]):
        t.rows[0].cells[c].text = name
    for i, pid in enumerate(order, 1):
        p = people_by_id.get(pid, {})
        a = s["attend"].get(pid)
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = str(i), p.get("dept", ""), p.get("name", "")
        if a and a.get("sign"):
            raw = base64.b64decode(a["sign"].split(",", 1)[1])
            cells[3].paragraphs[0].add_run().add_picture(io.BytesIO(raw), width=Cm(2.6))
            cells[4].text = a["at"][11:16]
        else:
            cells[3].text = "불참" if not a else "(수기 확인)"
            cells[4].text = a["at"][11:16] if a else ""
    d.add_paragraph(f"참석 {len(s['attend'])}명 / 대상 {len(s['targets'])}명 · 다이수 전자 출석으로 작성")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def signsheet_blocks(s: dict, people_by_id: dict, org: dict) -> list:
    order = list(s["targets"]) + [p for p in s["attend"] if p not in s["targets"]]
    rows = []
    for i, pid in enumerate(order, 1):
        p = people_by_id.get(pid, {})
        a = s["attend"].get(pid)
        rows.append([str(i), p.get("dept", ""), p.get("name", ""), ("전자서명" if a.get("sign") else "수기 확인") if a else "불참",
                     a["at"][11:16] if a else ""])
    return [("title", f"{s['title']} 참석자 서명부"), ("p", f"일시: {_time(s)} / 장소: {s.get('place') or '-'}"),
            ("table", ["연번", "소속", "성명", "서명", "시각"], rows, [1, 3, 2, 2, 1]),
            ("note", f"참석 {len(s['attend'])}명 / 대상 {len(s['targets'])}명 · 서명 이미지는 워드 서명부에 들어 있습니다.")]
