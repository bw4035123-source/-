"""보고·안내 산출물: 이수 현황 보고서, 교육 결과보고서, 미이수자 안내문, 부서장 통보, 이수현황 엑셀, 증빙 묶음."""
from __future__ import annotations

import datetime as dt
import io
import os
import re
import zipfile

from .util import dept_core, due_label, fmt_h, kdate

STATUS_KO = {"done": "이수", "partial": "부족", "none": "미이수", "leave": "휴직", "na": "-", "left": "퇴사"}


def _d(s):
    return dt.date.fromisoformat(s) if s else None


def open_items(res: dict, pid: str, mandatory_only=True) -> list[dict]:
    """한 직원의 아직 못 마친 교육 목록."""
    out = []
    for c in res["columns"]:
        if mandatory_only and not c["mandatory"]:
            continue
        cell = res["cells"][pid][c["key"]]
        if cell["s"] in ("none", "partial"):
            out.append({"key": c["key"], "label": c["label"], "remain": round(cell["need"] - cell["h"], 2), "need": cell["need"],
                        "due": cell["due"], "dday": cell.get("dday"), "overdue": cell.get("overdue")})
    return out


# ───────────── 미이수자 안내문·부서장 통보
def reminders(res: dict, org: dict, col_key: str | None = None, only_due_within: int | None = None) -> list[dict]:
    base = _d(res["base"])
    msgs = []
    for p in res["people"]:
        items = [i for i in open_items(res, p["id"]) if not col_key or i["key"] == col_key]
        if only_due_within is not None:
            items = [i for i in items if i["overdue"] or (i["dday"] is not None and i["dday"] <= only_due_within)]
        if not items:
            continue
        items.sort(key=lambda i: i["due"])
        lines = []
        for i in items:
            due = _d(i["due"])
            tail = "기한이 지났습니다. 빠른 이수 바랍니다" if i["overdue"] else f"기한 {kdate(due, False)}(D-{(due - base).days})"
            part = f"{fmt_h(i['remain'])}시간 남음" if i["remain"] < i["need"] else f"{fmt_h(i['need'])}시간"
            lines.append(f"· {i['label']} {part} – {tail}")
        text = (f"[{org.get('name', '')} {org.get('dept', '')}] {p['name']}님, 아직 이수하지 않은 법정의무교육을 알려드립니다.\n"
                + "\n".join(lines)
                + "\n사이버교육으로 이수하시거나 안내되는 집합교육에 참석해 주세요. 문의: " + f"{org.get('dept', '')} {org.get('manager', '교육담당')}")
        msgs.append({"id": p["id"], "name": p["name"], "dept": p["dept"], "team": p.get("team", ""), "items": items, "text": text,
                     "overdue": any(i["overdue"] for i in items)})
    return msgs


def boss_notes(res: dict, org: dict, col_key: str | None = None) -> list[dict]:
    msgs = reminders(res, org, col_key)
    out = []
    for dept in res["depts"]:
        mine = [m for m in msgs if m["dept"] == dept]
        if not mine:
            continue
        rows = "\n".join(f"  - {m['name']}: " + ", ".join(i["label"] + ("(기한 지남)" if i["overdue"] else "") for i in m["items"])
                         for m in mine)
        n_over = sum(1 for m in mine if m["overdue"])
        out.append({"dept": dept, "count": len(mine), "overdue": n_over,
                    "text": f"{dept}장님, 팀원의 법정의무교육 미이수 현황을 알려드립니다(기준일 {kdate(_d(res['base']))}).\n"
                            f"미이수 {len(mine)}명" + (f", 이 중 기한 지남 {n_over}명" if n_over else "") + "\n" + rows
                            + f"\n팀원들이 기한 내 이수할 수 있도록 협조 부탁드립니다. – {org.get('dept', '')}"})
    return out


# ───────────── 만족도 조사 요약
THEMES = [
    ("교육 시간", ["길", "시간이", "나눠", "짧", "집중"]),
    ("근무 일정", ["야간", "비번", "근무 후", "교대"]),
    ("실습·현장 연계", ["실습", "장비", "현장", "우리 팀", "업무와"]),
    ("사례·영상", ["사례", "영상"]),
    ("강사", ["강사"]),
    ("교육 자료", ["자료"]),
    ("교육장 환경", ["덥", "춥", "교육장", "장소"]),
]


def _tokens(s):
    return {t for t in re.findall(r"[가-힣A-Za-z0-9]+", s or "") if len(t) >= 2 and t not in ("2026년", "2026", "교육")}


def survey_for(res: dict, title: str) -> list[dict]:
    want = _tokens(title)
    rounds = set(re.findall(r"\d+회차", title))
    out = []
    for r in res["survey"]:
        got = _tokens(r["course"])
        if rounds and set(re.findall(r"\d+회차", r["course"])) not in (set(), rounds):
            continue
        if len(want & got) >= 2:
            out.append(r)
    return out


def survey_summary(rows: list[dict]) -> dict | None:
    if not rows:
        return None
    items: dict[str, list] = {}
    for r in rows:
        for k, v in r["scores"].items():
            items.setdefault(k, []).append(v)
    avg = {k: round(sum(v) / len(v), 2) for k, v in items.items()}
    allv = [v for vs in items.values() for v in vs]
    comments = [r["comment"] for r in rows if r["comment"]]
    themes = []
    for name, kws in THEMES:
        hits = [c for c in comments if any(k in c for k in kws)]
        if hits:
            themes.append({"theme": name, "count": len(hits), "examples": list(dict.fromkeys(hits))[:2]})
    themes.sort(key=lambda t: -t["count"])
    low = min(avg, key=avg.get)
    return {"n": len(rows), "avg": avg, "overall": round(sum(allv) / len(allv), 2), "all_comments": comments,
            "satisfied": round(100 * sum(1 for v in allv if v >= 4) / len(allv), 1), "themes": themes,
            "comments": len(comments), "lowest": {"item": low, "avg": avg[low]}}


# 개선 사항은 '불만·요청' 표현이 있을 때만 만든다(“짧고 좋았다” 같은 칭찬은 제외)
COMPLAINTS = [
    (["길어", "길다", "나눠", "집중이 어려", "너무 길"], "장시간 교육은 4시간 이내 회차로 나누어 운영"),
    (["야간 근무 후", "비번", "야간 후"], "야간 근무 다음 날(비번) 배정 금지 – 다이수 일정 추천에 이미 반영"),
    (["실습하면", "실습이 필요", "장비로"], "팀 업무와 연결된 현장 실습·장비 시연 포함"),
    (["덥", "춥", "좁"], "교육장 냉난방·좌석 점검"),
    (["미리 나눠", "자료를 미리", "자료가 없"], "교육 자료 사전 배포"),
    (["사례가 더", "사례를 더"], "현장 사례 보강"),
]


def improvements(sv: dict | None, comments: list[str] | None = None) -> list[str]:
    if not sv:
        return []
    out = []
    if sv["lowest"]["avg"] < 3.5:
        out.append(f"‘{re.sub(r'[(（].*', '', sv['lowest']['item']).strip()}’ 점수가 가장 낮음({sv['lowest']['avg']}점) – 우선 개선")
    texts = comments if comments is not None else sv.get("all_comments", [])
    for kws, action in COMPLAINTS:
        n = sum(1 for c in texts if any(k in c for k in kws))
        if n:
            out.append(f"{action} (관련 의견 {n}건)")
    return out


# ───────────── 보고서 블록
def status_blocks(res: dict, org: dict, plan_result: dict | None = None) -> list:
    base = _d(res["base"])
    active = [p for p in res["people"] if not (p.get("left") and p["left"] <= res["base"])]
    on_leave = sum(1 for p in active if p.get("leave") and p["leave"][0] <= res["base"] <= p["leave"][1])
    B = [("title", f"{res['year']}년 법정의무교육 이수 현황 보고"),
         ("p", f"기준일 {kdate(base)} · 재직 {len(active)}명(휴직 {on_leave}명 포함) · {org.get('name', '')} {org.get('dept', '')}"),
         ("h1", "1. 교육별 이수 현황")]
    rows = []
    for s in res["summary"]:
        if not s["mandatory"]:
            continue
        due = _d(s["due"]) if s["due"] else None
        rows.append([s["label"], str(s["target"]), str(s["done"]), f"{s['rate']}%" if s["rate"] is not None else "-",
                     str(s["partial"] + s["none"]) + (f"(부족 {s['partial']})" if s["partial"] else ""),
                     kdate(due, False) if due else due_label(s), f"기한 지남 {s['overdue']}명" if s["overdue"] else ""])
    B.append(("table", ["교육", "대상", "이수", "이수율", "미이수", "기한", "비고"], rows, [5, 1, 1, 1.4, 1.8, 1.8, 2]))
    opt = [s for s in res["summary"] if not s["mandatory"]]
    if opt:
        B.append(("note", "※ 권장 교육: " + ", ".join(f"{s['label']} {s['rate']}%" for s in opt)))

    B.append(("h1", "2. 즉시 조치가 필요한 사항"))
    over = [s for s in res["summary"] if s["overdue"]]
    if not over and not res["issues"]:
        B.append(("p", "기한이 지난 교육이 없습니다."))
    for s in over:
        names = [p["name"] for p in res["people"] if res["cells"][p["id"]][s["key"]].get("overdue")]
        B.append(("bullet", f"{s['label']}: 기한이 지났는데 미이수 {len(names)}명 – {', '.join(names[:12])}{' 외' if len(names) > 12 else ''}"))
    for i in res["issues"]:
        B.append(("bullet", f"[확인 필요] {i['text']}"))

    B.append(("h1", "3. 부서별 미이수 현황"))
    rows = []
    for dept in res["depts"]:
        ps = [p for p in active if p["dept"] == dept]
        miss = [(p, open_items(res, p["id"])) for p in ps]
        miss = [(p, it) for p, it in miss if it]
        cnt: dict[str, int] = {}
        for _, it in miss:
            for i in it:
                cnt[i["label"]] = cnt.get(i["label"], 0) + 1
        top = sorted(cnt.items(), key=lambda x: -x[1])[:3]
        rows.append([dept, str(len(ps)), str(len(miss)), ", ".join(f"{k} {v}명" for k, v in top) or "-"])
    B.append(("table", ["부서", "인원", "미이수자", "주요 미이수 교육"], rows, [2, 1, 1, 6]))

    B.append(("h1", "4. 향후 계획"))
    B.append(("bullet", "미이수자 본인 안내(기한 30일 전), 부서장 통보(기한 14일 전), 기한 경과 시 보고"))
    if plan_result and plan_result.get("sessions"):
        B.append(("bullet", f"{plan_result['col']['label']} 집합교육 {len(plan_result['sessions'])}회 운영(교대근무표 반영, 대상 {plan_result['targets']}명)"))
        B.append(("table", ["회차", "일시", "인원", "부서별"],
                  [[str(k), s["label"], str(s["count"]), ", ".join(f"{dept_core(d)} {n}" for d, n in s["by_dept"].items())]
                   for k, s in enumerate(plan_result["sessions"], 1)], [1, 4, 1, 5]))
    B.append(("note", f"※ 다이수가 자료 {len(res['files'])}개(직원 명단·이수 기록·서명부·수료증)를 근거로 작성했습니다. "
                      "직원별 판정 근거는 ‘이수현황 엑셀’에서 확인할 수 있습니다."))
    return B


def result_blocks(res: dict, org: dict, sign: dict) -> list:
    rows = [r for r in res["records"] if r["kind"] == "sign" and r["file"] == sign["file"] and r["where"].startswith(sign["where"] + " ")]
    attended = [r for r in rows if r["done"]]
    absent = [r for r in rows if not r["done"]]
    sv = survey_summary(survey_for(res, sign["title"] + " " + sign.get("context", "")))
    when = _d(sign["date"])
    B = [("title", f"{sign['title']} 결과 보고"),
         ("p", f"{org.get('name', '')} {org.get('dept', '')}"),
         ("h1", "1. 교육 개요"),
         ("table", ["구분", "내용"], [["교육명", sign["title"]], ["일시", (kdate(when) if when else "-") + (f" {sign['time']}" if sign.get("time") else "") + f" ({fmt_h(sign['hours'])}시간)"],
                                    ["장소", sign.get("place") or "-"],
                                    ["대상", f"{len(rows)}명"],
                                    ["참석", f"{len(attended)}명 (참석률 {round(100 * len(attended) / max(1, len(rows)), 1)}%)"],
                                    ["근거 자료", f"{sign['file']}"]], [1, 4])]
    B.append(("h1", "2. 만족도 조사 결과"))
    if sv:
        B.append(("p", f"응답 {sv['n']}명 · 전체 평균 {sv['overall']}점(5점 만점) · 만족(4점 이상) 비율 {sv['satisfied']}%"))
        B.append(("table", ["문항", "평균"], [[k, f"{v}점"] for k, v in sv["avg"].items()], [4, 1]))
        if sv["themes"]:
            B.append(("h2", "주요 의견"))
            for t in sv["themes"][:5]:
                B.append(("bullet", f"{t['theme']}({t['count']}건): “{t['examples'][0]}”"))
    else:
        B.append(("p", "이 교육과 연결된 만족도 응답을 찾지 못했습니다."))
    B.append(("h1", "3. 개선 사항"))
    imp = improvements(sv)
    for x in imp or ["특이사항 없음"]:
        B.append(("bullet", x))
    B.append(("h1", "4. 불참자 후속 조치"))
    if absent:
        B.append(("p", f"불참 {len(absent)}명: " + ", ".join(f"{r['name']}({dept_core(r['dept'])})" for r in absent)))
        B.append(("bullet", "다음 회차 또는 사이버교육으로 이수하도록 개별 안내"))
    else:
        B.append(("p", "불참자 없음"))
    B.append(("note", "※ 다이수가 서명부와 만족도 응답을 근거로 작성했습니다."))
    return B


def plan_blocks(res: dict, org: dict, pl: dict) -> list:
    B = [("title", f"{pl['col']['label']} 집합교육 일정 안내"),
         ("p", f"{org.get('name', '')} {org.get('dept', '')} · 대상 {pl['targets']}명 · 회차당 {fmt_h(pl['hours'])}시간"),
         ("h1", "1. 운영 원칙")] + [("bullet", r) for r in pl["rules"]]
    B.append(("h1", "2. 회차별 일정과 대상자"))
    for k, s in enumerate(pl["sessions"], 1):
        B.append(("h2", f"{k}회차 – {s['label']} ({s['count']}명)"))
        B.append(("p", ", ".join(f"{p['name']}({dept_core(p['dept'])}{('·' + p['team'] + '조') if p['team'] else ''})" for p in s["people"])))
    if pl["unplaced"]:
        B.append(("h1", "3. 집합교육 배정이 어려운 직원(사이버교육 안내)"))
        B.append(("table", ["성명", "부서", "사유"], [[u["name"], u["dept"], u["why"]] for u in pl["unplaced"]], [1, 1, 4]))
    return B


# ───────────── 엑셀·증빙 묶음
FILL = {"done": "DCEFE3", "partial": "FFF1C2", "none": "F8D7D3", "leave": "E3E8EF", "left": "EEEEEE", "na": "FFFFFF"}


def excel(res: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "교육별 요약"
    ws.append([f"{res['year']}년 법정의무교육 이수 현황 (기준일 {res['base']})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["교육", "근거", "대상", "이수", "부족", "미이수", "휴직", "이수율(%)", "기한", "기한 지남"])
    for s in res["summary"]:
        ws.append([s["label"] + ("" if s["mandatory"] else " (권장)"), s["law"], s["target"], s["done"], s["partial"], s["none"],
                   s["leave"], s["rate"], s["due"] or due_label(s), s["overdue"]])
    for col, w in zip("ABCDEFGHIJ", (30, 46, 7, 7, 7, 8, 7, 10, 12, 10)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("이수현황")
    head = ["사번", "성명", "소속", "근무", "직무"] + [c["label"] for c in res["columns"]]
    ws2.append(head)
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for p in res["people"]:
        row = [p["id"], p["name"], p["dept"], p["shift"] + (f"({p['team']})" if p.get("team") else ""), p["job"]]
        for c in res["columns"]:
            cell = res["cells"][p["id"]][c["key"]]
            if cell["s"] in ("done", "partial", "none"):
                row.append(f"{STATUS_KO[cell['s']]} {fmt_h(cell['h'])}/{fmt_h(cell['need'])}h" + (" ⚠기한지남" if cell.get("overdue") else ""))
            else:
                row.append(STATUS_KO[cell["s"]])
        ws2.append(row)
        for j, c in enumerate(res["columns"]):
            s = res["cells"][p["id"]][c["key"]]["s"]
            ws2.cell(ws2.max_row, 6 + j).fill = PatternFill("solid", fgColor=FILL[s])
    ws2.freeze_panes = "F2"
    for j in range(len(head)):
        ws2.column_dimensions[get_column_letter(j + 1)].width = 10 if j < 5 else 16
    ws2.row_dimensions[1].height = 45

    ws3 = wb.create_sheet("미이수자")
    ws3.append(["사번", "성명", "소속", "근무조", "교육", "남은 시간", "기한", "기한 지남"])
    for p in res["people"]:
        for i in open_items(res, p["id"], mandatory_only=False):
            ws3.append([p["id"], p["name"], p["dept"], p.get("team", ""), i["label"], i["remain"], i["due"], "예" if i["overdue"] else ""])

    ws4 = wb.create_sheet("판정 근거")
    ws4.append(["기록번호", "성명", "소속(기록)", "과정명", "이수일", "인정시간", "이수", "매칭", "판정 교육", "파일", "위치"])
    names = {r["id"]: r["name"] for r in res["rules"]["courses"]}
    for r in res["records"]:
        ws4.append([r["rid"], r["name"], r.get("dept", ""), r["title"], r.get("date") or "", r["hours"], "이수" if r["done"] else "미이수",
                    r.get("match") or "명단 없음", ", ".join(names[x] for x in r.get("rules", [])) or ("미분류" if not r.get("classified") else "의무 아님"),
                    r["file"], r["where"]])

    ws5 = wb.create_sheet("확인 필요")
    ws5.append(["구분", "내용"])
    for i in res["issues"]:
        ws5.append([i["type"], i["text"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reminders_excel(msgs: list[dict]) -> bytes:
    """문자·메신저 대량 발송 프로그램에 올릴 수 있는 형식(성명·소속·내용)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "안내문"
    ws.append(["사번", "성명", "소속", "근무조", "기한 지남", "내용"])
    for m in msgs:
        ws.append([m["id"], m["name"], m["dept"], m["team"], "예" if m["overdue"] else "", m["text"]])
    ws.column_dimensions["F"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def evidence_zip(res: dict, files_by_rel: dict[str, str]) -> bytes:
    """교육별 폴더에 이수를 인정한 원본 자료를 모으고, 증빙 목록 엑셀을 함께 넣는다(경영평가·감사 대비)."""
    from openpyxl import Workbook

    rec = {r["rid"]: r for r in res["records"]}
    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "증빙목록"
    ws.append(["교육", "근거", "대상", "이수", "이수율(%)", "증빙 파일", "이 파일로 인정된 인원"])
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for k, s in enumerate(res["summary"], 1):
            used: dict[str, set] = {}
            for p in res["people"]:
                for rid in res["cells"][p["id"]][s["key"]]["recs"]:
                    used.setdefault(rec[rid]["file"], set()).add(p["id"])
            folder = f"{k:02d}_{re.sub(r'[\\\\/:*?\"<>|]', '', s['label'])}"
            for rel, ids in sorted(used.items()):
                src = files_by_rel.get(rel)
                if src and os.path.exists(src):
                    z.write(src, f"{folder}/{os.path.basename(rel)}")
                ws.append([s["label"], s["law"], s["target"], s["done"], s["rate"], rel, len(ids)])
            if not used:
                ws.append([s["label"], s["law"], s["target"], s["done"], s["rate"], "(증빙 없음)", 0])
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["F"].width = 60
        xb = io.BytesIO()
        wb.save(xb)
        z.writestr("00_증빙목록.xlsx", xb.getvalue())
    return buf.getvalue()
