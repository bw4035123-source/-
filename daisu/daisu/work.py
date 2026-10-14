"""업무 일정(간트)과 업무별 체크리스트.

넣은 자료를 분석해 교육담당자가 1년 동안 할 업무를 기간 막대로 펼친다(계획·예산 → 법정·필수 교육 → 안내·독려 →
집합교육 운영 → 외부교육 → 보고·평가). 막대를 누르면 그 업무의 체크리스트를 만든다. 체크리스트 항목 중 자료로
확인되는 것(예: 미이수자 0명, 결과보고 제출)은 자동으로 체크되고, 나머지는 담당자가 체크한다.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io

from .util import kdate

LANES = ["전년도 이월", "계획·예산", "법정교육", "필수교육", "안내·독려", "집합교육 운영", "외부(위탁)교육", "보고·평가", "직접 추가"]
STATUS = {"done": "완료", "late": "지연", "doing": "진행 중", "todo": "예정", "past": "기간 지남"}


def _d(s):
    return dt.date.fromisoformat(s) if s else None


def _iso(d):
    return d.isoformat()


def _h(*parts) -> str:
    return hashlib.sha1("|".join(map(str, parts)).encode()).hexdigest()[:8]


def _status(start, end, base, done=None, late=False):
    if done:
        return "done"
    if late:
        return "late"
    if start <= base <= end:
        return "doing"
    return "todo" if base < start else "past"


# ───────────── 업무 분석 → 업무 목록
def analyze(project: dict, res: dict, ext: dict, remind: dict, spend: dict | None = None) -> list[dict]:
    y, base = res["year"], _d(res["base"])
    regs = res.get("regs") or {}
    sessions = project.get("sessions", [])
    exported = project.get("exported", {})
    out: list[dict] = []

    def add(id_, lane, title, start, end, detail="", link="", kind="generic", done=None, late=False, progress=None, ctx=None, marks=None):
        if end < start:
            start, end = end, start
        out.append({"id": id_, "lane": lane, "title": title, "start": _iso(start), "end": _iso(end), "detail": detail, "link": link,
                    "kind": kind, "status": _status(start, end, base, done, late), "progress": progress, "ctx": ctx or {},
                    "marks": marks or [], "base": res["base"]})

    # 전년도 이월(연도 릴레이): 지난해를 마감할 때 남은 미이수자를 올해 1분기 보충 업무로
    prev = (project.get("years") or {}).get(str(y - 1))
    if prev and prev["metrics"].get("unresolved_total"):
        un = prev["metrics"]["unresolved"]
        add(f"carry:{y - 1}", LANES[0], f"{y - 1}년 미이수자 보충·사유 정리 ({prev['metrics']['unresolved_total']}명·건)",
            dt.date(y, 1, 2), dt.date(y, 3, 31), " · ".join(f"{u['label']} {u['count']}명" for u in un[:4]) + (" 외" if len(un) > 4 else ""),
            "matrix", "carry", ctx={"unresolved": un, "prev": y - 1, "closed_at": prev.get("closed_at", "")})

    # 계획·예산
    gov, kp, req = regs.get("gov") or {}, regs.get("kpis") or [], regs.get("required") or []
    add("plan:annual", LANES[1], f"{y}년 교육훈련계획 수립", dt.date(y, 1, 2), dt.date(y, 1, 31),
        "전년 실적·정부 계획·편람 지표를 반영해 연간 계획을 세웁니다", "regs", "plan",
        ctx={"gov": bool(gov.get("priorities") or gov.get("rec_hours")), "kpis": len(kp), "required": len(req),
             "budget": bool(regs.get("budget"))})
    budget = (regs.get("budget") or {}).get("total")
    if budget:
        rate = round(100 * (spend or {}).get("total", 0) / budget, 1)
        add("budget", LANES[1], "교육훈련비 집행 관리", dt.date(y, 1, 1), dt.date(y, 12, 31),
            f"예산 {budget:,}원 · 집행 {(spend or {}).get('total', 0):,}원({rate}%)", "results", "budget", progress=min(100, rate),
            ctx={"budget": budget, "spent": (spend or {}).get("total", 0), "rate": rate,
                 "external": ext["total_cost"], "travel": ext["total_travel"] or 0})

    # 법정·필수 교육
    for s in res["summary"]:
        lane = LANES[3] if s.get("category") == "필수" else LANES[2]
        left = s["partial"] + s["none"]
        ctx = {k: s.get(k) for k in ("target", "done", "partial", "none", "leave", "overdue", "rate", "law", "group", "key", "due", "cycle")}
        ctx["left"] = left
        if s["cycle"] in ("hire", "promo"):
            add(f"course:{s['key']}", lane, s["label"], dt.date(y, 1, 1), dt.date(y, 12, 31),
                f"{'입사' if s['cycle'] == 'hire' else '승진'}일 기준 수시 · 대상 {s['target']}명 · 남은 {left}명",
                f"matrix/{s['key']}", "course_event", done=s["target"] > 0 and left == 0, late=bool(s["overdue"]),
                progress=s["rate"], ctx=ctx)
            continue
        start, due = _d(s["start"]), _d(s["due"] or s["end"])
        add(f"course:{s['key']}", lane, s["label"], start, due,
            f"대상 {s['target']}명 · 이수 {s['done']}명 · 남은 {left}명" + (f" · 기한 지남 {s['overdue']}명" if s["overdue"] else ""),
            f"matrix/{s['key']}", "course", done=s["target"] > 0 and left == 0, late=bool(s["overdue"]), progress=s["rate"], ctx=ctx,
            marks=[{"date": _iso(due), "label": "기한"}])

    # 안내·독려(기한별로 묶음)
    first, boss = remind.get("first_days", 30), remind.get("boss_days", 14)
    by_due: dict[str, list] = {}
    for s in res["summary"]:
        if s["due"] and s["mandatory"] and s["cycle"] not in ("hire", "promo"):
            by_due.setdefault(s["due"], []).append(s)
    for due, ss in sorted(by_due.items()):
        dd = _d(due)
        left = sum(s["partial"] + s["none"] for s in ss)
        names = ", ".join(s["label"] for s in ss[:3]) + (" 외" if len(ss) > 3 else "")
        add(f"remind:{due}", LANES[4], f"미이수자 안내·부서장 통보 ({kdate(dd, False)} 기한)", dd - dt.timedelta(days=first), dd,
            f"{names} · 남은 {left}명·건", "notify", "remind", done=left == 0, late=left > 0 and dd < base,
            ctx={"left": left, "first": first, "boss": boss, "names": names, "due": due},
            marks=[{"date": _iso(dd - dt.timedelta(days=boss)), "label": "부서장 통보"}])

    # 집합교육 운영: 다이수에서 만든 회차 + 자료로 들어온 서명부
    for s in sessions:
        sd = _d(s["date"])
        ex = {k: f"{s['id']}:{k}" in exported for k in ("plan", "result", "fee")}
        inst = s.get("instructor") or {}
        done = bool(s.get("attend")) and ex["result"]
        add(f"session:{s['id']}", LANES[5], s["title"], sd - dt.timedelta(days=14), sd + dt.timedelta(days=5),
            f"{kdate(sd)} {s.get('start', '')}~{s.get('end', '')} {s.get('place') or ''} · 대상 {len(s.get('targets', []))}명",
            f"session/{s['id']}", "session", done=done, late=not done and sd + dt.timedelta(days=5) < base,
            ctx={"exported": ex, "instructor": inst.get("name", ""), "grade": inst.get("grade", ""), "income": inst.get("income", ""),
                 "place": s.get("place", ""), "attend": len(s.get("attend", {})), "targets": len(s.get("targets", [])),
                 "survey": len(s.get("survey", []))},
            marks=[{"date": s["date"], "label": "교육일"}])
    seen = set()
    for g in res["signs"]:
        if not g.get("date") or (g["title"], g["date"]) in seen:
            continue
        seen.add((g["title"], g["date"]))
        sd = _d(g["date"])
        same = [x for x in res["signs"] if x["title"] == g["title"] and x.get("date") == g["date"]]
        n, att = sum(x["n"] for x in same), sum(x["attended"] for x in same)
        add(f"sign:{_h(g['title'], g['date'])}", LANES[5], g["title"], sd - dt.timedelta(days=14), sd + dt.timedelta(days=5),
            f"{kdate(sd)} 실시 · 출석 {att}/{n}명 · {g['file']}", "files", "sign", done=sd + dt.timedelta(days=5) < base,
            ctx={"n": n, "attended": att, "file": g["file"], "survey": bool(res["survey"])}, marks=[{"date": g["date"], "label": "교육일"}])

    # 외부(위탁)교육: 같은 과정·같은 기간은 한 업무로
    groups: dict[tuple, list] = {}
    for r in ext["rows"]:
        groups.setdefault((r["title"], r["start"] or r["end"], r["end"]), []).append(r)
    over = {x["pid"] for x in ext["over_limit"]}
    for (title, st, en), rows in sorted(groups.items(), key=lambda kv: kv[0][1] or ""):
        sd, ed = _d(st), _d(en)
        if not sd or not ed:
            continue
        dues = [_d(r["report_due"]) for r in rows if r.get("report_due")]
        end = max([ed] + dues)
        states = [r.get("report_state", "") for r in rows]
        missing = sum("미제출" in x for x in states)
        late_n = sum("지연" in x for x in states)
        submitted = all(x.startswith("기한 내") or "지연 제출" in x for x in states) if states and all(states) else False
        add(f"ext:{_h(title, st, en)}", LANES[6], title, sd - dt.timedelta(days=21), end,
            f"{rows[0]['inst']} · {len(rows)}명 · {kdate(sd)}~{kdate(ed)}" + (f" · 결과보고 기한 {kdate(max(dues))}" if dues else ""),
            "results", "external", done=submitted, late=missing > 0,
            ctx={"people": [r["name"] for r in rows], "n": len(rows), "cost": sum(r["cost"] or 0 for r in rows),
                 "travel": sum(r["travel"] or 0 for r in rows) if ext["has_travel_rule"] else None,
                 "groups": sorted({r["group"] for r in rows if r.get("group")}), "report_days": ext["report_days"], "limit": ext["limit"],
                 "over": [r["name"] for r in rows if r["pid"] in over], "missing": missing, "late": late_n,
                 "certs": True, "inst": rows[0]["inst"]},
            marks=[{"date": en, "label": "교육 종료"}] + ([{"date": _iso(max(dues)), "label": "결과보고 기한"}] if dues else []))

    # 보고·평가
    add("plan:next", LANES[7], f"{y + 1}년 교육훈련계획(안) 수립", dt.date(y, 11, 1), dt.date(y, 12, 31),
        "올해 실적·미달 지표·정부 중점 과제를 반영합니다", "report", "next")
    add("report:annual", LANES[7], f"{y}년 교육훈련 실적 보고", dt.date(y, 12, 1), dt.date(y + 1, 1, 15),
        "분야별 실적·외부교육·예산 집행·경영평가 지표를 보고서로 정리합니다", "report", "annual",
        ctx={"missing": len(ext["missing"])})
    add("report:kpi", LANES[7], "경영평가 교육훈련 지표 실적·증빙 정리", dt.date(y + 1, 1, 2), dt.date(y + 1, 2, 28),
        "제출 일정은 평가 편람·평가원 안내에 맞춰 바꾸세요", "report", "kpi", ctx={"kpis": len(kp)})
    if any(s["group"] == "폭력예방" for s in res["summary"]):
        add("report:violence", LANES[7], f"{y}년 폭력예방교육 실적 등록", dt.date(y + 1, 1, 2), dt.date(y + 1, 2, 28),
            "주무부처 시스템 등록", "report", "violence")

    # 사용자가 바꾼 기간·직접 추가한 업무
    w = project.get("work", {})
    for c in w.get("custom", []):
        add(c["id"], LANES[8], c["title"], _d(c["start"]), _d(c["end"]), c.get("detail", ""), "", "custom")
    for it in out:
        dates = w.get("dates", {}).get(it["id"])
        if dates:
            s, e = _d(dates[0]), _d(dates[1])
            it.update(start=dates[0], end=dates[1], moved=True)
            if it["status"] not in ("done", "late"):
                it["status"] = _status(s, e, base)
        ck = w.get("checks", {}).get(it["id"])
        if ck:
            items = checklist_view(it, ck)
            n_done = sum(1 for x in items if x["done"])
            it["check"] = {"done": n_done, "total": len(items)}
            if it["kind"] == "custom" or it["progress"] is None:
                it["progress"] = round(100 * n_done / max(1, len(items)))
                if items and n_done == len(items):
                    it["status"] = "done"
    return out


# ───────────── 체크리스트 틀(업무 종류별) – auto는 자료로 확인되는 항목
def _tpl(it: dict) -> list[dict]:
    c, k = it["ctx"], it["kind"]
    T = lambda text, hint="", auto=None, link="": {"text": text, "hint": hint, "auto": auto, "link": link}  # noqa: E731
    if k == "carry":
        un = c["unresolved"]
        return [T(f"{c['prev']}년 미이수자 명단 확인", f"{c['prev']}년 마감({c['closed_at']}) 시 {sum(u['count'] for u in un)}명·건", True, "matrix"),
                T("미이수 사유 확인(휴직·파견·퇴사 예정 등)"),
                T("보충 교육 일정 수립(사이버·집합)", "", None, "plan"),
                T("대상자 본인·부서장 안내", "", None, "notify"),
                T("보충 이수 결과 확인·기록"),
                T(f"{c['prev']}년 실적 보고서에 보충 결과 반영", "", None, "report")]
    if k == "plan":
        return [T("전년도 교육 실적·만족도·미달 지표 분석", link="results"),
                T("정부(행안부) 교육훈련 운영 계획의 중점 과제 반영", "규정·기준 탭에서 읽은 중점 과제", c.get("gov"), "regs"),
                T("경영평가 편람 교육훈련 지표·산식 확인", f"편람에서 지표 {c.get('kpis', 0)}개를 읽었습니다", bool(c.get("kpis")), "regs"),
                T("법정·필수 교육 대상과 시간 확정(내규 별표 포함)", f"내규 필수교육 {c.get('required', 0)}개 과정", bool(c.get("required")), "regs"),
                T("교육훈련 예산 편성", "예산표 확인", c.get("budget")),
                T("1인당 교육시간 목표 설정(내규·정부 권장 비교)", link="results"),
                T("연간 계획(안) 기안·결재"),
                T("전 직원 공지(연간 교육 일정)")]
    if k == "budget":
        return [T("분기별 집행 현황 점검", f"집행률 {c['rate']}%"),
                T("외부교육비 집행 내역 정리", f"{c['external']:,}원", True, "results"),
                T("교육 여비 산정·정산(여비규정)", f"{c['travel']:,}원"),
                T("강사료 지급 내역 정리", link="run"),
                T("연말 집행률 확인·불용 사유 정리", f"현재 {c['rate']}%", c["rate"] >= 90)]
    if k == "course":
        left, over = c["left"], c["overdue"]
        items = [T("대상자 명단 확정(휴직·퇴사·신규 반영)", f"대상 {c['target']}명" + (f", 휴직 {c['leave']}명 제외" if c["leave"] else ""),
                   c["target"] > 0, "matrix/" + c["key"]),
                 T("근거 법령·내규의 시간·주기 확인", c.get("law") or "", None, "regs"),
                 T("교육 방법 결정(사이버·집합·외부강사)"),
                 T("실시계획 기안·결재"),
                 T("교대 근무자 집합교육 일정 편성(근무표 반영)", "일정 짜기에서 근무표를 피해 회차를 만듭니다", None, "plan/" + c["key"]),
                 T("대상자 수강 안내 발송", "", None, "notify"),
                 T("미이수자 본인 안내(기한 전)", f"남은 {left}명", left == 0, "notify"),
                 T("부서장 미이수 통보", f"남은 {left}명", left == 0, "notify"),
                 T("이수 증빙 수합(수료증·서명부)", f"이수 {c['done']}/{c['target']}명", c["target"] > 0 and left == 0, "report")]
        if c.get("due") and c["due"] < it.get("base", "9999"):
            items.append(T("기한 경과 미이수자 사유 확인·보충 교육", f"기한 지남 {over}명", over == 0, "matrix/over"))
        if c.get("group") == "폭력예방":
            items.append(T("주무부처 시스템 실적 등록", "다음 해 2월 말까지", None, "report"))
        items.append(T("이수 실적 결과 보고", "", None, "report"))
        return items
    if k == "course_event":
        return [T("대상자(입사·승진) 발생 확인", f"대상 {c['target']}명", c["target"] > 0, "matrix/" + c["key"]),
                T("배치 전·기한 내 교육 일정 안내"),
                T("교육 실시·이수 확인", f"남은 {c['left']}명", c["left"] == 0, "matrix/" + c["key"]),
                T("기한 지난 대상자 보충", f"기한 지남 {c['overdue']}명", c["overdue"] == 0),
                T("이수 기록·증빙 보관")]
    if k == "remind":
        return [T("미이수자 명단 추출", f"{c['names']} · 남은 {c['left']}명", True, "notify"),
                T(f"본인 안내문 발송(기한 {c['first']}일 전)", "", c["left"] == 0, "notify"),
                T(f"부서장 통보(기한 {c['boss']}일 전)", "", c["left"] == 0, "notify"),
                T("교대 근무자 별도 회차 안내", "", None, "plan"),
                T("기한 3일 전 재확인·개별 연락", "", c["left"] == 0),
                T("기한 당일 최종 점검", f"남은 {c['left']}명", c["left"] == 0, "board")]
    if k == "session":
        ex = c["exported"]
        items = [T("실시계획 기안·결재", "", ex["plan"], it["link"]),
                 T("강사 섭외·강사수당 등급 확인", (f"{c['instructor']} · {c['grade']}" if c["instructor"] else "강사 미정"), bool(c["instructor"]), it["link"]),
                 T("교육장·기자재 예약", c["place"] or "장소 미정", bool(c["place"])),
                 T("대상자 선정·근무 조정 요청", f"대상 {c['targets']}명", c["targets"] > 0),
                 T("교육 안내 발송", "", None, "notify"),
                 T("출석 키오스크·서명부 준비", "", None, it["link"]),
                 T("교육 실시·출석 확인", f"출석 {c['attend']}/{c['targets']}명", c["attend"] > 0, it["link"]),
                 T("만족도 조사", f"응답 {c['survey']}건", c["survey"] > 0, it["link"]),
                 T("결과 보고 기안", "", ex["result"], it["link"])]
        if c["instructor"] and c["income"] != "근로":
            items.append(T("강사료 지급 기안(원천징수)", "", ex["fee"], it["link"]))
        items.append(T("이수 기록 반영", "", c["attend"] > 0))
        return items
    if k == "sign":
        return [T("서명부 수합", f"출석 {c['attended']}/{c['n']}명 · {c['file']}", True, "files"),
                T("불참자 확인·보충 교육 안내", f"불참 {c['n'] - c['attended']}명", c["n"] == c["attended"], "matrix/over"),
                T("만족도 조사 결과 정리", "", c["survey"] or None, "report"),
                T("결과 보고 기안", "", None, "report"),
                T("이수 기록 반영", "", True, "matrix")]
    if k == "external":
        tr = c.get("travel")
        return [T("교육 수요 조사·대상자 선정", ", ".join(c["people"][:6]) + (" 외" if c["n"] > 6 else ""), True),
                T("교육 신청·출장(복무) 기안", c.get("inst", "")),
                T("교육비 지출 결의", f"{c['cost']:,}원" + (f" · 1인 한도 {c['limit']:,}원" if c.get("limit") else "")),
                T("지원 한도 초과 여부 확인", ("초과: " + ", ".join(c["over"])) if c["over"] else "초과자 없음", not c["over"], "results"),
                T("교육 여비 산정(여비규정)", (f"{tr:,}원 · " + "/".join(c["groups"])) if tr is not None else "여비규정을 넣으면 계산합니다",
                  tr is not None, "results"),
                T("수료증 수합", "", None, "files"),
                T(f"결과보고서 제출 확인(종료 후 {c['report_days']}일)" if c.get("report_days") else "결과보고서 제출 확인",
                  (f"미제출 {c['missing']}명" if c["missing"] else "") + (f" 지연 {c['late']}명" if c["late"] else ""),
                  c["missing"] == 0 and c.get("report_days") is not None, "results"),
                T("여비 정산(교통비 실비 증빙)", "", None, "report"),
                T("교육시간 실적 반영", "", True, "results")]
    if k == "annual":
        return [T("분야별(법정·필수·직무·외부) 실적 집계", "", True, "results"),
                T("외부교육 결과보고 누락 점검", f"미제출 {c['missing']}건", c["missing"] == 0, "results"),
                T("교육훈련비 집행 실적 정리", "", None, "results"),
                T("실적 보고서 작성(다이수 보고서)", "", None, "report"),
                T("보고서 기안·결재")]
    if k == "kpi":
        return [T("편람 지표·산식 확인", f"지표 {c['kpis']}개", bool(c["kpis"]), "regs"),
                T("지표별 실적 산출·근거 정리", "", True, "results"),
                T("증빙 자료 묶음(이수 증빙 ZIP)", "", None, "report"),
                T("제출 일정·양식 확인(평가원 안내)"),
                T("실적 제출")]
    if k == "next":
        return [T("올해 미달 지표·미이수 원인 분석", "", None, "results"),
                T("정부 중점 과제 반영", "", None, "regs"),
                T("내년 계획(안) 작성(다이수 보고서)", "", None, "report"),
                T("예산 요구"),
                T("계획(안) 기안·결재")]
    if k == "violence":
        return [T("4대 폭력예방교육 이수 현황 엑셀", "", None, "report"), T("주무부처 시스템 등록"), T("등록 결과 보관")]
    return [T("계획 수립"), T("기안·결재"), T("실시"), T("결과 보고")]


def build_checklist(it: dict) -> dict:
    items = []
    for i, t in enumerate(_tpl(it), 1):
        items.append({"id": f"t{i}", "text": t["text"], "hint": t["hint"], "link": t["link"], "auto": t["auto"] is not None, "checked": False})
    return {"created": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "items": items}


def checklist_view(it: dict, ck: dict) -> list[dict]:
    """저장된 체크리스트 + 지금 자료로 다시 확인한 자동 항목."""
    live = {t["text"]: t for t in _tpl(it)}
    out = []
    for x in ck["items"]:
        t = live.get(x["text"])
        auto_ok = bool(t and t["auto"])
        hint = (t["hint"] if t else x.get("hint", "")) or ""
        out.append(dict(x, hint=hint, auto_ok=auto_ok, done=bool(x.get("checked") or auto_ok)))
    return out


# ───────────── 엑셀(간트 + 체크리스트)
def excel(items: list[dict], checks: dict, year: int, base: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = {"done": "03A564", "late": "F04452", "doing": "3182F6", "todo": "9FC5FB", "past": "B0B8C1"}
    months = [(year, m) for m in range(1, 13)] + [(year + 1, 1), (year + 1, 2)]
    wb = Workbook()
    ws = wb.active
    ws.title = "업무 일정"
    ws.append([f"{year}년 교육훈련 업무 일정 (기준일 {base})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["업무 구분", "업무", "시작", "끝", "상태", "진행(%)", "체크리스트"] + [f"{yy % 100}.{m}월" for yy, m in months])
    for x in items:
        ck = x.get("check")
        row = [x["lane"], x["title"], x["start"], x["end"], STATUS[x["status"]], x["progress"] if x["progress"] is not None else "",
               f"{ck['done']}/{ck['total']}" if ck else ""]
        ws.append(row)
        r = ws.max_row
        s, e = x["start"][:7], x["end"][:7]
        for j, (yy, m) in enumerate(months):
            ym = f"{yy}-{m:02d}"
            if s <= ym <= e:
                ws.cell(r, 8 + j).fill = PatternFill("solid", fgColor=fill[x["status"]])
    for col, wdt in zip("ABCDEFG", (14, 44, 11, 11, 9, 8, 10)):
        ws.column_dimensions[col].width = wdt
    for j in range(len(months)):
        ws.column_dimensions[ws.cell(2, 8 + j).column_letter].width = 6
    ws.freeze_panes = "H3"
    ws2 = wb.create_sheet("체크리스트")
    ws2.append(["업무", "기간", "번호", "항목", "확인", "참고"])
    by_id = {x["id"]: x for x in items}
    for iid, ck in checks.items():
        it = by_id.get(iid)
        if not it:
            continue
        for n, c in enumerate(checklist_view(it, ck), 1):
            ws2.append([it["title"], f"{it['start']}~{it['end']}", n, c["text"], "☑" if c["done"] else "☐",
                        ("자료로 확인 · " if c["auto_ok"] else "") + (c["hint"] or "")])
    if ws2.max_row == 1:
        ws2.append(["(아직 만든 체크리스트가 없습니다 – 업무 일정 화면에서 막대를 눌러 만드세요)"])
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["D"].width = 44
    ws2.column_dimensions["F"].width = 50
    for row in ws2.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
