"""계산 결과 → 제출 양식(한글 hwpx · 워드 docx)의 표에 채워 문서 만들기.

기관·연도마다 제출 양식이 조금씩 달라서, 양식을 코드에 박아 두지 않고 '매핑'으로 연결한다.
  · 양식 파일(hwpx/docx)을 그대로 받는다. 표는 '머리글 행의 글자'로 찾는다(위치·순서가 달라도 됨).
  · 머리글 바로 아래 행을 '서식 행'으로 보고, 데이터 수만큼 복제해 채운다(글꼴·테두리·칸 너비 유지).
  · 본문의 {{기관명}} 같은 자리표시는 값으로 바꾼다.
  · 양식 파일이 없으면 build_default_template() 로 [예시] 양식을 만든다(실제 제출 양식 아님 – 기관 양식으로 바꿔 쓰세요).

매핑 예(TemplateMap)
  placeholders: {"{{기관명}}": "org", "{{연도}}": "year"}
  tables: [{"source": "depts", "columns": {"부서": "dept", "인원": "headcount", ...}, "total_row": True}]
"""
from __future__ import annotations

import copy
import io
import re
import zipfile
from decimal import Decimal

from pydantic import BaseModel, Field

from .calc import PersonResult, dept_summary
from .schema import Category

HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"


class TableMap(BaseModel):
    source: str = Field(pattern="^(depts|persons)$")
    columns: dict[str, str] = Field(description="머리글 글자 → 데이터 필드")
    total_row: bool = False
    match_min: int = Field(default=3, ge=1, description="머리글 중 몇 개 이상 일치해야 이 표로 볼지")


class TemplateMap(BaseModel):
    placeholders: dict[str, str] = Field(default_factory=dict)
    tables: list[TableMap] = Field(default_factory=list)


DEPT_COLS = {"연번": "no", "부서": "dept", "인원": "headcount", "기본교육": "basic", "전문교육": "prof", "자기개발": "self",
             "총 이수시간": "total_hours", "1인당 평균": "avg_hours", "의무시간 충족": "met_count", "충족률(%)": "met_rate"}
PERSON_COLS = {"연번": "no", "부서": "dept", "성명": "name", "직렬": "series", "직급": "grade", "기본교육": "basic",
               "전문교육": "prof", "자기개발": "self", "합계": "total", "의무시간": "required", "미달시간": "shortfall",
               "가점": "bonus"}
DEFAULT_MAP = TemplateMap(
    placeholders={"{{기관명}}": "org", "{{연도}}": "year", "{{기준일}}": "base_date", "{{총인원}}": "people",
                  "{{총이수시간}}": "total_hours", "{{1인당평균}}": "avg_hours", "{{충족률}}": "met_rate"},
    tables=[TableMap(source="depts", columns=DEPT_COLS, total_row=True), TableMap(source="persons", columns=PERSON_COLS)])


def _fmt(v) -> str:
    if v is None:
        return "-"
    if isinstance(v, Decimal):
        v = v.normalize()
        return f"{v:,f}" if v == v.to_integral() else f"{v:,f}".rstrip("0").rstrip(".")
    if isinstance(v, (int, float)):
        return f"{v:,}"
    return str(v)


# ───────────── 계산 결과 → 표 데이터
def report_data(results: list[PersonResult], org: str, year: int, base_date: str) -> dict:
    ds = dept_summary(results)
    n = len(results)
    tot = sum((r.total for r in results), Decimal(0))
    judged = [r for r in results if r.met is not None]
    met = sum(1 for r in judged if r.met)
    depts = [{"no": i, "dept": d.dept, "headcount": d.headcount, "basic": d.by_category[Category.BASIC],
              "prof": d.by_category[Category.PROFESSIONAL], "self": d.by_category[Category.SELF_DEV], "total_hours": d.total_hours,
              "avg_hours": d.avg_hours, "met_count": d.met_count, "met_rate": d.met_rate} for i, d in enumerate(ds, 1)]
    total = {"no": "", "dept": "합계", "headcount": n,
             **{k: sum((d[k] for d in depts), Decimal(0)) for k in ("basic", "prof", "self", "total_hours")},
             "avg_hours": (tot / n).quantize(Decimal("0.01")) if n else Decimal(0), "met_count": met,
             "met_rate": (Decimal(100) * met / len(judged)).quantize(Decimal("0.1")) if judged else None}
    persons = [{"no": i, "dept": r.profile.dept, "name": r.profile.name, "series": r.profile.series, "grade": r.profile.grade,
                "basic": r.by_category[Category.BASIC], "prof": r.by_category[Category.PROFESSIONAL], "self": r.by_category[Category.SELF_DEV],
                "total": r.total, "required": r.total_required if r.total_required is not None else (sum(r.required.values(), Decimal(0)) if r.required else None),
                "shortfall": r.total_shortfall, "bonus": r.bonus_points} for i, r in enumerate(results, 1)]
    return {"org": org, "year": str(year), "base_date": base_date, "people": n, "total_hours": total["total_hours"],
            "avg_hours": total["avg_hours"], "met_rate": total["met_rate"], "depts": depts, "depts_total": total, "persons": persons}


def _rows_for(tm: TableMap, data: dict) -> list[dict]:
    rows = list(data[tm.source])
    if tm.total_row and tm.source == "depts":
        rows.append(data["depts_total"])
    return rows


def _norm(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def _match(header_cells: list[str], tm: TableMap) -> dict[int, str] | None:
    """머리글 칸 위치 → 필드. 일치 개수가 match_min 미만이면 None."""
    cols = {_norm(k): v for k, v in tm.columns.items()}
    pos = {i: cols[_norm(h)] for i, h in enumerate(header_cells) if _norm(h) in cols}
    return pos if len(pos) >= min(tm.match_min, len(cols)) else None


def _replace_all(text: str, mapping: TemplateMap, data: dict) -> str:
    for ph, key in mapping.placeholders.items():
        if ph in text:
            text = text.replace(ph, _fmt(data.get(key)))
    return text


# ───────────── 워드(docx)
def fill_docx(template: bytes, data: dict, mapping: TemplateMap = DEFAULT_MAP) -> bytes:
    from docx import Document

    doc = Document(io.BytesIO(template))

    def fix_paragraph(p):
        full = "".join(r.text for r in p.runs)
        new = _replace_all(full, mapping, data)
        if new != full and p.runs:
            p.runs[0].text = new  # 자리표시가 여러 run 으로 나뉘어도 첫 run 서식으로 합침
            for r in p.runs[1:]:
                r.text = ""

    for p in doc.paragraphs:
        fix_paragraph(p)
    used = set()
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    fix_paragraph(p)
        for hi, row in enumerate(t.rows[:3]):
            hdr = [c.text for c in row.cells]
            for k, tm in enumerate(mapping.tables):
                if k in used:
                    continue
                pos = _match(hdr, tm)
                if not pos:
                    continue
                used.add(k)
                tpl = t.rows[hi + 1]._tr if len(t.rows) > hi + 1 else copy.deepcopy(row._tr)
                tpl = copy.deepcopy(tpl)
                for extra in list(t.rows[hi + 1:]):
                    t._tbl.remove(extra._tr)
                for rec in _rows_for(tm, data):
                    tr = copy.deepcopy(tpl)
                    t._tbl.append(tr)
                    cells = t.rows[-1].cells
                    for i, cell in enumerate(cells):
                        val = _fmt(rec.get(pos[i])) if i in pos else ""
                        ps = cell.paragraphs
                        if ps and ps[0].runs:
                            ps[0].runs[0].text = val
                            for r in ps[0].runs[1:]:
                                r.text = ""
                        elif ps:
                            ps[0].add_run(val)
                        for extra_p in ps[1:]:
                            extra_p._p.getparent().remove(extra_p._p)
                break
            else:
                continue
            break
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ───────────── 한글(hwpx)
def _t_texts(el) -> str:
    return "".join(t.text or "" for t in el.iter(f"{{{HP}}}t"))


def _set_cell(tc, value: str) -> None:
    ts = list(tc.iter(f"{{{HP}}}t"))
    if ts:
        ts[0].text = value
        for t in ts[1:]:
            t.text = ""
        return
    run = next(tc.iter(f"{{{HP}}}run"), None)
    if run is not None:
        t = run.makeelement(f"{{{HP}}}t", {})
        t.text = value
        run.append(t)


def fill_hwpx(template: bytes, data: dict, mapping: TemplateMap = DEFAULT_MAP) -> bytes:
    from lxml import etree

    zin = zipfile.ZipFile(io.BytesIO(template))
    out = io.BytesIO()
    used: set[int] = set()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            raw = zin.read(info.filename)
            if re.fullmatch(r"Contents/section\d+\.xml", info.filename):
                root = etree.fromstring(raw)
                # 자리표시: 문단 안에서 run 이 나뉘어 있어도 합쳐서 바꿈
                for p in root.iter(f"{{{HP}}}p"):
                    ts = [t for t in p.iter(f"{{{HP}}}t") if t.getparent().getparent() is p]
                    full = "".join(t.text or "" for t in ts)
                    if "{{" in full:
                        new = _replace_all(full, mapping, data)
                        if new != full and ts:
                            ts[0].text = new
                            for t in ts[1:]:
                                t.text = ""
                for tbl in list(root.iter(f"{{{HP}}}tbl")):
                    trs = tbl.findall(f"{{{HP}}}tr")
                    for hi, tr in enumerate(trs[:3]):
                        hdr = [_t_texts(tc) for tc in tr.findall(f"{{{HP}}}tc")]
                        hit = next(((k, tm, _match(hdr, tm)) for k, tm in enumerate(mapping.tables)
                                    if k not in used and _match(hdr, tm)), None)
                        if not hit:
                            continue
                        k, tm, pos = hit
                        used.add(k)
                        tpl = copy.deepcopy(trs[hi + 1] if len(trs) > hi + 1 else tr)
                        for extra in trs[hi + 1:]:
                            tbl.remove(extra)
                        anchor = tr
                        for n, rec in enumerate(_rows_for(tm, data), hi + 1):
                            new = copy.deepcopy(tpl)
                            for i, tc in enumerate(new.findall(f"{{{HP}}}tc")):
                                _set_cell(tc, _fmt(rec.get(pos[i])) if i in pos else "")
                                addr = tc.find(f"{{{HP}}}cellAddr")
                                if addr is not None:
                                    addr.set("rowAddr", str(n))
                            anchor.addnext(new)
                            anchor = new
                        tbl.set("rowCnt", str(len(tbl.findall(f"{{{HP}}}tr"))))
                        break
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            # mimetype 은 압축 없이 맨 앞(한글 프로그램 요구)
            ct = zipfile.ZIP_STORED if info.filename == "mimetype" else zipfile.ZIP_DEFLATED
            zout.writestr(zipfile.ZipInfo(info.filename, date_time=info.date_time), raw, compress_type=ct)
    return out.getvalue()


# ───────────── [예시] 양식 만들기 · 한 번에 렌더링
def build_default_template(fmt: str = "hwpx", mapping: TemplateMap = DEFAULT_MAP) -> bytes:
    from ..render import FORMATS

    dcols = next(t for t in mapping.tables if t.source == "depts").columns
    pcols = next(t for t in mapping.tables if t.source == "persons").columns
    blocks = [("title", "{{연도}}년 교육훈련 실적 제출 – {{기관명}}"),
              ("note", "※ [예시] 양식입니다. 실제 제출 양식(행정안전부·주무부처 안내 서식)을 받으면 그 파일을 template 으로 넣으세요."),
              ("p", "기준일: {{기준일}} · 총 인원 {{총인원}}명 · 총 이수시간 {{총이수시간}}시간 · 1인당 평균 {{1인당평균}}시간 · 의무시간 충족률 {{충족률}}%"),
              ("h2", "1. 부서별 실적"), ("table", list(dcols), [[""] * len(dcols)]),
              ("h2", "2. 개인별 실적"), ("table", list(pcols), [[""] * len(pcols)])]
    return FORMATS[fmt][0](blocks)


def render(results: list[PersonResult], org: str, year: int, base_date: str, fmt: str = "hwpx",
           template: bytes | None = None, mapping: TemplateMap = DEFAULT_MAP) -> bytes:
    if fmt not in ("hwpx", "docx"):
        raise ValueError("형식은 hwpx 또는 docx 입니다.")
    data = report_data(results, org, year, base_date)
    tpl = template or build_default_template(fmt, mapping)
    return fill_hwpx(tpl, data, mapping) if fmt == "hwpx" else fill_docx(tpl, data, mapping)
