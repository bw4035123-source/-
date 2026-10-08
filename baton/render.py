"""문서 렌더러: 블록 목록 하나로 한글(.hwpx)·워드(.docx)·마크다운(.md)을 모두 만든다.

블록 형식
  ("title", "문서 제목")
  ("h1" | "h2", "제목")
  ("p", "문단")          # **굵게** 표기 지원
  ("bullet", "항목")
  ("table", 머리글목록, 행목록[, 열너비비율])
  ("note", "※ 참고")
"""
from __future__ import annotations

import io
import re
import warnings

RICH = re.compile(r"(\*\*.+?\*\*)")


def _segments(text: str):
    for part in RICH.split(str(text or "")):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            yield part[2:-2], True
        else:
            yield part, False


def to_md(blocks) -> bytes:
    out = []
    for b in blocks:
        kind = b[0]
        if kind == "title":
            out += [f"# {b[1]}", ""]
        elif kind in ("h1", "h2"):
            out += ["#" * (int(kind[1]) + 1) + " " + b[1], ""]
        elif kind == "p":
            out += [b[1], ""]
        elif kind == "bullet":
            out.append(f"- {b[1]}")
        elif kind == "note":
            out += [f"> {b[1]}", ""]
        elif kind == "table":
            esc = lambda s: str(s).replace("|", "\\|").replace("\n", "<br>")  # noqa: E731
            out.append("| " + " | ".join(esc(h) for h in b[1]) + " |")
            out.append("|" + "---|" * len(b[1]))
            out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in b[2]]
            out.append("")
    return ("\n".join(out) + "\n").encode("utf-8")


def to_docx(blocks) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(10.5)

    def rich(par, text, size=None):
        for seg, bold in _segments(text):
            r = par.add_run(seg)
            r.bold = bold or None
            if size:
                r.font.size = Pt(size)

    for b in blocks:
        kind = b[0]
        if kind == "title":
            d.add_heading(b[1], level=0)
        elif kind in ("h1", "h2"):
            d.add_heading(b[1], level=int(kind[1]))
        elif kind == "p":
            for line in str(b[1]).split("\n"):
                rich(d.add_paragraph(), line)
        elif kind == "note":
            p = d.add_paragraph()
            r = p.add_run(b[1])
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif kind == "bullet":
            rich(d.add_paragraph(style="List Bullet"), b[1])
        elif kind == "table":
            header, rows = b[1], b[2]
            t = d.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for c, hname in enumerate(header):
                t.rows[0].cells[c].text = str(hname)
                for r in t.rows[0].cells[c].paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for c in range(len(header)):
                    cells[c].text = ""
                    rich(cells[c].paragraphs[0], row[c] if c < len(row) and row[c] is not None else "", size=9.5)
            d.add_paragraph()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_hwpx(blocks) -> bytes:
    """한글(.hwpx). python-hwpx(Apache-2.0)로 만들어 한글 프로그램에서 바로 열린다."""
    from hwpx import HwpxDocument

    warnings.filterwarnings("ignore", category=DeprecationWarning)
    doc = HwpxDocument.new()
    styles: dict = {}

    def style(bold=False, size=None):
        key = (bold, size)
        if key not in styles:
            styles[key] = doc.ensure_run_style(bold=bold, size=size)
        return styles[key]

    def rich(par, text, size=None, bold_all=False):
        for seg, bold in _segments(text):
            par.add_run(seg, char_pr_id_ref=style(bold or bold_all, size))

    first = True
    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = doc.paragraphs[0] if first else doc.add_paragraph("", include_run=False)
            rich(p, b[1], size=16, bold_all=True)
        elif kind in ("h1", "h2"):
            p = doc.add_paragraph("", include_run=False)
            rich(p, b[1], size=13 if kind == "h1" else 11, bold_all=True)
        elif kind in ("p", "note", "bullet"):
            text = ("• " + b[1]) if kind == "bullet" else b[1]
            for line in str(text).split("\n"):
                rich(doc.add_paragraph("", include_run=False), line, size=9 if kind == "note" else None)
        elif kind == "table":
            header, rows = b[1], b[2]
            t = doc.add_table(len(rows) + 1, len(header))
            for c, hname in enumerate(header):
                t.set_cell_text(0, c, str(hname))
                t.set_cell_shading(0, c, "EAF1FB")
            for ri, row in enumerate(rows, 1):
                for c in range(len(header)):
                    val = str(row[c]) if c < len(row) and row[c] is not None else ""
                    t.set_cell_text(ri, c, val.replace("**", ""))
            if len(b) > 3 and b[3]:
                try:
                    total = sum(b[3])
                    t.set_column_widths([int(42520 * w / total) for w in b[3]])  # A4 본문 폭(HWPUNIT)
                except Exception:
                    pass
            doc.add_paragraph("")
        first = False
    return doc.to_bytes()


FORMATS = {
    "hwpx": (to_hwpx, "application/hwp+zip"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "md": (to_md, "text/markdown; charset=utf-8"),
}
