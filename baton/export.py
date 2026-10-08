"""인수인계서 내보내기: Markdown · Word(DOCX, 한글에서도 열림) · 인쇄용 HTML.

각 항목 뒤에 각주 번호를 달고, 문서 끝에 '근거 목록(파일·위치)'을 붙여 출처를 끝까지 추적할 수 있게 한다.
"""
from __future__ import annotations

import html
import io

from .ingest import KIND_LABEL

TRUST_LABEL = {"official": "공식", "doc": "문서", "mail": "메일", "memo": "메모", "oral": "구술", "none": "근거없음"}
STATUS_LABEL = {"verified": "확인", "edited": "수정", "ai": "미검수", "added": "추가"}


def _collect(project: dict):
    """(섹션 목록, 근거 번호표) – 삭제·근거없음 항목 제외."""
    draft = project["draft"]
    chunk_map = {c["id"]: c for c in project["chunks"]}
    qmap = {q["id"]: q for q in draft["questions"]}
    notes: dict[str, int] = {}
    secs = []
    for s in draft["sections"]:
        items = []
        for it in s["items"]:
            if it["status"] in ("deleted", "unsupported"):
                continue
            nums = []
            for sid in it["sources"]:
                if sid in chunk_map or sid in qmap:
                    notes.setdefault(sid, len(notes) + 1)
                    nums.append(notes[sid])
            items.append((it, nums))
        secs.append((s, items))
    refs = []
    for sid, n in sorted(notes.items(), key=lambda x: x[1]):
        if sid in qmap:
            refs.append((n, sid, f"전임자 구술 – {qmap[sid]['q'][:60]} ({qmap[sid].get('answered_at', '')})"))
        else:
            c = chunk_map[sid]
            refs.append((n, sid, f"{c['file']} · {c['where']} [{KIND_LABEL.get(c['kind'], '')}]"))
    return secs, refs


def _meta_rows(project: dict) -> list[tuple[str, str]]:
    p, d = project, project["draft"]
    h = p.get("handover", {})
    integ = p.get("integrity", {})
    return [
        ("기관·부서", f"{p.get('org', '')} {p.get('dept', '')}".strip()),
        ("업무명", p.get("name", "")),
        ("전임자 → 후임자", f"{p.get('from_name', '')} → {p.get('to_name', '')}"),
        ("인수인계 기준일", p.get("base_date", "")),
        ("분석 자료", f"{len(p['docs'])}개 파일 / 근거조각 {len(p['chunks'])}개"),
        ("초안 생성", f"{d.get('generated_by', '')} ({d.get('generated_at', '')})"),
        ("원본 무결성", "변경 없음 확인" if integ.get("ok") else "확인 전"),
        ("전임자 확인", h.get("from_signed_at", "미확인")),
        ("후임자 수령", h.get("to_signed_at", "미수령")),
    ]


def _item_text(it: dict) -> str:
    t = it["text"]
    m = it["meta"]
    if m.get("next") and m["next"] not in t:
        t += f" → 다음 할 일: {m['next']}"
    if m.get("due"):
        t += f" (기한: {m['due']})"
    if m.get("resolved"):
        t += f" ⇒ 전임자 확인: {m['resolved']}"
    return t


def to_markdown(project: dict) -> str:
    secs, refs = _collect(project)
    out = [f"# {project.get('title') or '업무 인수인계서'} – {project.get('name', '')}", ""]
    out += ["| 구분 | 내용 |", "|---|---|"] + [f"| {k} | {v} |" for k, v in _meta_rows(project)] + [""]
    for s, items in secs:
        out.append(f"## {s['title']}")
        if not items:
            out.append("- (해당 없음)")
        for it, nums in items:
            foot = "".join(f"[^{n}]" for n in nums)
            out.append(f"- [{TRUST_LABEL[it['trust']]}·{STATUS_LABEL.get(it['status'], '')}] {_item_text(it)}{foot}")
        out.append("")
    qs = [q for q in project["draft"]["questions"] if q.get("answer")]
    if qs:
        out.append("## 부록. 전임자 인터뷰 기록")
        for q in qs:
            out += [f"- **Q.** {q['q']}", f"  - **A.** {q['answer']} ({q.get('answered_at', '')})"]
        out.append("")
    out.append("## 근거 목록")
    out += [f"[^{n}]: {sid} – {label}" for n, sid, label in refs]
    return "\n".join(out) + "\n"


def to_html(project: dict) -> str:
    secs, refs = _collect(project)
    e = html.escape
    rows = "".join(f"<tr><th>{e(k)}</th><td>{e(str(v))}</td></tr>" for k, v in _meta_rows(project))
    body = []
    for s, items in secs:
        lis = "".join(
            f"<li><span class='tag {it['trust']}'>{TRUST_LABEL[it['trust']]}</span> {e(_item_text(it))}"
            + "".join(f"<sup>{n}</sup>" for n in nums) + "</li>" for it, nums in items) or "<li>(해당 없음)</li>"
        body.append(f"<h2>{e(s['title'])}</h2><ul>{lis}</ul>")
    qs = [q for q in project["draft"]["questions"] if q.get("answer")]
    if qs:
        body.append("<h2>부록. 전임자 인터뷰 기록</h2><dl>" + "".join(
            f"<dt>Q. {e(q['q'])}</dt><dd>A. {e(q['answer'])}</dd>" for q in qs) + "</dl>")
    body.append("<h2>근거 목록</h2><ol class='refs'>" + "".join(f"<li value='{n}'>{e(sid)} – {e(label)}</li>" for n, sid, label in refs) + "</ol>")
    h = project.get("handover", {})
    sign = (f"<table class='sign'><tr><th>전임자</th><td>{e(project.get('from_name', ''))} {e(h.get('from_signed_at', ''))}</td>"
            f"<th>후임자</th><td>{e(project.get('to_name', ''))} {e(h.get('to_signed_at', ''))}</td></tr></table>")
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>{e(project.get('name', '인수인계서'))}</title>
<style>body{{font-family:'Malgun Gothic','Apple SD Gothic Neo',sans-serif;max-width:860px;margin:32px auto;padding:0 16px;line-height:1.6;color:#111}}
h1{{font-size:24px;border-bottom:3px solid #1f3a8a;padding-bottom:8px}}h2{{font-size:17px;margin-top:28px;color:#1f3a8a}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #bbb;padding:6px 8px;font-size:14px;text-align:left}}th{{background:#eef2fb;width:150px}}
li{{margin:4px 0}}sup{{color:#1f3a8a}}.tag{{font-size:11px;border-radius:4px;padding:1px 5px;background:#eee}}.official{{background:#dcfce7}}.memo{{background:#fef3c7}}
.mail{{background:#dbeafe}}.oral{{background:#ede9fe}}.refs{{font-size:12px;color:#444}}.sign{{margin-top:32px}}dt{{font-weight:bold;margin-top:8px}}
@media print{{body{{margin:0}}}}</style></head><body>
<h1>{e(project.get('title') or '업무 인수인계서')} – {e(project.get('name', ''))}</h1><table>{rows}</table>{''.join(body)}{sign}</body></html>"""


def to_docx(project: dict) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    secs, refs = _collect(project)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(10.5)
    doc.add_heading(f"{project.get('title') or '업무 인수인계서'} – {project.get('name', '')}", 0)
    meta = _meta_rows(project)
    t = doc.add_table(rows=len(meta), cols=2)
    t.style = "Table Grid"
    for r, (k, v) in enumerate(meta):
        t.cell(r, 0).text, t.cell(r, 1).text = k, str(v)
    for s, items in secs:
        doc.add_heading(s["title"], 1)
        if not items:
            doc.add_paragraph("(해당 없음)", style="List Bullet")
        for it, nums in items:
            p = doc.add_paragraph(style="List Bullet")
            tag = p.add_run(f"[{TRUST_LABEL[it['trust']]}] ")
            tag.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
            tag.font.size = Pt(9)
            p.add_run(_item_text(it))
            if nums:
                sup = p.add_run("".join(f"[{n}]" for n in nums))
                sup.font.superscript = True
    qs = [q for q in project["draft"]["questions"] if q.get("answer")]
    if qs:
        doc.add_heading("부록. 전임자 인터뷰 기록", 1)
        for q in qs:
            doc.add_paragraph(f"Q. {q['q']}").runs[0].bold = True
            doc.add_paragraph(f"A. {q['answer']}")
    doc.add_heading("근거 목록", 1)
    for n, sid, label in refs:
        p = doc.add_paragraph(f"[{n}] {sid} – {label}")
        p.runs[0].font.size = Pt(8.5)
    h = project.get("handover", {})
    doc.add_heading("인계·인수 확인", 1)
    t = doc.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    for c, v in enumerate(["구분", "성명", "확인 일시"]):
        t.cell(0, c).text = v
    t.cell(1, 0).text = "전임자 / 후임자"
    t.cell(1, 1).text = f"{project.get('from_name', '')} / {project.get('to_name', '')}"
    t.cell(1, 2).text = f"{h.get('from_signed_at', '')} / {h.get('to_signed_at', '')}"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
