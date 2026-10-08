"""모의데이터 생성기 – 가상 기관 '가상시 스마트정보과' 전임자의 업무 폴더.

모든 인물·기관·연락처는 가상이며 실제 개인정보를 포함하지 않는다.
의도적으로 (1) 공문과 메일의 기한 불일치, (2) 결론 없는 협의, (3) 메모에만 있는 노하우,
(4) 메모 속 비밀번호(민감정보 가림 시험)를 넣어 두었다.

실행: python sample_data/make_samples.py
"""
from __future__ import annotations

import os
import zipfile
from email.message import EmailMessage
from email.utils import format_datetime
import datetime as dt

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "전임자_업무폴더")


def path(*p):
    full = os.path.join(ROOT, *p)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    return full


# ───────────── 1. 사무분장표(엑셀)
def make_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "사무분장"
    ws.append(["담당자", "직급", "담당업무", "비중(%)", "대체자"])
    rows = [
        ["김바통", "주무관", "정보화 예산 편성·집행 및 정보화사업 총괄", 30, "이어달"],
        ["김바통", "주무관", "정보보안 업무(보안점검의 날, 보안교육, 정보보안 관리실태 평가)", 30, "이어달"],
        ["김바통", "주무관", "개인정보보호 관리수준진단 대응 및 개인정보 파일 관리", 20, "이어달"],
        ["김바통", "주무관", "홈페이지 유지관리 용역 감독", 20, "최서버"],
        ["최서버", "주무관", "서버·네트워크 운영, 업무용 PC 보급", 60, "김바통"],
        ["최서버", "주무관", "행정정보시스템(새올) 운영 지원", 40, "김바통"],
    ]
    for r in rows:
        ws.append(r)
    ws2 = wb.create_sheet("용역계약현황")
    ws2.append(["계약명", "업체", "계약기간", "계약금액(원)", "감독관", "업체 담당"])
    ws2.append(["2026년 홈페이지 유지관리 용역", "(주)가상솔루션", "2026. 1. 1. ~ 2026. 12. 31.", 48000000, "김바통 주무관",
                "정유지 PM (02-1234-5678)"])
    ws2.append(["정보보안 관제 서비스", "가상보안(주)", "2026. 3. 1. ~ 2027. 2. 28.", 36000000, "김바통 주무관",
                "한관제 책임 (02-2345-6789)"])
    wb.save(path("01_업무분장", "2026_스마트정보과_사무분장표.xlsx"))


# ───────────── 2. 연간 추진계획(HWPX)
HWPX_PARAS = [
    "2026년 정보보안 업무 추진계획",
    "가상시 스마트정보과",
    "1. 추진 목적",
    "사이버 위협에 대응하고 개인정보 유출 사고를 예방하기 위하여 연간 정보보안 업무를 체계적으로 추진한다.",
    "2. 연간 주요 일정",
    "매월 10일 ‘보안점검의 날’ 운영: 부서별 PC 보안점검 결과를 취합하여 과장 결재 후 온나라로 보고",
    "매년 3월 정보보안 관리실태 자체점검 실시 및 결과 보고",
    "2026. 5. 29.(금)까지 개인정보보호 관리수준진단 실적 자료 제출(행정안전부)",
    "매년 6월 상반기 전 직원 정보보안 교육(집합교육 1회, 온라인 병행)",
    "2027년 정보화 예산 요구서는 6월 12일까지 기획예산과 제출",
    "매년 10월 정보보안 관리실태 평가 수감(국가정보원)",
    "매년 12월 개인정보 파일 보유 현황 정비 및 개인정보 처리방침 갱신",
    "분기별 홈페이지 취약점 점검 실시(용역사 수행, 감독관 확인)",
    "3. 추진 체계",
    "총괄: 스마트정보과장 / 실무: 정보보안 담당 주무관 / 기술지원: (주)가상솔루션, 가상보안(주)",
    "4. 협조 사항",
    "각 부서는 보안점검의 날 점검 결과를 매월 8일까지 스마트정보과로 제출",
]


def make_hwpx():
    def p(t, i):
        t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return (f'<hp:p id="{i}" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
                f'<hp:run charPrIDRef="0"><hp:t>{t}</hp:t></hp:run></hp:p>')

    section = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
               'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
               + "".join(p(t, i) for i, t in enumerate(HWPX_PARAS)) + "</hs:sec>")
    container = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<ocf:container xmlns:ocf="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<ocf:rootfiles><ocf:rootfile full-path="Contents/content.hpf" media-type="application/hwpml-package+xml"/>'
                 '</ocf:rootfiles></ocf:container>')
    hpf = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<opf:package xmlns:opf="http://www.idpf.org/2007/opf/" version="" unique-identifier="" id="">'
           '<opf:metadata><opf:title>2026년 정보보안 업무 추진계획</opf:title></opf:metadata>'
           '<opf:manifest><opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/></opf:manifest>'
           '<opf:spine><opf:itemref idref="section0" linear="yes"/></opf:spine></opf:package>')
    with zipfile.ZipFile(path("02_계획", "2026년_정보보안_업무추진계획.hwpx"), "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip")
        z.writestr("META-INF/container.xml", container, zipfile.ZIP_DEFLATED)
        z.writestr("Contents/content.hpf", hpf, zipfile.ZIP_DEFLATED)
        z.writestr("Contents/section0.xml", section, zipfile.ZIP_DEFLATED)
        z.writestr("Preview/PrvText.txt", "\n".join(HWPX_PARAS), zipfile.ZIP_DEFLATED)


# ───────────── 3. 공문(PDF)
def make_pdf():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    lines = [
        ("가 상 도", 18),
        ("", 11),
        ("수신  가상시장(스마트정보과장)", 11),
        ("(경유)", 11),
        ("제목  2026년 개인정보보호 관리수준진단 실적 제출 안내", 11),
        ("", 11),
        ("1. 행정안전부 개인정보보호정책과-1234(2026. 4. 20.)호와 관련입니다.", 11),
        ("2. 2026년 개인정보보호 관리수준진단을 아래와 같이 실시하오니", 11),
        ("   기한 내 실적 자료를 제출하여 주시기 바랍니다.", 11),
        ("   가. 진단 대상: 시·군 본청 및 직속기관", 11),
        ("   나. 제출 기한: 2026. 5. 29.(금)까지", 11),
        ("   다. 제출 방법: 개인정보보호 종합지원시스템 등록 및 증빙자료 첨부", 11),
        ("   라. 주요 변경: 개인정보 파일 등록 정확성 지표 신설(배점 10점)", 11),
        ("3. 증빙자료 누락 시 감점되므로 유의하시기 바랍니다.", 11),
        ("", 11),
        ("붙임  2026년 관리수준진단 지표 및 증빙 목록 1부.  끝.", 11),
        ("", 11),
        ("가상도지사", 16),
        ("", 11),
        ("주무관 최보안   정보보안팀장 전결 2026. 4. 24.", 10),
        ("시행  정보보안팀-567 (2026. 4. 24.)", 10),
        ("가상도 정보보안팀 / 전화 055-123-4567 / 공개구분 공개", 10),
    ]
    c = canvas.Canvas(path("03_공문", "[공문]2026년_개인정보보호_관리수준진단_실적제출_안내.pdf"))
    y = 790
    for t, size in lines:
        c.setFont("HYGothic-Medium", size)
        c.drawString(60, y, t)
        y -= size + 12
    c.save()


# ───────────── 4. 회의록(DOCX)
def make_docx():
    from docx import Document

    d = Document()
    d.add_heading("2026년 1분기 정보화추진협의회 회의록", 1)
    for t in [
        "일시: 2026. 3. 18.(수) 14:00 / 장소: 본관 3층 소회의실",
        "참석: 스마트정보과장, 김바통 주무관, 최서버 주무관, 기획예산과 박예산 주무관, (주)가상솔루션 정유지 PM",
        "1. 노후 서버 교체: 2018년 도입한 홈페이지 WEB 서버의 유지보수 기간이 만료됨. 교체 필요성은 공감하였으나 올해 예산이 부족하여 보류, 2027년 예산 요구 시 반영 검토 중.",
        "2. 홈페이지 웹 접근성 개선: 접근성 인증 만료(2026. 9. 30.) 전 개선 필요. (주)가상솔루션에서 견적 제출 예정이며 현재 회신 대기 중.",
        "3. 업무용 PC 교체: 수요조사 결과 42대 교체 필요. 최서버 주무관이 조달 구매 진행 중.",
        "4. 후속조치: 2027년 정보화 예산 요구 시 서버 교체비 포함 여부를 6월 초까지 결정 필요.",
    ]:
        d.add_paragraph(t)
    d.save(path("05_회의록", "2026_1분기_정보화추진협의회_회의록.docx"))


# ───────────── 5. 메일(EML)
ME = "김바통 주무관 <kim.baton@gasang.go.kr>"


def eml(fname, frm, to, subject, when, body, cc=None):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = frm, to, subject
    if cc:
        m["Cc"] = cc
    m["Date"] = format_datetime(when)
    m.set_content(body)
    with open(path("04_메일", fname), "wb") as f:
        f.write(bytes(m))


def make_mails():
    kst = dt.timezone(dt.timedelta(hours=9))
    eml("01_2027_정보화예산_요구일정.eml", "박예산 주무관 <park.budget@gasang.go.kr>", ME,
        "2027년도 정보화 예산 요구 일정 안내", dt.datetime(2026, 4, 28, 10, 12, tzinfo=kst),
        "김바통 주무관님, 기획예산과 박예산입니다.\n\n"
        "2027년도 정보화사업 예산 요구서는 6월 12일까지 제출 부탁드립니다.\n"
        "올해부터 정보화사업은 사전 협의(5월 중) 후 요구서를 받습니다. 협의 일정은 따로 연락드릴게요.\n"
        "작년 양식이 아닌 첨부한 2027년 양식을 꼭 사용해 주세요. 작년 양식으로 내시면 반려됩니다.\n\n"
        "박예산 드림 (기획예산과, 내선 2231)")
    eml("02_수준진단_도_취합일정.eml", "최보안 주무관 <choi.sec@gasang-do.go.kr>", ME,
        "RE: 개인정보 관리수준진단 실적 제출 관련 문의", dt.datetime(2026, 4, 30, 16, 40, tzinfo=kst),
        "안녕하세요, 가상도 정보보안팀 최보안입니다.\n\n"
        "문의하신 건 답변드립니다. 시·군 실적은 도에서 먼저 검토 후 행정안전부에 올리기 때문에\n"
        "증빙자료는 5월 22일까지 도로 먼저 보내주셔야 합니다.\n"
        "올해 신설된 개인정보 파일 등록 정확성 지표는 감점이 크니 파일 목록을 꼭 재점검해 주세요.\n\n"
        "-----Original Message-----\n"
        "수준진단 제출은 공문 기한(5. 29.)에 맞추면 되는지요?")
    eml("03_홈페이지_접근성_견적.eml", "정유지 PM <jung.pm@gasang-sol.co.kr>", ME,
        "홈페이지 웹 접근성 개선 견적 관련", dt.datetime(2026, 5, 6, 9, 5, tzinfo=kst),
        "주무관님 안녕하세요, (주)가상솔루션 정유지입니다.\n\n"
        "웹 접근성 개선 견적은 내부 검토 중이며 다음 주 중 회신드리겠습니다.\n"
        "개선 범위(민원 신청 화면 포함 여부)는 아직 결정이 필요합니다. 범위가 확정되어야 정확한 견적이 가능합니다.\n"
        "장애 발생 시에는 유지관리 콜센터(02-1234-5678)로 먼저 연락 주세요.\n\n"
        "정유지 PM 드림", cc="최서버 주무관 <choi.server@gasang.go.kr>")
    eml("04_보안교육_강사섭외.eml", ME, "한관제 책임 <han.soc@gasangsec.co.kr>",
        "상반기 정보보안 교육 강사 섭외 요청", dt.datetime(2026, 5, 11, 13, 30, tzinfo=kst),
        "한관제 책임님, 가상시 스마트정보과 김바통입니다.\n\n"
        "6월 중 전 직원 정보보안 교육(약 1시간, 집합)을 계획하고 있어 강사 지원 가능 여부를 여쭙니다.\n"
        "작년처럼 피싱메일 모의훈련 결과 공유도 포함되면 좋겠습니다. 가능한 일정 회신 부탁드립니다.\n\n"
        "김바통 드림")


# ───────────── 6. 개인 메모(TXT)
MEMO = """[업무 메모 – 김바통, 수시로 적어둔 것]

* 예산
- 예산요구서는 기획예산과 박예산 주무관이 매년 새 양식을 줌. 작년 양식 쓰면 반려되니 꼭 새 양식 받을 것!
- 서버 교체비(약 3천만원) 2027년 예산에 넣을지 과장님과 아직 결론 못 냄.

* 보안점검의 날
- 매월 10일 결과 보고. 부서 제출이 늦으니 7일쯤 메신저로 미리 독촉하는 게 요령.
- 과장님은 오전 결재를 선호하심. 오후 늦게 올리면 다음날로 밀림.

* 수준진단
- 5/22까지 도 제출. 증빙은 공유폴더 > 정보보안 > 2026_수준진단 폴더에 모아둠.
- 개인정보 파일 목록은 부서마다 이름을 다르게 적어서 반드시 대조표로 확인해야 함.

* 시스템
- 보안관제 포털 계정 PW: Gasang!2026 (후임자에게 구두로 전달하고 바로 변경할 것)
- 온나라 문서 분류는 '정보보안 > 보안점검'으로 해야 나중에 찾기 쉬움.

* 사람
- 최서버 주무관: 서버·PC 담당, 장애 나면 제일 먼저 상의.
- 도 정보보안팀 최보안 주무관은 메일보다 전화(055-123-4567)를 선호.
"""


def make_memo():
    with open(path("99_개인메모", "업무메모_김바통.txt"), "w", encoding="utf-8") as f:
        f.write(MEMO)


# ───────────── 7. 이전 담당자의 바통 파일(지식 릴레이 시연용)
def make_relay():
    import json

    data = {
        "format": "baton/1",
        "meta": {"name": "정보보안 담당 인수인계", "org": "가상시", "dept": "스마트정보과",
                 "from_name": "박선배 주무관", "to_name": "김바통 주무관", "base_date": "2024-07-01"},
        "lineage": [{"name": "박선배 주무관", "handed_to": "김바통 주무관", "date": "2024-07-01", "work": "정보보안 담당"}],
        "sections": [
            {"kind": "calendar", "title": "시기별 할 일", "items": [
                {"text": "[매년 1월] 보안관제 서비스 계약은 매년 2월 말 만료되므로 1월 초에 다음 연도 계약 준비를 시작해야 함", "trust": "official"},
                {"text": "[매년 10월] 정보보안 관리실태 평가 수감 전, 전년도 지적사항 조치 결과를 9월까지 정리해 둘 것", "trust": "oral"}]},
            {"kind": "tips", "title": "노하우·주의사항", "items": [
                {"text": "국정원 관리실태 평가 때 가장 많이 지적되는 것은 퇴직자 계정 미삭제 – 인사발령 공문 볼 때마다 계정 삭제 요청할 것", "trust": "oral"},
                {"text": "보안교육 출석부는 교육 당일 받아야 함. 나중에 받으면 누락자가 생김", "trust": "oral"}]},
        ],
        "interview": [{"q": "후임자가 가장 실수하기 쉬운 부분은?", "a": "관리실태 평가 증빙을 연말에 몰아서 챙기는 것. 매달 보안점검의 날 결과를 증빙 폴더에 바로 넣어 두면 10월이 편합니다."}],
        "exported_at": "2024-06-28 17:40",
    }
    with open(path("00_이전인수인계", "정보보안_박선배_2024.baton"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    make_relay()
    make_xlsx()
    make_hwpx()
    make_pdf()
    make_docx()
    make_mails()
    make_memo()
    print("모의데이터 생성 완료:", ROOT)
