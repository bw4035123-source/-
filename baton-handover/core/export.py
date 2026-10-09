"""산출물 내보내기: 한글(.hwpx), 워드(.docx), 마크다운(.md).

앱은 아래 형식의 '블록 목록' 하나만 만들면 세 형식으로 모두 저장할 수 있다.
  ("title", "문서 제목")
  ("h1"|"h2"|"h3", "제목")
  ("p", "문단")            # __밑줄__, **굵게** 표기 지원
  ("bullet", "항목")
  ("table", header_list, rows_list[, col_widths])
  ("note", "※ 참고")
"""
import io
import re
import warnings

RICH = re.compile(r"(__.+?__|\*\*.+?\*\*)")


def _segments(text):
    """'__밑줄__', '**굵게**' 표기를 (텍스트, 굵게, 밑줄) 조각으로 나눈다."""
    out = []
    for part in RICH.split(text or ""):
        if not part:
            continue
        if part.startswith("__") and part.endswith("__"):
            out.append((part[2:-2], False, True))
        elif part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        else:
            out.append((part, False, False))
    return out


def plain(text):
    return "".join(s for s, _, _ in _segments(text))


# ---------------------------------------------------------------- Markdown

def to_markdown(blocks):
    lines = []
    for b in blocks:
        kind = b[0]
        if kind == "title":
            lines += [f"# {b[1]}", ""]
        elif kind in ("h1", "h2", "h3"):
            lines += ["#" * (int(kind[1]) + 1) + " " + b[1], ""]
        elif kind == "p":
            lines += [b[1], ""]
        elif kind == "bullet":
            lines.append(f"- {b[1]}")
        elif kind == "note":
            lines += [f"> {b[1]}", ""]
        elif kind == "table":
            header, rows = b[1], b[2]
            esc = lambda s: str(s).replace("|", "\\|").replace("\n", "<br>")
            lines.append("| " + " | ".join(esc(h) for h in header) + " |")
            lines.append("|" + "---|" * len(header))
            for r in rows:
                lines.append("| " + " | ".join(esc(c) for c in r) + " |")
            lines.append("")
    return "\n".join(lines).encode("utf-8")


# ---------------------------------------------------------------- HWPX

def to_hwpx(blocks):
    from hwpx import HwpxDocument
    warnings.filterwarnings("ignore", category=DeprecationWarning)
    doc = HwpxDocument.new()
    styles = {}

    def style(bold=False, underline=False, size=None):
        key = (bold, underline, size)
        if key not in styles:
            styles[key] = doc.ensure_run_style(bold=bold, underline=underline, size=size)
        return styles[key]

    def rich_para(para, text, size=None, bold_all=False):
        for seg, bold, und in _segments(text):
            para.add_run(seg, char_pr_id_ref=style(bold or bold_all, und, size))

    first = True
    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = doc.add_paragraph("", include_run=False) if not first else doc.paragraphs[0]
            rich_para(p, b[1], size=16, bold_all=True)
        elif kind in ("h1", "h2", "h3"):
            size = {"h1": 14, "h2": 12, "h3": 11}[kind]
            p = doc.add_paragraph("", include_run=False)
            rich_para(p, b[1], size=size, bold_all=True)
        elif kind in ("p", "note", "bullet"):
            text = b[1] if kind != "bullet" else "• " + b[1]
            for line in str(text).split("\n"):
                p = doc.add_paragraph("", include_run=False)
                rich_para(p, line)
        elif kind == "table":
            header, rows = b[1], b[2]
            ncol = max([len(header)] + [len(r) for r in rows])
            t = doc.add_table(len(rows) + 1, ncol)
            for c in range(ncol):
                t.set_cell_text(0, c, str(header[c]) if c < len(header) else "")
                t.set_cell_shading(0, c, "EAF1FB")
            for ri, r in enumerate(rows, 1):
                for c in range(ncol):
                    val = str(r[c]) if c < len(r) and r[c] is not None else ""
                    if RICH.search(val):
                        cell = t.cell(ri, c)
                        cell.set_text("")
                        lines = val.split("\n")
                        para = cell.paragraphs[0]
                        for li, line in enumerate(lines):
                            if li:
                                para = cell.add_paragraph("")
                            for seg, bold, und in _segments(line):
                                para.add_run(seg, char_pr_id_ref=style(bold, und))
                    else:
                        t.set_cell_text(ri, c, val)
            if len(b) > 3 and b[3]:
                try:
                    total = sum(b[3])
                    width = 42520  # A4 본문 폭(HWPUNIT)
                    t.set_column_widths([int(width * w / total) for w in b[3]])
                except Exception:
                    pass
            doc.add_paragraph("")
        first = False
    return doc.to_bytes()


# ---------------------------------------------------------------- DOCX

def to_docx(blocks):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(10.5)

    def rich(par, text):
        for seg, bold, und in _segments(text):
            r = par.add_run(seg)
            r.bold = bold or None
            r.underline = und or None

    for b in blocks:
        kind = b[0]
        if kind == "title":
            d.add_heading(b[1], level=0)
        elif kind in ("h1", "h2", "h3"):
            d.add_heading(b[1], level=int(kind[1]))
        elif kind in ("p", "note"):
            for line in str(b[1]).split("\n"):
                rich(d.add_paragraph(), line)
        elif kind == "bullet":
            rich(d.add_paragraph(style="List Bullet"), b[1])
        elif kind == "table":
            header, rows = b[1], b[2]
            ncol = max([len(header)] + [len(r) for r in rows])
            t = d.add_table(rows=1, cols=ncol)
            t.style = "Table Grid"
            for c in range(ncol):
                t.rows[0].cells[c].text = str(header[c]) if c < len(header) else ""
            for r in rows:
                cells = t.add_row().cells
                for c in range(ncol):
                    val = str(r[c]) if c < len(r) and r[c] is not None else ""
                    cells[c].text = ""
                    lines = val.split("\n")
                    par = cells[c].paragraphs[0]
                    for li, line in enumerate(lines):
                        if li:
                            par = cells[c].add_paragraph()
                        rich(par, line)
            d.add_paragraph()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


FORMATS = {
    "hwpx": (to_hwpx, "application/hwp+zip"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "md": (to_markdown, "text/markdown; charset=utf-8"),
}


def export(blocks, fmt):
    if fmt not in FORMATS:
        from .server import HttpError
        raise HttpError(400, f"지원하지 않는 파일 형식입니다: '{fmt}'. 한글(hwpx)·워드(docx)·마크다운(md) 중에서 골라 주세요.")
    fn, ctype = FORMATS[fmt]
    return fn(blocks), ctype
