"""여러 형식의 문서를 읽어 '블록(위치가 붙은 문장 단위)'으로 바꾼다.

지원 형식: 한글(.hwpx, .hwp), PDF, 엑셀(.xlsx/.xlsm), 워드(.docx), 파워포인트(.pptx), 메일(.eml), 텍스트(.txt/.md/.csv)
원본 파일은 읽기만 하고 절대 수정·이동·삭제하지 않는다.
"""
import csv
import email
import email.policy
import hashlib
import io
import json
import re
import struct
import zipfile
import zlib
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field, asdict
from pathlib import Path

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
SUPPORTED = {".hwpx", ".hwp", ".pdf", ".xlsx", ".xlsm", ".docx", ".pptx", ".eml", ".txt", ".md", ".csv", ".baton", ".zip"} | IMAGE_EXT

# 사진·스캔본 글자 읽기(core/ocr.py 의 make()가 돌려주는 함수). 앱이 시작할 때 넣는다. 없으면 사진은 '글자 못 읽음'으로 남는다
OCR = None
MAX_BLOCKS = 3000        # 한 파일에서 읽는 줄(문단·표 행) 수 상한
OCR_LIMIT = 200          # 한 번에 글자를 읽을 사진·스캔 쪽 수 상한
ZIP_MAX_BYTES = 300 * 1024 * 1024
ZIP_MAX_FILES = 3000
JUNK = {"thumbs.db", "desktop.ini", ".ds_store"}
OLD_FORMATS = {
    ".ppt": "파워포인트 97~2003 형식(.ppt)은 읽지 못합니다. PowerPoint에서 .pptx로 저장해 넣어 주세요",
    ".xls": "엑셀 97~2003 형식(.xls)은 읽지 못합니다. 엑셀에서 .xlsx로 저장해 넣어 주세요",
    ".doc": "워드 97~2003 형식(.doc)은 읽지 못합니다. 워드에서 .docx로 저장해 넣어 주세요",
    ".heic": "아이폰 사진(.heic)은 읽지 못합니다. JPG로 바꿔 넣어 주세요(사진 앱에서 '호환성 우선' 설정)",
    ".hwt": "한글 서식 파일(.hwt)은 내용이 없는 틀이라 건너뜁니다",
    ".7z": "7z 압축은 풀지 못합니다. ZIP으로 압축해 넣어 주세요",
    ".rar": "RAR 압축은 풀지 못합니다. ZIP으로 압축해 넣어 주세요",
    ".alz": "알집(.alz) 압축은 풀지 못합니다. ZIP으로 압축해 넣어 주세요",
    ".egg": "알집(.egg) 압축은 풀지 못합니다. ZIP으로 압축해 넣어 주세요",
}


def skip_reason(ext):
    return OLD_FORMATS.get(ext.lower(), "지원하지 않는 형식")


@dataclass
class Block:
    """원문의 한 조각(문단·표 행·페이지 줄). 출처 표시의 최소 단위."""
    id: str
    doc_id: str
    file: str
    loc: str
    text: str

    def cite(self, quote=None):
        q = quote if quote is not None else self.text
        if len(q) > 160:
            q = q[:157] + "..."
        return {"block_id": self.id, "file": self.file, "loc": self.loc, "quote": q}


@dataclass
class Document:
    id: str
    name: str
    relpath: str
    ext: str
    blocks: list = field(default_factory=list)
    tables: list = field(default_factory=list)  # [{"name":..., "header":[...], "rows":[[...]], "loc_prefix":...}]
    meta: dict = field(default_factory=dict)
    error: str = ""

    @property
    def text(self):
        return "\n".join(b.text for b in self.blocks)

    def to_dict(self, with_blocks=False):
        d = {
            "id": self.id, "name": self.name, "relpath": self.relpath, "ext": self.ext,
            "meta": self.meta, "error": self.error, "n_blocks": len(self.blocks),
            "chars": sum(len(b.text) for b in self.blocks),
        }
        if with_blocks:
            d["blocks"] = [asdict(b) for b in self.blocks]
        return d


def _doc_id(relpath):
    return "d" + hashlib.md5(relpath.encode("utf-8")).hexdigest()[:8]


def _decode(data):
    for enc in ("utf-8-sig", "cp949", "euc-kr", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _clean(s):
    s = s.replace(" ", " ").replace("\t", " ")
    s = re.sub(r"[ 　]+", " ", s)
    return s.strip()


def _local(tag):
    return tag.rsplit("}", 1)[-1]


# ---------------------------------------------------------------- 형식별 읽기

def _read_text(data, name):
    out = []
    for i, line in enumerate(_decode(data).splitlines(), 1):
        t = _clean(line)
        if t:
            out.append((f"{i}행", t))
    return out, []


def _read_csv(data, name):
    text = _decode(data)
    rows = [r for r in csv.reader(io.StringIO(text))]
    rows = [[_clean(c) for c in r] for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    header = rows[0]
    out = []
    for i, r in enumerate(rows[1:], 2):
        out.append((f"{i}행", _row_text(header, r)))
    return out, [{"name": Path(name).stem, "header": header, "rows": rows[1:], "loc_prefix": "",
                  "row_locs": [f"{i}행" for i in range(2, len(rows) + 1)]}]


def _row_text(header, row):
    parts = []
    for i, v in enumerate(row):
        if v in (None, ""):
            continue
        h = header[i] if i < len(header) and header[i] else ""
        parts.append(f"{h}: {v}" if h else str(v))
    return " / ".join(parts)


def _xml_paragraphs(root, p_tag, t_tag, tbl_tag, tr_tag, tc_tag):
    """문서 순서대로 (종류, 텍스트) 를 낸다. 표는 행 단위로 '셀 | 셀' 로 합친다."""
    def para_text(p):
        return "".join(t.text or "" for t in p.iter() if _local(t.tag) == t_tag)

    def walk(node):
        for child in node:
            tag = _local(child.tag)
            if tag == tbl_tag:
                rows = []
                for tr in child.iter():
                    if _local(tr.tag) != tr_tag:
                        continue
                    cells = []
                    for tc in tr:
                        if _local(tc.tag) != tc_tag:
                            continue
                        cells.append(" / ".join(t for t in (
                            _clean(para_text(p)) for p in tc.iter() if _local(p.tag) == p_tag) if t))
                    rows.append(cells)
                yield ("table", rows)
            elif tag == p_tag:
                # 문단 안에 표가 들어있는 경우(한글) 표를 먼저 처리
                inner_tbls = top_tables(child)
                if inner_tbls:
                    own = []
                    for run in child:
                        for t in run:
                            if _local(t.tag) == t_tag and t.text:
                                own.append(t.text)
                    if _clean("".join(own)):
                        yield ("p", _clean("".join(own)))
                    for tbl in inner_tbls:
                        yield from walk_tbl(tbl)
                else:
                    t = _clean(para_text(child))
                    if t:
                        yield ("p", t)
            else:
                yield from walk(child)

    def top_tables(node):
        found = []
        for c in node:
            if _local(c.tag) == tbl_tag:
                found.append(c)
            else:
                found.extend(top_tables(c))
        return found

    def walk_tbl(tbl):
        wrapper = ET.Element("w")
        wrapper.append(tbl)
        yield from walk(wrapper)

    yield from walk(root)


def _emit_xml(items, tables, tname):
    out = []
    pi = ti = 0
    for kind, val in items:
        if kind == "p":
            pi += 1
            out.append((f"문단 {pi}", val))
        else:
            ti += 1
            rows = [r for r in val if any(r)]
            if not rows:
                continue
            header = rows[0]
            tables.append({"name": f"{tname} 표{ti}", "header": header, "rows": rows[1:],
                           "loc_prefix": f"표{ti} ",
                           "row_locs": [f"표{ti} {ri}행" for ri in range(2, len(rows) + 1)]})
            for ri, r in enumerate(rows, 1):
                txt = " | ".join(c for c in r if c)
                if txt:
                    out.append((f"표{ti} {ri}행", txt))
    return out


def _read_hwpx(data, name):
    z = zipfile.ZipFile(io.BytesIO(data))
    secs = sorted([n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n)],
                  key=lambda n: int(re.findall(r"\d+", n)[-1]))
    out, tables = [], []
    for s in secs:
        root = ET.fromstring(z.read(s))
        items = list(_xml_paragraphs(root, "p", "t", "tbl", "tr", "tc"))
        out.extend(_emit_xml(items, tables, Path(name).stem))
    return out, tables


def _read_docx(data, name):
    z = zipfile.ZipFile(io.BytesIO(data))
    root = ET.fromstring(z.read("word/document.xml"))
    tables = []
    items = list(_xml_paragraphs(root, "p", "t", "tbl", "tr", "tc"))
    return _emit_xml(items, tables, Path(name).stem), tables


def _pptx_slides(z):
    """발표 순서대로 슬라이드 파일 경로를 낸다(presentation.xml의 sldIdLst 순서)."""
    names = set(z.namelist())
    try:
        rels = ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
        target = {r.get("Id"): "ppt/" + r.get("Target").lstrip("/").removeprefix("ppt/") for r in rels}
        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        order = [target.get(e.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
                 for e in pres.iter() if _local(e.tag) == "sldId"]
        order = [o for o in order if o in names]
        if order:
            return order
    except (KeyError, ET.ParseError):
        pass
    return sorted((n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                  key=lambda n: int(re.findall(r"\d+", n)[-1]))


def _read_pptx(data, name):
    z = zipfile.ZipFile(io.BytesIO(data))
    names = set(z.namelist())
    out, tables = [], []
    for si, path in enumerate(_pptx_slides(z), 1):
        sub = []
        items = list(_xml_paragraphs(ET.fromstring(z.read(path)), "p", "t", "tbl", "tr", "tc"))
        for loc, text in _emit_xml(items, sub, f"{Path(name).stem} 슬라이드{si}"):
            out.append((f"슬라이드 {si} {loc}", text))
        for t in sub:
            t["loc_prefix"] = f"슬라이드 {si} " + t["loc_prefix"]
            t["row_locs"] = [f"슬라이드 {si} " + x for x in t["row_locs"]]
        tables.extend(sub)
        # 발표자 메모: 슬라이드 rels에서 notesSlide를 찾는다
        rel = path.replace("slides/", "slides/_rels/") + ".rels"
        if rel in names:
            for r in ET.fromstring(z.read(rel)):
                if r.get("Type", "").endswith("/notesSlide"):
                    npath = "ppt/" + r.get("Target").replace("../", "")
                    if npath in names:
                        ni = 0
                        for e in ET.fromstring(z.read(npath)).iter():
                            if _local(e.tag) == "p":
                                t = _clean("".join(x.text or "" for x in e.iter() if _local(x.tag) == "t"))
                                if t and not t.isdigit():  # 쪽 번호 자리 제외
                                    ni += 1
                                    out.append((f"슬라이드 {si} 메모 {ni}", t))
    return out, tables


# HWP 5.0 바이너리: 레코드 구조에서 문단 텍스트(HWPTAG_PARA_TEXT=67)만 꺼낸다.
_HWP_CHAR_CTRL = {0, 10, 13, 24, 25, 26, 27, 28, 29, 30, 31}


def _hwp_para_text(payload):
    chars = []
    i = 0
    n = len(payload) // 2
    while i < n:
        c = struct.unpack_from("<H", payload, i * 2)[0]
        if c < 32:
            if c in _HWP_CHAR_CTRL:
                if c in (10, 13):
                    chars.append("\n")
                i += 1
            else:
                i += 8  # 인라인/확장 컨트롤은 8 WCHAR 차지
                if c == 9:
                    chars.append(" ")
            continue
        chars.append(chr(c))
        i += 1
    return "".join(chars)


def _read_hwp(data, name):
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(data))
    header = ole.openstream("FileHeader").read()
    flags = struct.unpack_from("<I", header, 36)[0]
    compressed = bool(flags & 1)
    if flags & 2:
        raise ValueError("암호가 걸린 한글 파일은 읽을 수 없습니다")
    secs = sorted([e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"],
                  key=lambda e: int(re.findall(r"\d+", e[1])[0]))
    out = []
    pi = 0
    for e in secs:
        raw = ole.openstream(e).read()
        if compressed:
            raw = zlib.decompress(raw, -15)
        pos = 0
        while pos + 4 <= len(raw):
            h = struct.unpack_from("<I", raw, pos)[0]
            tag = h & 0x3FF
            size = (h >> 20) & 0xFFF
            pos += 4
            if size == 0xFFF:
                size = struct.unpack_from("<I", raw, pos)[0]
                pos += 4
            if tag == 67:
                for line in _hwp_para_text(raw[pos:pos + size]).split("\n"):
                    t = _clean(line)
                    if t:
                        pi += 1
                        out.append((f"문단 {pi}", t))
            pos += size
    return out, []


def _read_pdf(data, name):
    """PDF 글자. 글자가 없는 쪽(스캔본)은 쪽 안의 그림을 글자 인식(OCR)으로 읽는다."""
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    out, scanned, engines, errors = [], 0, set(), []
    for pno, page in enumerate(r.pages, 1):
        text = page.extract_text() or ""
        if not text.strip():
            try:
                imgs = list(page.images)
            except Exception:
                imgs = []
            if imgs:
                scanned += 1
                for img in imgs[:3]:
                    res = _ocr(img.data)
                    if res["engine"]:
                        engines.add(res["engine"])
                    elif res["error"]:
                        errors.append(res["error"])
                    for li, line in enumerate(res["text"].splitlines(), 1):
                        t = _clean(line)
                        if t:
                            out.append((f"p.{pno} 스캔 {li}줄", t))
                continue
        for li, line in enumerate(text.splitlines(), 1):
            t = _clean(line)
            if t:
                out.append((f"p.{pno} {li}줄", t))
    meta = {}
    if scanned:
        meta = {"scanned_pages": scanned, "ocr_engine": ", ".join(sorted(engines)),
                "ocr_error": "" if engines else (errors[0] if errors else "스캔본 글자를 읽지 못했습니다")}
    return out, [], meta


_ocr_budget = [OCR_LIMIT]


def _ocr(data):
    if OCR is None:
        return {"text": "", "engine": "", "error": "글자 인식(OCR) 기능이 꺼져 있습니다"}
    if _ocr_budget[0] <= 0:
        return {"text": "", "engine": "", "error": f"사진·스캔 쪽이 {OCR_LIMIT}개를 넘어 나머지는 글자를 읽지 않았습니다"}
    _ocr_budget[0] -= 1
    return OCR(data)


def _read_image(data, name):
    """사진(현장 사진, 화이트보드, 메모지, 휴대폰으로 찍은 공문 등): 글자 인식 결과를 줄 단위로."""
    res = _ocr(data)
    items = [(f"사진 {i}줄", t) for i, t in enumerate((_clean(x) for x in res["text"].splitlines()), 1) if t]
    meta = {"image": True, "ocr_engine": res.get("engine", ""), "ocr_error": "" if items else (res.get("error") or "사진에서 글자를 찾지 못했습니다"),
            "taken": res.get("taken"), "size": f"{res.get('width')}x{res.get('height')}" if res.get("width") else ""}
    return items, [], meta


def _read_xlsx(data, name):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, tables = [], []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if v is None else _clean(str(v)) for v in r]
            while cells and cells[-1] == "":
                cells.pop()
            rows.append(cells)
        # 앞쪽 빈 행 제외, 첫 비어있지 않은 행을 머리행으로 본다
        idx = [i for i, r in enumerate(rows) if any(r)]
        if not idx:
            continue
        # 제목 행(셀 1개)은 건너뛰고, 값이 충분히 찬 첫 행을 머리행으로 본다
        widest = max(sum(1 for c in rows[i] if c) for i in idx)
        hi = next(i for i in idx if sum(1 for c in rows[i] if c) >= max(2, widest // 2))
        for i in idx:
            if i < hi:
                out.append((f"{ws.title}!{i + 1}행", " ".join(c for c in rows[i] if c)))
        header = rows[hi]
        body = [(i + 1, rows[i]) for i in idx if i > hi]
        tables.append({"name": ws.title, "header": header, "rows": [r for _, r in body],
                       "loc_prefix": f"{ws.title}!", "row_locs": [f"{ws.title}!{rn}행" for rn, _ in body]})
        out.append((f"{ws.title}!{hi + 1}행", " | ".join(c for c in header if c)))
        for rn, r in body:
            out.append((f"{ws.title}!{rn}행", _row_text(header, r)))
    return out, tables


def _strip_html(s):
    s = re.sub(r"(?is)<(script|style).*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>", "\n", s)
    s = re.sub(r"<[^>]+>", " ", s)
    return s


def _read_eml(data, name):
    msg = email.message_from_bytes(data, policy=email.policy.default)
    meta = {k: str(msg.get(k, "")) for k in ("From", "To", "Cc", "Date", "Subject")}
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = _strip_html(text)
    attachments = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    meta["attachments"] = attachments
    out = []
    for k in ("Subject", "From", "To", "Date"):
        if meta.get(k):
            out.append((f"머리글 {k}", f"{k}: {meta[k]}"))
    for i, line in enumerate(text.splitlines(), 1):
        t = _clean(line)
        if t and not t.startswith(">"):
            out.append((f"본문 {i}행", t))
    return out, [], meta


# ---------------------------------------------------------------- 바통 파일(.baton): 지식 릴레이

BATON_FORMAT = "baton-handover/1"
BATON_SECTIONS = {"rnr": "담당업무", "schedule": "일정", "contacts": "연락처", "issues": "현안"}
_BATON_FIELDS = {
    "rnr": {"duty": str, "detail": str},
    "schedule": {"task": str, "related": str, "action": str, "recurring": str, "when": str, "date": str,
                 "months": list, "day": int, "year": int, "approx": bool},
    "contacts": {"name": str, "title": str, "org": str, "phones": list, "emails": list, "topics": list},
    "issues": {"title": str, "state": str, "next_action": str, "due": str, "topic": str},
}
_BATON_MAIN = {"rnr": "duty", "schedule": "task", "contacts": "name", "issues": "title"}
# 손상·악의적 파일로 앱이 느려지지 않게 하는 상한
BATON_MAX_BYTES = 5 * 1024 * 1024
BATON_MAX_LINEAGE = 50
BATON_MAX_ITEMS = 1000
BATON_MAX_TEXT = 2000


def _baton_item(sec, raw):
    """바통 파일의 항목을 정해진 필드·형식만 남겨 읽는다(손으로 고친 파일도 앱이 깨지지 않게)."""
    if not isinstance(raw, dict):
        return None
    out = {}
    for k, typ in _BATON_FIELDS[sec].items():
        v = raw.get(k)
        if typ is str:
            out[k] = str(v).strip()[:BATON_MAX_TEXT] if v not in (None, "") else ""
        elif typ is list:
            out[k] = [str(x).strip()[:BATON_MAX_TEXT] for x in (v if isinstance(v, list) else [])[:50] if str(x).strip()]
        elif typ is int:
            out[k] = int(v) if isinstance(v, (int, str)) and str(v).isdigit() else None
        else:
            out[k] = bool(v)
    if sec == "schedule":
        out["months"] = sorted({int(m) for m in out["months"] if m.isdigit() and 1 <= int(m) <= 12})
        if out["day"] is not None and not 1 <= out["day"] <= 31:
            out["day"] = None
        if out["year"] is not None and not 2000 <= out["year"] <= 2100:
            out["year"] = None
        if not out["months"]:
            return None
        out["date"] = out["date"] if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", out["date"]) else None
    if not out[_BATON_MAIN[sec]]:
        return None
    return out


def _baton_text(sec, it):
    if sec == "rnr":
        return f"[담당업무] {it['duty']}" + (f" — {it['detail']}" if it["detail"] else "")
    if sec == "schedule":
        when = "매월" if it["recurring"] == "매월" else ", ".join(f"{m}월" for m in it["months"])
        when += f" {it['day']}일" if it["day"] else ""
        return f"[일정] {when}{' (' + it['recurring'] + ')' if it['recurring'] and it['recurring'] != '매월' else ''}: {it['task']}"
    if sec == "contacts":
        return (f"[연락처] {it['name']} {it['title']} ({it['org'] or '-'}) {', '.join(it['phones'] + it['emails']) or '연락처 없음'}"
                + (f" — {', '.join(it['topics'][:2])}" if it["topics"] else ""))
    return (f"[현안] {it['title']} [{it['state'] or '-'}] 다음 할 일: {it['next_action'] or '-'} / 기한: {it['due'] or '-'}")


def _read_baton(data, name):
    """이전 담당자가 저장한 인수인계서(.baton)를 블록으로 읽고, 정리된 항목·계보·질의응답은 meta로 넘긴다."""
    if len(data) > BATON_MAX_BYTES:
        raise ValueError(f"바통 파일이 너무 큽니다({len(data) // 1024 // 1024}MB, 최대 5MB)")
    obj = json.loads(_decode(data))
    if not isinstance(obj, dict) or obj.get("format") != BATON_FORMAT:
        raise ValueError("업무바통 파일(.baton) 형식이 아닙니다")
    as_list = lambda v: v if isinstance(v, list) else []
    lineage = [{k: str(x.get(k) or "")[:100] for k in ("name", "handed_to", "date", "work")}
               for x in as_list(obj.get("lineage")) if isinstance(x, dict) and x.get("name")][-BATON_MAX_LINEAGE:]
    # 전임자가 확인 완료하지 않고 저장한 바통이면 모든 항목을 '이전 담당자 미확인'으로 이어받는다
    file_confirmed = obj.get("confirmed") is not False
    holder = str(obj.get("holder") or (lineage[-1]["name"] if lineage else "") or "이전 담당자")[:50]
    gen = len(lineage) or 1
    out, items = [], []
    chain = " → ".join(f"{i}대 {x['name']}" + (f"({x.get('date')})" if x.get("date") else "") for i, x in enumerate(lineage, 1))
    out.append(("계보", f"업무 계보: {chain or holder} → {obj.get('handed_to') or '다음 담당자'} · {obj.get('work', '')}"))
    sections = obj.get("sections") if isinstance(obj.get("sections"), dict) else {}
    for sec, label in BATON_SECTIONS.items():
        for i, raw in enumerate(as_list(sections.get(sec))[:BATON_MAX_ITEMS], 1):
            it = _baton_item(sec, raw)
            if not it:
                continue
            since = raw.get("since") if isinstance(raw.get("since"), dict) else {}
            since = ({"gen": int(since["gen"]), "name": str(since.get("name") or holder)}
                     if str(since.get("gen", "")).isdigit() else {"gen": gen, "name": holder})
            srcs = [s for s in as_list(raw.get("sources")) if isinstance(s, dict)][:2]
            text = _baton_text(sec, it) + (" (근거: " + "; ".join(f"{s.get('file', '')} {s.get('loc', '')}".strip() for s in srcs) + ")" if srcs else "")
            unverified = not file_confirmed or str(raw.get("status") or "") in ("초안", "")
            items.append({"section": sec, "item": it, "block": len(out), "since": since, "unverified": unverified})
            out.append((f"{label} {i}", text))
    interview = []
    for i, x in enumerate(as_list(obj.get("interview"))[:BATON_MAX_ITEMS], 1):
        if isinstance(x, dict) and x.get("q") and x.get("a"):
            x = {"q": str(x["q"])[:BATON_MAX_TEXT], "a": str(x["a"])[:BATON_MAX_TEXT], "by": str(x.get("by") or holder)[:50], "at": str(x.get("at") or "")[:30]}
            interview.append(x)
            out.append((f"질의응답 {i}", f"질문: {x['q']} / 답변({x['by']}): {x['a']}"))
    notes = []
    for i, x in enumerate(as_list(obj.get("notes"))[:BATON_MAX_ITEMS], 1):
        if isinstance(x, dict) and x.get("text"):
            x = {"text": str(x["text"])[:BATON_MAX_TEXT], "by": str(x.get("by") or holder)[:50], "at": str(x.get("at") or "")[:30]}
            notes.append(x)
            out.append((f"전임자 메모 {i}", f"{x['by']} 메모: {x['text']}"))
    meta = {"title": f"{holder} 인수인계서(바통 파일)", "holder": holder, "handed_to": str(obj.get("handed_to") or ""),
            "work": str(obj.get("work") or ""), "lineage": lineage, "baton_items": items,
            "interview": interview, "notes": notes}
    return out, [], meta


READERS = {
    **{e: _read_image for e in IMAGE_EXT},
    ".txt": _read_text, ".md": _read_text, ".csv": _read_csv,
    ".hwpx": _read_hwpx, ".hwp": _read_hwp, ".docx": _read_docx, ".pptx": _read_pptx,
    ".pdf": _read_pdf, ".xlsx": _read_xlsx, ".xlsm": _read_xlsx,
}


def _why(e, ext):
    """읽기 실패 이유를 사람이 알아듣는 말로(영어 오류 문구 대신)."""
    msg = str(e)
    if isinstance(e, zipfile.BadZipFile) or "not a zip file" in msg.lower():
        return f"파일이 깨졌거나 {ext} 형식이 아닙니다(다른 프로그램으로 열어 다시 저장해 보세요)"
    if "password" in msg.lower() or "encrypt" in msg.lower():
        return "암호가 걸린 파일이라 열 수 없습니다"
    if isinstance(e, (json.JSONDecodeError, UnicodeDecodeError)) or "Expecting" in msg:
        return "파일 내용이 깨져 있어 읽을 수 없습니다"
    if "xml" in type(e).__name__.lower() or "XMLSyntax" in msg:
        return "문서 내부 구조가 깨져 있어 읽을 수 없습니다"
    if "EOF" in msg or "startxref" in msg:
        return "PDF가 끝까지 저장되지 않았거나 깨져 있습니다"
    return f"파일을 열 수 없습니다({type(e).__name__})"


def load_bytes(data, relpath):
    """파일 내용(bytes)과 상대경로로 Document를 만든다."""
    relpath = relpath.replace("\\", "/")
    ext = Path(relpath).suffix.lower()
    did = _doc_id(relpath)
    doc = Document(id=did, name=Path(relpath).name, relpath=relpath, ext=ext)
    try:
        if ext in (".eml", ".baton"):
            items, tables, meta = (_read_eml if ext == ".eml" else _read_baton)(data, relpath)
            doc.meta.update(meta)
        elif ext in READERS:
            res = READERS[ext](data, relpath)
            items, tables = res[0], res[1]
            if len(res) > 2:
                doc.meta.update(res[2])
        else:
            doc.error = "지원하지 않는 형식"
            return doc
    except Exception as e:  # 한 파일이 깨져도 나머지는 계속 읽는다
        doc.error = f"읽기 실패: {_why(e, ext)}"
        return doc
    if len(items) > MAX_BLOCKS:  # 아주 큰 파일(로그·덤프 등)은 앞부분만 읽어 전체가 멈추지 않게 한다
        doc.meta["truncated"] = f"내용이 너무 길어 앞부분 {MAX_BLOCKS:,}줄만 읽었습니다(전체 {len(items):,}줄)"
        items = items[:MAX_BLOCKS]
    doc.tables = tables
    if ext in IMAGE_EXT:
        doc.image_bytes = data  # 포토 아카이브용(문서 목록·저장 데이터에는 들어가지 않음)
    for i, (loc, text) in enumerate(items):
        doc.blocks.append(Block(id=f"{did}-{i}", doc_id=did, file=relpath, loc=loc, text=text))
    return doc


def iter_folder(folder):
    """폴더 안 파일을 (상대경로, 절대경로)로 낸다. 숨김·임시 파일은 건너뛴다."""
    folder = Path(folder)
    for p in sorted(folder.rglob("*")):
        if not p.is_file():
            continue
        rel = p.relative_to(folder).as_posix()
        if any(part.startswith((".", "~$")) for part in p.relative_to(folder).parts):
            continue
        yield rel, p


def _zip_name(zi):
    """한국어 Windows에서 만든 ZIP은 파일 이름이 CP949로 들어 있다(UTF-8 표시가 없으면 CP949로 다시 읽음)."""
    if zi.flag_bits & 0x800:
        return zi.filename
    try:
        return zi.filename.encode("cp437").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return zi.filename


def _unzip(data, rel, skipped):
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except UnicodeDecodeError:
        try:  # 이름이 UTF-8이라고 잘못 표시된 압축(일부 압축 프로그램): 한국어 Windows 인코딩으로 다시 읽는다
            z = zipfile.ZipFile(io.BytesIO(data), metadata_encoding="cp949")
        except Exception:
            skipped.append({"file": rel, "reason": "압축 안 파일 이름을 읽지 못함"})
            return
    except Exception:
        skipped.append({"file": rel, "reason": "깨진 압축 파일"})
        return
    total = 0
    for i, zi in enumerate(z.infolist()):
        if zi.is_dir():
            continue
        name = "/".join(x for x in _zip_name(zi).replace("\\", "/").split("/") if x not in ("", ".", ".."))
        if i >= ZIP_MAX_FILES:
            skipped.append({"file": rel, "reason": f"압축 안 파일이 {ZIP_MAX_FILES}개를 넘어 나머지는 건너뜀"})
            return
        if zi.flag_bits & 0x1:
            skipped.append({"file": f"{rel}/{name}", "reason": "암호가 걸린 압축이라 열 수 없음"})
            continue
        total += zi.file_size
        if total > ZIP_MAX_BYTES:
            skipped.append({"file": rel, "reason": "압축을 푼 크기가 너무 커서(300MB 초과) 나머지는 건너뜀"})
            return
        try:
            yield f"{rel}/{name}", z.open(zi).read(zi.file_size + 1)
        except Exception as e:
            skipped.append({"file": f"{rel}/{name}", "reason": f"압축 풀기 실패: {e}"})


COPY_RX = re.compile(r"복사본|사본|copy|\(\d+\)\.[^.]+$", re.I)


def load_entries(entries):
    """(상대경로, 내용 읽는 함수) 목록을 읽는다. 뒤죽박죽 자료를 정리하며 읽는다:
    ZIP은 풀어서(두 단계까지), 같은 내용 파일·빈 파일·임시/시스템 파일은 빼고, 옛 형식은 바꾸는 방법을 안내한다."""
    import hashlib as _h
    docs, skipped, seen = [], [], {}
    _ocr_budget[0] = OCR_LIMIT

    def handle(rel, data, depth):
        name = rel.rsplit("/", 1)[-1]
        ext = Path(name).suffix.lower()
        if name.lower() in JUNK or name.startswith(("~$", "._")):
            skipped.append({"file": rel, "reason": "임시·시스템 파일(자동으로 생긴 파일)이라 건너뜀"})
            return
        if not data:
            skipped.append({"file": rel, "reason": "빈 파일"})
            return
        if ext == ".zip":
            if depth >= 2:
                skipped.append({"file": rel, "reason": "압축 안의 압축은 두 단계까지만 풉니다"})
                return
            for inner, inner_data in list(_unzip(data, rel, skipped)):
                handle(inner, inner_data, depth + 1)
            return
        if ext not in SUPPORTED:
            skipped.append({"file": rel, "reason": skip_reason(ext)})
            return
        digest = _h.sha1(data).hexdigest()
        if digest in seen:
            old = seen[digest]
            if COPY_RX.search(old.rsplit("/", 1)[-1]) and not COPY_RX.search(name):
                # 먼저 읽은 쪽이 '복사본'이면 원래 이름 파일을 남긴다
                docs[:] = [x for x in docs if x.relpath != old]
                skipped[:] = [x for x in skipped if x["file"] != old]
                skipped.append({"file": old, "reason": f"같은 내용의 파일이 이미 있어 건너뜀({rel})"})
            else:
                skipped.append({"file": rel, "reason": f"같은 내용의 파일이 이미 있어 건너뜀({old})"})
                return
        seen[digest] = rel
        doc = load_bytes(data, rel)
        if doc.error:
            skipped.append({"file": rel, "reason": doc.error})
        elif doc.meta.get("truncated"):
            skipped.append({"file": rel, "reason": doc.meta["truncated"]})
        docs.append(doc)

    for rel, read in entries:
        handle(rel, read(), 0)
    return docs, skipped


def load_folder(folder):
    return load_entries((rel, (lambda p=p: p.read_bytes())) for rel, p in iter_folder(folder))
