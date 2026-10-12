"""과제① 업무바통 모의데이터 생성 (모두 가상의 기관·인물·연락처).

전화번호는 실제 번호와 겹치지 않도록 국번 000(02-000-xxxx, 031-000-xxxx)·휴대전화 010-0000-xxxx 대역만 쓴다.
"""
import email.message
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from core.export import to_hwpx, to_docx  # noqa: E402

OUT = ROOT / "sample_data" / "전임자_업무폴더_시설운영"


def w(rel, data):
    p = OUT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))


def safety_plan():
    return to_hwpx([
        ("title", "2026년 체육센터 안전점검 계획"),
        ("p", "공공기관 시설운영팀 (문서번호: 시설운영팀-312, 2026. 1. 20.)"),
        ("h1", "1. 목적"),
        ("p", "체육센터 이용자의 안전 확보를 위하여 시설물 정기·수시 안전점검을 체계적으로 실시하고자 함"),
        ("h1", "2. 점검 일정"),
        ("bullet", "해빙기 안전점검: 2월 중 자체점검 실시 (자체점검표 활용)"),
        ("bullet", "상반기 정기 안전점검: 2026. 3. 16. ~ 3. 27. 외부 전문기관 위탁"),
        ("bullet", "우기 대비 점검: 6월 말까지 배수로·옥상 방수 점검"),
        ("bullet", "하반기 정기 안전점검: 2026. 9. 14. ~ 9. 25. 외부 전문기관 위탁"),
        ("bullet", "동절기 대비 점검: 11월 중 보일러·동파 취약 배관 점검"),
        ("bullet", "수영장 수질검사: 매월 10일까지 검사 결과를 구청 체육진흥과에 제출"),
        ("h1", "3. 결과 보고"),
        ("p", "정기 안전점검 완료 후 14일 이내 이사장에게 결과를 보고하고, 구청 체육진흥과 박서연 주무관(02-000-3345)에게 통보한다."),
        ("h1", "4. 추진 체계"),
        ("table", ["구분", "담당", "연락처"], [
            ["총괄", "시설운영팀장 이정훈", "02-000-7710"],
            ["실무", "시설운영팀 김도윤 주임", "02-000-7714"],
            ["위탁기관", "(주)한결안전기술 최현우 과장", "031-000-2201"],
            ["관계기관", "구청 체육진흥과 박서연 주무관", "02-000-3345"],
        ]),
    ])


def contracts():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "계약현황"
    ws.append(["2026년 시설 유지보수 용역 계약 현황 (시설운영팀)"])
    ws.append([])
    ws.append(["연번", "계약명", "업체명", "업체 담당자", "연락처", "계약기간", "계약금액(원)", "진행상황", "비고"])
    rows = [
        [1, "승강기 유지관리 용역", "(주)대한엘리베이터서비스", "정민재 대리", "02-000-1830", "2026.01.01~2026.12.31", 14400000, "진행 중", "매월 정기점검, 12월 재계약 필요"],
        [2, "소방시설 종합정밀점검", "(주)안전소방", "한지훈 과장", "02-000-6612", "2026.04.01~2026.04.30", 3850000, "완료", "결과보고서 5월 제출 완료"],
        [3, "냉난방기 교체공사", "(주)그린공조", "오세영 팀장", "031-000-7890", "2026.10.20~2026.12.10", 86000000, "착공 준비", "착공일 업체와 협의 중"],
        [4, "수영장 여과기 정비", "(주)맑은물테크", "서지우 대리", "02-000-9087", "2026.07.01~2026.08.31", 12100000, "완료", "하자보수기간 1년(2027.8.31.까지)"],
        [5, "청소 용역", "(주)깨끗한세상", "윤나래 실장", "02-000-2244", "2026.01.01~2026.12.31", 182000000, "진행 중", "11월 중 2027년 입찰공고 필요"],
    ]
    for r in rows:
        ws.append(r)
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def water_pdf():
    import io
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(TTFont("K", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", subfontIndex=0))
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    lines = [
        ("수영장 수질검사 결과 보고 (2026년 9월)", 15),
        ("공공기관 시설운영팀", 10),
        ("", 10),
        ("1. 검사 개요", 12),
        ("검사일: 2026. 9. 8.   검사기관: 보건환경연구원", 10),
        ("대상: 체육센터 성인풀, 유아풀", 10),
        ("", 10),
        ("2. 검사 결과", 12),
        ("유리잔류염소: 0.3mg/L (기준 0.4~1.0mg/L) - 부적합", 10),
        ("탁도: 1.4 NTU (기준 2.8 NTU 이하) - 적합", 10),
        ("수소이온농도: 7.2 (기준 5.8~8.6) - 적합", 10),
        ("", 10),
        ("3. 조치 사항", 12),
        ("염소 투입량 조정 후 2026. 9. 22. 재검사 실시 결과 적합 판정", 10),
        ("10월 검사 시 탁도 추이를 재확인할 것 (여과기 정비 후 첫 가을철)", 10),
        ("", 10),
        ("4. 향후 계획", 12),
        ("다음 정기검사: 2026. 10. 6. 예정", 10),
        ("결과는 매월 10일까지 구청 체육진흥과에 제출", 10),
    ]
    y = 800
    for text, size in lines:
        c.setFont("K", size)
        c.drawString(60, y, text)
        y -= size + 10
    c.save()
    return buf.getvalue()


def mail(frm, to, date, subject, body):
    m = email.message.EmailMessage()
    m["From"], m["To"], m["Date"], m["Subject"] = frm, to, date, subject
    m.set_content(body)
    return m.as_bytes()


MEMO = """인수인계 메모 (김도윤, 개인 정리용)

- 냉난방기 교체공사 착공은 10월 27일로 알고 있음. 업체랑 한번 더 확인 필요
- 청소용역 2027년 입찰공고는 아마 11월 중순? 계약팀 강민호 주임이랑 협의해야 함 (내선 7731)
- 승강기 업체 정민재 대리 번호 바뀜 010-0000-9081
- 수영장 수질 9월에 염소 부적합 나와서 재검사 적합 받음. 10월엔 탁도 꼭 다시 볼 것
- 매년 12월에 시설물 보험 갱신 (총무팀 문가은 주임과 협의). 올해 보험사는 미정
- 민원 많은 것: 샤워실 온수, 주차장. 민원 답변은 접수 후 7일 이내
- 해빙기 점검은 매년 2월. 자체점검표 양식은 공유폴더 '안전점검' 폴더에 있음
- 승강기 유지관리 재계약은 12월 초에 계약팀에 의뢰해야 1월 공백 없음
- 예산 전용·변경은 결재 올리기 전에 경영지원팀 정하은 주임과 사전 협의 필수 (협의 없이 올리면 반려됨)
- 그린공조는 일정이 자주 바뀌니까 마감 3일 전에 꼭 재확인
- 구의회 행정사무감사(11월) 전에 의회 전문위원실 자료요구가 옴. 시설 안전점검 결과·민원 처리 현황은 미리 정리해 둘 것
- 샤워실 온수 민원은 같은 이용자가 반복 제기함. 회신 전에 팀장 보고, 통화 내용은 민원 대장에 기록
"""


def rnr_docx():
    return to_docx([
        ("title", "시설운영팀 업무분장표 (2026. 1. 1. 기준)"),
        ("table", ["담당자", "직위", "담당 업무"], [
            ["이정훈", "팀장", "시설운영팀 업무 총괄 / 예산 및 계약 결재"],
            ["김도윤", "주임", "체육센터 시설물 유지관리\n안전점검(해빙기·정기·우기·동절기) 계획 및 시행\n시설 유지보수 용역 계약 관리\n수영장 수질관리 및 결과 보고\n시설 관련 민원 처리\n소관 예산 요구 및 집행"],
            ["박하늘", "주임", "문화센터 시설물 유지관리 / 주차장 운영"],
        ]),
        ("p", "※ 김도윤 주임 부재 시 대행: 박하늘 주임"),
    ])


def contacts_xlsx():
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "연락처"
    ws.append(["이름", "소속", "직위", "연락처", "이메일", "주요 협의 업무"])
    for r in [
        ["박서연", "구청 체육진흥과", "주무관", "02-000-3345", "psy@public.example", "수질검사 결과 보고, 안전점검 결과 통보"],
        ["최현우", "(주)한결안전기술", "과장", "031-000-2201", "choi@hangyeol.example", "정기 안전점검 위탁"],
        ["강민호", "기관 계약팀", "주임", "02-000-7731", "", "입찰·계약 의뢰"],
        ["정하은", "기관 경영지원팀", "주임", "02-000-7702", "budget@public.example", "예산 요구·배정"],
        ["문가은", "기관 총무팀", "주임", "02-000-7705", "", "시설물 보험 갱신"],
        ["오세영", "(주)그린공조", "팀장", "031-000-7890", "oh.sy@greenhvac.example", "냉난방기 교체공사"],
        ["정민재", "(주)대한엘리베이터서비스", "대리", "02-000-1830", "", "승강기 유지관리"],
    ]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def relay_baton():
    """지식 릴레이 시연용: 1대 담당자(최선우)가 2024년에 김도윤에게 넘긴 바통 파일."""
    import json
    src = lambda f, loc, q: [{"file": f, "loc": loc, "quote": q}]
    since = {"gen": 1, "name": "최선우"}
    data = {
        "format": "baton-handover/1", "app": "업무바통", "exported_at": "2024-06-28 17:40",
        "work": "시설운영팀 시설관리 업무", "holder": "최선우", "handed_to": "김도윤", "confirmed": True,
        "lineage": [{"gen": 1, "name": "최선우", "handed_to": "김도윤", "date": "2024-07-01", "work": "시설운영팀 시설관리 업무"}],
        "sections": {
            "rnr": [
                {"duty": "수영장 수질관리 및 결과 보고", "detail": "부적합 판정이 나오면 당일 염소 투입량을 조정하고 재검사를 의뢰",
                 "status": "확인", "since": since, "sources": src("수질관리 업무노트.hwp", "문단 4", "부적합 시 당일 염소 투입량 조정 후 재검사 의뢰")},
                {"duty": "시설 관련 민원 처리", "detail": "겨울철 샤워실 온수 민원이 몰리므로 11월에 보일러 점검을 먼저",
                 "status": "확인", "since": since, "sources": src("민원 처리 대장 2023.xlsx", "민원!12행", "샤워실 온수 미지급 민원 (12~2월 집중)")},
            ],
            "schedule": [
                {"task": "정기 안전점검 외부 전문기관 위탁 계약 체결(상·하반기 점검 대비)", "months": [1], "day": None,
                 "recurring": "매년", "when": "매년 1월", "related": "정기 안전점검", "action": "계약", "status": "확인", "since": since,
                 "sources": src("2024년 체육시설 안전점검 계획.hwpx", "문단 9", "정기 안전점검은 1월 중 외부 전문기관과 위탁계약 체결")},
                {"task": "수영장 정기 휴장 기간 여과기·배관 점검", "months": [7], "day": None, "recurring": "매년",
                 "when": "매년 7월(휴장 기간)", "related": "수영장", "action": "점검", "status": "확인", "since": since,
                 "sources": src("수영장 운영 계획.docx", "문단 15", "7월 정기 휴장 기간 중 여과기 및 배관 점검")},
                {"task": "수영장 수질검사 결과를 구청 체육진흥과에 제출", "months": list(range(1, 13)), "day": 10, "recurring": "매월",
                 "when": "매월 10일", "related": "", "action": "제출", "status": "확인", "since": since,
                 "sources": src("2024년 체육시설 안전점검 계획.hwpx", "문단 21", "수영장 수질검사: 매월 10일까지 결과 제출")},
            ],
            "contacts": [
                {"name": "이정훈", "title": "팀장", "org": "시설운영팀", "phones": ["02-000-7710"], "emails": [],
                 "topics": ["결재권자", "예산·계약 사전 협의"], "status": "확인", "since": since,
                 "sources": src("시설운영팀 업무분장표.docx", "표1 2행", "팀장 이정훈 / 팀 업무 총괄·결재")},
            ],
            "issues": [
                {"title": "주차장 바닥 균열 보수", "state": "예산 요구 중", "next_action": "2025년 예산 반영 결과를 보고 보수공사 발주",
                 "due": "", "topic": "주차장", "status": "확인", "since": since,
                 "sources": src("2025년 예산요구서(시설운영팀).xlsx", "요구!7행", "주차장 바닥 균열 보수 18,000천원")},
            ],
        },
        "interview": [
            {"q": "수질검사에서 부적합이 나오면 어떻게 하나요?",
             "a": "당일 염소 투입량을 조정하고 보건환경연구원에 재검사를 의뢰하세요. 결과는 구청 체육진흥과에 바로 알리면 됩니다.",
             "by": "최선우", "at": "2024-06-27 10:12"},
        ],
        "notes": [
            {"text": "공유폴더 '안전점검' 폴더의 자체점검표 양식은 매년 1월에 바뀐 기준으로 고쳐 두세요.", "by": "최선우", "at": "2024-06-28 17:35"},
        ],
        "note": "업무바통 '자료 넣기'에 이 파일을 함께 넣으면 다음 담당자가 이 내용을 근거로 이어받습니다.",
    }
    return json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8")


# ---------------------------------------------------------------- 뒤죽박죽 최소자료(개떡같이 넘겨도 찰떡같이)

MESSY = ROOT / "sample_data" / "전임자_뒤죽박죽_최소자료"
MEMO_MESSY = "업체 연락처는 엑셀 참고(없으면 계약서 보기)\n정산은 11월 말까지\n모르면 연락 주세요 010-0000-3321 (박정우)\n복합기 업체는 계약 끝나기 한 달 전에 꼭 연락해야 함(자동연장 조항 있음)\n"


def _font(size):
    """사진 흉내용 글꼴: 맑은 고딕 → Pretendard(OTF) 순으로 찾는다(생성할 때만 필요)."""
    from PIL import ImageFont
    import glob
    for cand in [r"C:\Windows\Fonts\malgun.ttf", *glob.glob(str(ROOT.parent / "**" / "Pretendard-Regular.otf"), recursive=True),
                 *glob.glob("/tmp/**/Pretendard-Regular.otf", recursive=True), "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"]:
        try:
            return ImageFont.truetype(cand, size)
        except OSError:
            continue
    raise SystemExit("한글 글꼴을 찾지 못했습니다(맑은 고딕 또는 Pretendard OTF 필요)")


def _photo(lines, size=(1280, 900), bg=(236, 236, 230), date="2026:09:15 09:30:12", rotate=-1.2, fmt="JPEG"):
    """화이트보드·메모지를 휴대폰으로 찍은 듯한 사진(약간 기울어짐, 찍은 날짜 EXIF)."""
    import io
    from PIL import Image, ImageDraw, ImageFilter
    im = Image.new("RGB", size, bg)
    d = ImageDraw.Draw(im)
    y = 70
    for i, line in enumerate(lines):
        f = _font(46 if i == 0 else 36)
        d.text((70, y), line, fill=(30, 30, 40) if i else (20, 40, 120), font=f)
        y += 78 if i == 0 else 64
    im = im.rotate(rotate, expand=False, fillcolor=bg).filter(ImageFilter.GaussianBlur(0.6))
    buf = io.BytesIO()
    if fmt == "JPEG":
        exif = Image.Exif()
        exif[306] = date
        exif.get_ifd(0x8769)[36867] = date
        im.save(buf, "JPEG", quality=88, exif=exif)
    else:
        im.save(buf, fmt, resolution=150)
    return buf.getvalue()


def _exif_jpeg(im, date):
    import io
    from PIL import Image
    buf = io.BytesIO()
    exif = Image.Exif()
    exif[306] = date
    exif.get_ifd(0x8769)[36867] = date
    im.save(buf, "JPEG", quality=86, exif=exif)
    return buf.getvalue()


def scene_training():
    """교육장 현장 사진 흉내: 화면에 과정 이름, 책상 줄, 참석자는 얼굴 없는 실루엣."""
    from PIL import Image, ImageDraw, ImageFilter
    im = Image.new("RGB", (1280, 860), (222, 226, 232))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 560, 1280, 860], fill=(176, 160, 140))                  # 바닥
    d.rectangle([300, 70, 980, 420], fill=(250, 250, 252), outline=(90, 90, 100), width=6)   # 스크린
    d.text((350, 120), "2026 상반기 직무중심교육", fill=(20, 60, 160), font=_font(54))
    d.text((350, 210), "예산·회계 실무 과정", fill=(40, 40, 50), font=_font(42))
    d.text((350, 290), "5. 20.(수) 14:00 · 3층 교육장", fill=(70, 70, 80), font=_font(34))
    for row, y in enumerate((600, 700, 790)):
        for x in range(90, 1200, 230):
            d.ellipse([x + 60, y - 70, x + 120, y - 10], fill=(70, 74, 86))   # 머리(뒷모습)
            d.rectangle([x + 40, y - 12, x + 140, y + 40], fill=(70, 74, 86))
            d.rectangle([x, y + 20, x + 190, y + 50], fill=(120, 96, 70))      # 책상
    return _exif_jpeg(im.filter(ImageFilter.GaussianBlur(0.8)), "2026:05:20 14:12:05")


def scene_parking_crack():
    """주차장 바닥 균열 사진 흉내: 회색 바닥, 갈라진 선, 라바콘, 기둥 번호."""
    from PIL import Image, ImageDraw, ImageFilter
    import random
    random.seed(7)
    im = Image.new("RGB", (1280, 860), (150, 150, 146))
    d = ImageDraw.Draw(im)
    d.rectangle([980, 0, 1180, 860], fill=(200, 200, 196))                  # 기둥
    d.rectangle([1010, 120, 1150, 220], fill=(250, 210, 40))
    d.text((1028, 140), "B1 C-3", fill=(30, 30, 30), font=_font(40))
    x, y = 120, 760
    pts = [(x, y)]
    for _ in range(14):
        x += random.randint(40, 70)
        y -= random.randint(15, 55)
        pts.append((x, y))
    d.line(pts, fill=(40, 40, 40), width=7)
    d.line([(p[0] + 4, p[1] + 3) for p in pts[5:11]], fill=(60, 60, 60), width=3)
    d.polygon([(860, 700), (900, 520), (940, 700)], fill=(240, 110, 20))     # 라바콘
    d.rectangle([840, 700, 960, 720], fill=(30, 30, 30))
    d.text((60, 40), "주차장 바닥 균열 — B1 C-3 기둥 앞 (2026.9.3)", fill=(250, 250, 250), font=_font(36))
    return _exif_jpeg(im.filter(ImageFilter.GaussianBlur(0.7)), "2026:09:03 10:41:20")


def scene_shower_pipe():
    """샤워실 온수 배관 점검 사진 흉내: 타일 벽, 배관, 점검 꼬리표."""
    from PIL import Image, ImageDraw, ImageFilter
    im = Image.new("RGB", (1280, 860), (226, 232, 236))
    d = ImageDraw.Draw(im)
    for x in range(0, 1280, 80):
        d.line([(x, 0), (x, 860)], fill=(200, 208, 214), width=3)
    for y in range(0, 860, 80):
        d.line([(0, y), (1280, y)], fill=(200, 208, 214), width=3)
    d.rectangle([0, 300, 1280, 340], fill=(184, 115, 51))                    # 온수관(구리)
    d.rectangle([0, 420, 1280, 455], fill=(150, 160, 170))                   # 냉수관
    d.rectangle([600, 270, 640, 470], fill=(120, 120, 130))                  # 밸브
    d.rectangle([700, 360, 1080, 560], fill=(255, 250, 220), outline=(120, 100, 40), width=4)
    d.text((725, 380), "남자 샤워실 온수 배관", fill=(30, 30, 30), font=_font(34))
    d.text((725, 440), "점검 2026.10.5", fill=(30, 30, 30), font=_font(32))
    d.text((725, 495), "수온 38℃ (기준 40℃ 미달)", fill=(180, 30, 30), font=_font(30))
    return _exif_jpeg(im.filter(ImageFilter.GaussianBlur(0.6)), "2026:10:05 07:55:41")


def add_photos():
    """포토 아카이브 시연용 사진(두 모의자료에 각 2장)."""
    w("04_현장사진/주차장 바닥 균열(2026.9.3).jpg", scene_parking_crack())
    w("04_현장사진/샤워실 온수 배관 점검.jpg", scene_shower_pipe())
    (MESSY / "2026 상반기 직무중심교육_현장.jpg").write_bytes(scene_training())
    (MESSY / "직무중심교육 결과 메모.jpg").write_bytes(_photo([
        "직무중심교육 결과 (5/20)",
        "- 예산·회계 실무 과정, 참석 10명 / 대상 12명",
        "- 만족도 4.6점, 강사료 정산 완료",
        "- 하반기 과정 11월 중 실시 필요 (미실시)"], date="2026:05:20 16:40:00", rotate=0.9))


def messy_budget_xlsx():
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["2026년 기획예산팀 예산 집행현황"])
    ws.append(["(단위: 천원)"])
    ws.append([])
    ws.append(["사업명", "예산액", "집행액", "잔액", "비고"])
    for r in [["업무추진비", 12000, 8400, 3600, "매월 10일 홈페이지 공개"],
              ["청사 경비용역", 180000, 135000, 45000, "12월 정산"],
              ["복합기 임차", 9600, 8000, 1600, "11월 말 계약 종료, 재계약 검토 필요"],
              ["직원 교육", 5000, 1200, 3800, "하반기 교육 미실시, 11월 중 실시 필요"]]:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def messy_report_hwpx():
    return to_hwpx([
        ("title", "기획예산팀 회계 업무 현황(초안)"),
        ("p", "작성: 기획예산팀 박정우 주무관"),
        ("bullet", "담당 업무: 예산 편성·집행 관리 / 업무추진비 공개 / 계약 관리(경비·임차)"),
        ("bullet", "다음 달 할 일: 2027년 예산 2차 요구서 11월 14일까지 제출"),
        ("bullet", "관련 규정: 지방재정법, 지방자치단체 회계관리에 관한 훈령, 업무추진비 집행 규칙"),
        ("bullet", "시스템: e호조(회계), 온나라(문서). 자료는 공유폴더 \\\\nas\\기획예산\\회계 에 보관"),
    ])


def messy_rnr_docx():
    return to_docx([
        ("title", "기획예산팀 업무분장(2025)"),
        ("table", ["담당자", "직위", "담당 업무"], [
            ["이서진", "팀장", "팀 업무 총괄 / 예산 결재"],
            ["박정우", "주무관", "예산 편성·집행 관리\n업무추진비 공개\n위탁·임차 계약 관리"],
        ]),
    ])


def messy_zip():
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("업무분장_2025.docx", messy_rnr_docx())
        z.writestr("메모 - 복사본.txt", MEMO_MESSY)   # 바깥 메모와 같은 내용(중복)
        z.writestr("Thumbs.db", b"\x00" * 64)       # 윈도우가 만든 미리보기 파일
    return buf.getvalue()


def messy():
    """최소한의 자료를 아무렇게나 모은 폴더: 휴대폰 사진, 스캔 PDF, 제목 줄이 엉망인 엑셀, 한글, ZIP, 중복·빈 파일, 옛 형식."""
    def put(rel, data):
        p = MESSY / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    xlsx = messy_budget_xlsx()
    put("IMG_20260915_093012.jpg", _photo([
        "9월 업무 메모 (화이트보드)",
        "- 2027 예산 1차 요구: 10/31까지 기획예산팀 제출",
        "- 업무추진비 공개: 매월 10일 홈페이지 게시",
        "- 회계 마감 12/20, 이월 신청 12/15까지",
        "- 문의: 회계팀 한서윤 주무관 내선 2471"]))
    put("스캔_위탁계약 현황.pdf", _photo([
        "2026 위탁·임차 계약 현황",
        "청사 경비용역 / (주)든든경비 / 2026.01.01~2026.12.31",
        "  담당: 오민석 과장 031-000-4410",
        "복합기 임차 / 한빛오피스 / ~2026.11.30 (재계약 필요)",
        "  담당: 윤가람 대리 02-000-5512"], bg=(250, 250, 248), rotate=0.6, fmt="PDF"))
    put("예산 집행현황(최종)(진짜최종).xlsx", xlsx)
    put("예산 집행현황(최종)(진짜최종) - 복사본.xlsx", xlsx)   # 같은 파일 복사본
    put("업무 보고(초안).hwpx", messy_report_hwpx())
    put("메모.txt", MEMO_MESSY)
    put("옛날자료.zip", messy_zip())
    put("빈 문서.txt", b"")
    put("옛날 예산.xls", b"\xd0\xcf\x11\xe0" + b"\x00" * 60)   # 엑셀 97~2003 형식(읽지 못함 → 안내)
    put("회의 사진.heic", b"\x00\x00\x00\x18ftypheic" + b"\x00" * 40)
    print("뒤죽박죽 모의자료:", MESSY)


def main():
    w("00_이전인수인계/시설운영_1대_인계.baton", relay_baton())
    w("01_공식문서/2026년 체육시설 안전점검 계획.hwpx", safety_plan())
    w("01_공식문서/2026년 시설 유지보수 용역 계약 현황.xlsx", contracts())
    w("01_공식문서/수영장 수질검사 결과 보고(2026년 9월).pdf", water_pdf())
    w("01_공식문서/시설운영팀 업무분장표.docx", rnr_docx())
    w("02_메일/RE_ 냉난방기 교체공사 착공일 협의.eml", mail(
        "오세영 <oh.sy@greenhvac.example>", "김도윤 <kdy@public.example>",
        "Tue, 29 Sep 2026 14:12:00 +0900", "RE: 냉난방기 교체공사 착공일 협의",
        "김도윤 주임님, 안녕하세요. (주)그린공조 오세영 팀장입니다.\n\n"
        "말씀하신 일정 검토 결과 착공일은 2026년 11월 3일(화)로 확정하고자 합니다.\n"
        "공사 기간 중 3층 다목적실 사용 중지가 필요합니다.\n"
        "이용자 안내를 위해 10월 24일까지 공사 안내문 게시를 부탁드립니다.\n"
        "현장 대리인 연락처: 오세영 팀장 010-0000-7788\n\n감사합니다.\n"))
    w("02_메일/2027년 예산요구서 제출 요청.eml", mail(
        "정하은 <budget@public.example>", "시설운영팀 <facility@public.example>",
        "Mon, 21 Sep 2026 09:30:00 +0900", "[요청] 2027년 소관 예산요구서 제출",
        "시설운영팀 담당자님께\n\n2027년 예산요구서를 2026년 10월 16일까지 제출하여 주시기 바랍니다.\n"
        "냉난방기 교체 2단계(4~5층) 예산 반영 여부를 검토하여 주십시오.\n"
        "예산 확정은 12월 이사회 의결 후 통보 예정입니다.\n\n경영지원팀 정하은 주임 (02-000-7702)\n"))
    w("03_개인메모/인수인계 메모.txt", MEMO)
    w("03_개인메모/연락처 정리.xlsx", contacts_xlsx())
    print("업무바통 모의데이터:", OUT)


def main_all():
    main()
    messy()
    add_photos()


if __name__ == "__main__":
    if "--messy" in sys.argv:
        messy()
    elif "--photos" in sys.argv:
        add_photos()
    else:
        main_all()
